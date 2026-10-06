from __future__ import annotations

import sys
import tempfile
import unittest
import py_compile
import re
from collections import OrderedDict
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "public" / "py"))

from rd_cost_workflow_v1.engine import (  # noqa: E402
    LedgerRecord,
    OUTPUT_TEMPLATE,
    Project,
    SalaryRow,
    allocate_salary,
    build_application_candidates,
    allocate_by_displayed_rates,
    calculate_legacy_hours,
    classify_input_workbooks,
    combination_filename_matches,
    copy_summary_to_template,
    displayed_rates_with_residual,
    fee_source_content_last_row,
    hours_workbook_month,
    locate_combination_summary_sheet,
    normalize_combination_workbook,
    parse_combination_summary_sheet,
    quote_internal_sheet_references,
    read_expense_detail_workbooks,
    read_dimension_summary_workbook,
    parse_ledger_file,
    relocate_freight_visible_table,
    relocate_material_allocation_tables,
    relocate_standard_fee_allocation_table,
    standard_fee_allocation_table,
    write_template_expense_source,
    write_template_hours_sheets,
    write_template_material_allocation,
    write_visible_freight_allocation_table,
)


def records(fee_code: str, fee_name: str, count: int) -> list[LedgerRecord]:
    return [
        LedgerRecord(
            source_path=Path("input.xlsx"),
            sheet="Sheet1",
            fee_code=fee_code,
            fee_name=fee_name,
            project_code="RD141232",
            project_name="80V低压驱动器研发",
            posting_date="2026-01-01",
            fiscal_year=2026,
            period=1,
            voucher_word="记",
            voucher_no=index + 1,
            summary=f"测试明细{index + 1}",
            debit=Decimal("1.00"),
            credit=Decimal("0"),
            source_row=index + 2,
        )
        for index in range(count)
    ]


def make_categorized_detail_ledger(path: Path, fee_code: str, fee_name: str, amount: int | float) -> None:
    source_book = Workbook()
    sheet = source_book.active
    sheet.title = "Sheet1"
    sheet.append(["账簿:苏州溯驭", "币别:人民币", "期间:2026年01月01日 至 2026年01月31日"])
    sheet.append([
        "科目范围:[5301] 研发支出",
        f"核算维度:[费用项目]-[{fee_code}]{fee_name}/[研发项目]-[RD9999]研发公摊",
    ])
    sheet.append(["日期", "会计年度", "期间", "科目编码", "科目名称", "凭证字", "凭证号", "摘要", "借方", "贷方"])
    sheet.append(["2026/1/31", 2026, 1, "5301", "研发支出", "记", 1, f"{fee_name}测试", amount, None])
    source_book.save(path)


class PlatformRegressionTests(unittest.TestCase):
    def test_application_candidates_keep_parent_project_consolidated_in_ledger(self) -> None:
        candidates = build_application_candidates(
            OrderedDict([("RD100", Project("RD100", "母项目", Decimal("12")))]),
            {}, [], OrderedDict([("RD100-01", Project("RD100-01", "子项目"))]),
        )
        self.assertEqual(candidates[0]["project_code"], "RD100")
        self.assertEqual(candidates[0]["status"], "待确认")
        self.assertIn("子项目归集", candidates[0]["reason"])

    def test_browser_generated_calculation_script_compiles(self) -> None:
        """Prevent indentation regressions inside the R&D workbench Python template."""
        page_source = (PROJECT_ROOT / "app" / "rd-cost" / "page.tsx").read_text(encoding="utf-8")
        match = re.search(
            r"await runtime\.runPythonAsync\(\s*`(.*?)`\s*,?\s*\);",
            page_source,
            flags=re.DOTALL,
        )
        self.assertIsNotNone(match, "未找到浏览器端测算 Python 模板")
        calculation_source = (
            match.group(1)
            .replace(r"\n", "\n")
            .replace('${paths.join(",")}', 'Path("/work/input_1.xlsx")')
            .replace('${accountingYear}', '2025')
        )
        compile(calculation_source, "browser_calculation_template.py", "exec")

    def test_browser_rule_python_sources_compile(self) -> None:
        """网页端会直接加载这些源码；发布前必须拦截任何 Python 语法错误。"""
        for source in (
            PROJECT_ROOT / "public" / "py" / "rd_cost_workflow_v1" / "engine.py",
            PROJECT_ROOT / "public" / "py" / "rd_hours_v1" / "engine.py",
        ):
            with self.subTest(source=source.name):
                py_compile.compile(str(source), doraise=True)

    def test_freight_uses_displayed_rates_and_puts_residual_in_last_project(self) -> None:
        allocations = allocate_by_displayed_rates(
            Decimal("1433.02"),
            ["RD-A", "RD-B", "RD-C"],
            {
                "RD-A": Decimal("0.0733058939"),
                "RD-B": Decimal("0.0688091299"),
                "RD-C": Decimal("0.8578849762"),
            },
        )

        self.assertEqual(allocations["RD-A"], Decimal("105.04"))
        self.assertEqual(allocations["RD-B"], Decimal("98.59"))
        self.assertEqual(allocations["RD-C"], Decimal("1229.39"))
        self.assertEqual(sum(allocations.values()), Decimal("1433.02"))

        residual_allocations = allocate_by_displayed_rates(
            Decimal("1.00"),
            ["RD-A", "RD-B", "RD-C"],
            {
                "RD-A": Decimal("0.3333"),
                "RD-B": Decimal("0.3333"),
                "RD-C": Decimal("0.3334"),
            },
        )
        self.assertEqual(
            residual_allocations,
            {
                "RD-A": Decimal("0.33"),
                "RD-B": Decimal("0.33"),
                "RD-C": Decimal("0.34"),
            },
        )

    def test_displayed_rates_put_percentage_residual_in_last_project(self) -> None:
        rates = displayed_rates_with_residual(
            ["RD-A", "RD-B", "RD-C"],
            {"RD-A": Decimal("0.02031"), "RD-B": Decimal("0.01369"), "RD-C": Decimal("0.96599")},
        )

        self.assertEqual(rates, {"RD-A": Decimal("0.0203"), "RD-B": Decimal("0.0137"), "RD-C": Decimal("0.9660")})
        self.assertEqual(sum(rates.values()), Decimal("1"))

    def test_combination_filename_recognition_uses_keywords_not_full_name(self) -> None:
        self.assertTrue(combination_filename_matches(Path("2026核算维度_费用科目组合表_最终版.xlsx")))
        self.assertTrue(combination_filename_matches(Path("客户A核算收入与组合数据.xlsx")))
        self.assertFalse(combination_filename_matches(Path("科目组合.xlsx")))

    def test_hours_month_reads_chinese_month_header(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "hours.xlsx"
            source_book = Workbook()
            sheet = source_book.active
            sheet.title = "工时合计"
            sheet.append(["序号", "姓名", "六月", "RD141232"])
            source_book.save(source_path)

            month = hours_workbook_month(source_path, "工时合计")

        self.assertEqual(month, 6)

    def test_explicit_detail_period_drives_output_month_metadata(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "freight_detail.xlsx"
            source_book = Workbook()
            sheet = source_book.active
            sheet.title = "FYXM02_SYS_运费_RD9999_研发公摊"
            sheet["A1"] = "期间:2026年07月01日 至 2026年07月31日"
            sheet.append(["日期", "会计年度", "期间", "凭证字", "凭证号", "摘要", "借方"])
            sheet.append(["2026/7/1", 2026, 7, "记", 1, "运费测试", 100])
            source_book.save(source_path)

            detail_rows, month = read_expense_detail_workbooks([source_path], [])

        self.assertEqual(month, 7)
        self.assertEqual(len(detail_rows), 1)

    def test_internal_sheet_references_are_always_quoted(self) -> None:
        workbook = Workbook()
        source = workbook.active
        source.title = "研发项目主表"
        target = workbook.create_sheet("研发明细表")
        target["A1"] = "=研发项目主表!C21"
        target["A2"] = "='研发项目主表'!C21"

        quote_internal_sheet_references(workbook)

        self.assertEqual(target["A1"].value, "='研发项目主表'!C21")
        self.assertEqual(target["A2"].value, "='研发项目主表'!C21")

    def test_normalized_combination_workbook_is_valid_dimension_source(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "normalized_combination.xlsx"
            source_book = Workbook()
            sheet = source_book.active
            sheet.title = "Sheet"
            sheet.append(["费用项目", None, None, "研发项目", "费用金额（元）"])
            sheet.append(["FYXM02_SYS  运费", None, None, "RD9999  研发公摊", 2115.77])
            sheet.append(["FYXM02_SYS  运费 小计", None, None, None, 2115.77])
            source_book.save(source_path)

            located = locate_combination_summary_sheet(source_path)
            self.assertEqual(located, (source_path, "Sheet", 1))
            workbook = load_workbook(source_path, data_only=True)
            records, month = parse_combination_summary_sheet(
                source_path, workbook["Sheet"], located[2], []
            )

        self.assertEqual(month, 0)
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0].fee_code, "FYXM02-SYS")
        self.assertEqual(records[0].project_code, "RD9999")
        self.assertEqual(records[0].amount, Decimal("2115.77"))

    def test_classification_deduplicates_equivalent_combination_sources(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            hours_path = temp_path / "工时.xlsx"
            hours_book = Workbook()
            hours = hours_book.active
            hours.title = "工时"
            hours.append(["姓名", "1月", "RD141232", "RD140928"])
            hours.append(["测试员工", 100, 60, 40])
            hours_book.save(hours_path)

            salary_path = temp_path / "工资.xlsx"
            salary_book = Workbook()
            salary = salary_book.active
            salary.append(["序号", "人员名称", "1月", "1月工资费用"])
            salary.append(["001", "测试员工", 100, 1000])
            salary_book.save(salary_path)

            combo_paths = [temp_path / "核算维度与科目组合表1月.xlsx", temp_path / "核算维度与科目组合表_副本.xlsx"]
            for combo_path in combo_paths:
                combo_book = Workbook()
                combo = combo_book.active
                combo.append(["费用项目", None, None, "研发项目", "借方发生(原币)"])
                combo.append(["FYXM02_SYS  运费", None, None, "RD9999  研发公摊", 100])
                combo_book.save(combo_path)

            classified = classify_input_workbooks([hours_path, salary_path, *combo_paths])

        self.assertEqual(classified["dimension_ledger"], combo_paths[0])
        self.assertEqual(classified["expense_ledgers"], [])

    def test_invalid_hours_are_reported_as_a_rule_error_not_an_index_error(self) -> None:
        """A rejected hours workbook must never fall through to roles["hours"][0]."""
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "工时.xlsx"
            source_book = Workbook()
            sheet = source_book.active
            sheet.append(["姓名", "1月", "RD141232", "RD140928"])
            sheet.append(["测试员工", 100, 60, 60])
            source_book.save(source_path)

            with self.assertRaisesRegex(
                RuntimeError,
                "工时数据勾稽未通过",
            ) as raised:
                classify_input_workbooks([source_path])

        self.assertNotIn("IndexError", str(raised.exception))

    def test_legacy_hours_prefers_right_side_hours_total_over_month_capacity(self) -> None:
        """有“1月”标准工时和右侧“工时合计”时，必须使用后者作项目汇总。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "工时.xlsx"
            source_book = Workbook()
            sheet = source_book.active
            sheet.append(["序号", "姓名", "1月", "RD141232", "RD140928", "工时合计"])
            sheet.append([None, None, None, "低压驱动器", "一体化电控", None])
            # 项目工时超过标准月工时，但与显式的工时合计相等，属于有效加班数据。
            sheet.append([1, "测试员工", 176, 120, 84, 204])
            source_book.save(source_path)

            _, result, _ = calculate_legacy_hours(source_path, OrderedDict())

        self.assertEqual(result["people"][0]["final_hours"], Decimal("204"))
        self.assertEqual(result["projects"][0]["project_name"], "低压驱动器")

    def test_salary_without_hours_is_shared_by_global_hours_rate(self) -> None:
        """工资表有、工时表无的人员工资应按全员工时分配率进入各项目。"""
        hours_result = {
            "people": [],
            "projects": [
                {"project_code": "RD141232", "final_hours": Decimal("60")},
                {"project_code": "RD140928", "final_hours": Decimal("40")},
            ],
        }
        salary_rows = [SalaryRow(2, "034", "沈佳诚", Decimal("1272.66"))]
        warnings: list[str] = []

        _, totals, salary_total, matched_salary, unmatched_salary = allocate_salary(
            salary_rows, hours_result, ["RD141232", "RD140928"], warnings
        )

        self.assertEqual(totals, {"RD141232": Decimal("763.60"), "RD140928": Decimal("509.06")})
        self.assertEqual(salary_total, matched_salary)
        self.assertEqual(unmatched_salary, Decimal("0"))
        self.assertIn("沈佳诚", warnings[0])

    def test_categorized_dimension_detail_ledger_uses_preface_fee(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "核算维度明细账_运费.xlsx"
            make_categorized_detail_ledger(source_path, "FYXM02_SYS", "运费", 3284.9)

            parsed, month, source_sheets = parse_ledger_file(source_path, [])

        self.assertEqual(month, 1)
        self.assertEqual([(source.fee_code, source.fee_name) for source in source_sheets], [("FYXM02-SYS", "运费")])
        self.assertEqual(len(parsed), 1)
        self.assertEqual(parsed[0].fee_code, "FYXM02-SYS")
        self.assertEqual(parsed[0].project_code, "RD9999")
        self.assertEqual(parsed[0].amount, Decimal("3284.9"))

    def test_categorized_detail_ledgers_can_supply_dimension_source(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            hours_path = temp_path / "工时.xlsx"
            hours_book = Workbook()
            hours = hours_book.active
            hours.title = "工时"
            hours.append(["姓名", "1月", "RD141232", "RD140928"])
            hours.append(["测试员工", 100, 60, 40])
            hours_book.save(hours_path)

            salary_path = temp_path / "工资.xlsx"
            salary_book = Workbook()
            salary = salary_book.active
            salary.append(["序号", "人员名称", "1月", "1月工资费用"])
            salary.append(["001", "测试员工", 100, 1000])
            salary_book.save(salary_path)

            freight_path = temp_path / "核算维度明细账_运费.xlsx"
            depreciation_path = temp_path / "核算维度明细账_折旧.xlsx"
            make_categorized_detail_ledger(freight_path, "FYXM02_SYS", "运费", 100)
            make_categorized_detail_ledger(depreciation_path, "FYXM08_SYS", "折旧费用", 200)

            classified = classify_input_workbooks([hours_path, salary_path, freight_path, depreciation_path])
            warnings: list[str] = []
            records, month, source = read_dimension_summary_workbook(classified["dimension_ledger"], warnings)

        self.assertEqual(classified["dimension_ledger"], [freight_path, depreciation_path])
        self.assertEqual(month, 1)
        self.assertEqual(source[2], "generated")
        self.assertEqual(
            [(record.fee_code, record.project_code, record.amount) for record in records],
            [("FYXM02-SYS", "RD9999", Decimal("100.00")), ("FYXM08-SYS", "RD9999", Decimal("200.00"))],
        )
        self.assertTrue(any("费用明细账" in warning and "测算金额来源" in warning for warning in warnings))

    def test_normalizing_an_already_normalized_combination_is_idempotent(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "normalized_combination.xlsx"
            output_path = Path(temp_dir) / "normalized_combination_copy.xlsx"
            source_book = Workbook()
            sheet = source_book.active
            sheet.append(["费用项目", None, None, "研发项目", "费用金额（元）"])
            sheet.append(["FYXM02_SYS  运费", None, None, "RD9999  研发公摊", 100])
            source_book.save(source_path)

            result = normalize_combination_workbook(source_path, output_path)
            normalized = load_workbook(output_path, data_only=False)

        self.assertEqual(result["primary_amount_column"], "E")
        self.assertEqual(result["cleared_amount_columns"], [])
        self.assertEqual(normalized.active["E1"].value, "费用金额（元）")
        self.assertEqual(normalized.active["E2"].value, 100)

    def test_combination_source_is_copied_row_for_row_into_month_summary(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "任意客户_核算收入_科目组合_导出.xlsx"
            source_book = Workbook()
            source = source_book.active
            source.title = "导出数据"
            source.append(["报表名称", None, None, None, None])
            source.append(["期间", "2026-01", None, None, None])
            source.append(["费用项目", None, None, "研发项目", "借方发生(原币)"])
            source.append(["FYXM02_SYS  运费", None, None, "RD9999  研发公摊", 100])
            source.append(["FYXM02_SYS  运费 小计", None, None, None, 100])
            source.append(["FYXM28_SYS  直接投入(材料、燃料)", None, None, "RD1260301  生产自研工装", 50])
            source_book.save(source_path)

            workbook = load_workbook(OUTPUT_TEMPLATE, data_only=False)
            copy_summary_to_template(
                workbook,
                (source_path, "导出数据", "combination"),
                [],
                month=1,
            )
            summary = workbook["项目分类明细汇总表1月"]

        self.assertEqual(
            [summary.cell(1, col).value for col in range(1, 5)],
            ["序号", "费用项目", "研发项目", "费用金额（元）"],
        )
        self.assertEqual(
            [summary.cell(2, col).value for col in range(1, 5)],
            [1, "FYXM02_SYS  运费", "RD9999  研发公摊", 100],
        )
        self.assertEqual(
            [summary.cell(3, col).value for col in range(1, 5)],
            [None, "FYXM02_SYS  运费 小计", None, 100],
        )
        self.assertEqual(
            [summary.cell(4, col).value for col in range(1, 5)],
            [2, "FYXM28_SYS  直接投入(材料、燃料)", "RD1260301  生产自研工装", 50],
        )

    def test_oversized_combination_copy_filters_non_template_fee_rows(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "核算维度与科目组合表.xlsx"
            source_book = Workbook()
            source = source_book.active
            source.append(["费用项目", None, None, "研发项目", "借方发生(原币)"])
            source.append(["FYXM02_SYS  运费", None, None, "RD9999  研发公摊", 100])
            source.append(["FYXM02_SYS  运费 小计", None, None, None, 100])
            for index in range(24):
                source.append([f"FYXM21_SYS  企业基本养老保险{index}", None, None, "RD9999  研发公摊", 10])
                source.append([f"FYXM21_SYS  企业基本养老保险{index} 小计", None, None, None, 10])
            source.append(["FYXM08_SYS  折旧费用", None, None, "RD9999  研发公摊", 200])
            source.append(["FYXM08_SYS  折旧费用 小计", None, None, None, 200])
            source_book.save(source_path)

            workbook = load_workbook(OUTPUT_TEMPLATE, data_only=False)
            warnings: list[str] = []
            copy_summary_to_template(
                workbook,
                (source_path, "Sheet", "combination"),
                [],
                month=1,
                warnings=warnings,
            )
            summary = workbook["项目分类明细汇总表1月"]
            copied_values = [
                summary.cell(row, 2).value
                for row in range(2, summary.max_row + 1)
                if summary.cell(row, 2).value
            ]

        self.assertIn("FYXM02_SYS  运费", copied_values)
        self.assertIn("FYXM08_SYS  折旧费用", copied_values)
        self.assertTrue(all("企业基本养老保险" not in str(value) for value in copied_values))
        self.assertTrue(any("未进入固定费用页" in warning for warning in warnings))

    def test_material_allocation_overwrites_old_helper_addition(self) -> None:
        workbook = load_workbook(OUTPUT_TEMPLATE, data_only=False)
        projects = OrderedDict(
            [
                ("RD141232", Project("RD141232", "80V低压驱动器研发")),
                ("RD1260301", Project("RD1260301", "生产自研工装")),
            ]
        )
        sheet = workbook["14.材料明细账"]
        sheet["D52"] = "=G7+G12+G15"

        write_template_material_allocation(
            workbook,
            projects,
            direct_totals={"RD141232": Decimal("10"), "RD1260301": Decimal("30")},
            public_totals={"RD141232": Decimal("2"), "RD1260301": Decimal("3")},
            combined_totals={"RD141232": Decimal("12"), "RD1260301": Decimal("33")},
            month=1,
        )

        self.assertIsNone(sheet["D52"].value)
        self.assertEqual([sheet["C50"].value, sheet["G50"].value, sheet["K50"].value], [10.0, 2.0, 12.0])
        self.assertEqual([sheet["C51"].value, sheet["G51"].value, sheet["K51"].value], [30.0, 3.0, 33.0])
        self.assertEqual([sheet["C68"].value, sheet["G68"].value, sheet["K68"].value], [40.0, 5.0, 45.0])

    def test_direct_labor_total_uses_amount_number_format(self) -> None:
        workbook = load_workbook(OUTPUT_TEMPLATE, data_only=False)
        number_format = workbook["2.直接人工"]["C27"].number_format

        self.assertEqual(number_format, "#,##0.00")

    def test_embedded_freight_allocation_table_is_not_parsed_as_ledger(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "mixed_freight.xlsx"
            source_book = Workbook()
            sheet = source_book.active
            sheet.title = "运费"
            headers = ["日期", "会计年度", "期间", "凭证字", "凭证号", "摘要", "借方", "研发项目", "研发名称"]
            for col, value in enumerate(headers, start=1):
                sheet.cell(1, col).value = value
            values = ["2026-01-01", 2026, 1, "记", 1, "运输费", 100, "RD141232", "80V低压驱动器研发"]
            for col, value in enumerate(values, start=1):
                sheet.cell(2, col).value = value
            for col, value in enumerate(["项目编号", "工时合计", "工时分配率", "费用"], start=6):
                sheet.cell(3, col).value = value
            for col, value in enumerate(["RD141232", 100, 1, 100], start=6):
                sheet.cell(4, col).value = value
            source_book.save(source_path)

            parsed, _, _ = parse_ledger_file(source_path, [])

        self.assertEqual(len(parsed), 1)
        self.assertEqual(parsed[0].summary, "运输费")
        self.assertEqual(parsed[0].amount, Decimal("100"))

    def test_hours_tables_include_rd_and_shared_rd_columns(self) -> None:
        workbook = load_workbook(OUTPUT_TEMPLATE, data_only=False)
        projects = OrderedDict(
            [("RD141232", Project("RD141232", "80V低压驱动器研发", Decimal("100"), Decimal("1")))]
        )
        hours_result = {
            "projects": [
                {
                    "project_code": "RD141232",
                    "project_name": "80V低压驱动器研发",
                    "final_hours": Decimal("100"),
                }
            ],
            "people": [
                {
                    "row": 2,
                    "employee_id": "001",
                    "employee_name": "测试员工",
                    "final_hours": Decimal("100"),
                    "actual_hours": Decimal("60"),
                    "allocated_hours": Decimal("40"),
                    "final_by_project": {"RD141232": Decimal("100")},
                }
            ],
            "totals": {
                "final_hours": Decimal("100"),
                "actual_hours": Decimal("60"),
                "allocated_hours": Decimal("40"),
            },
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "hours.xlsx"
            source_book = Workbook()
            source_sheet = source_book.active
            source_sheet.title = "Hours"
            source_sheet.cell(1, 1).value = "序号"
            source_sheet.cell(1, 2).value = "姓名"
            source_sheet.cell(2, 1).value = "001"
            source_sheet.cell(2, 2).value = "测试员工"
            source_book.save(source_path)

            write_template_hours_sheets(
                workbook,
                source_path,
                SimpleNamespace(sheet="Hours"),
                hours_result,
                projects,
                month=1,
            )

        for sheet_name in ("研发人员工时统计", "研发人员工时统计合计"):
            sheet = workbook[sheet_name]
            self.assertEqual([sheet.cell(1, col).value for col in (30, 31, 32)], ["工时汇总", "RD工时", "公摊RD工时"])
        raw_sheet = workbook["研发人员工时统计"]
        self.assertEqual([raw_sheet.cell(2, col).value for col in (30, 31, 32)], [100.0, 60.0, 40.0])
        summary_sheet = workbook["研发人员工时统计合计"]
        self.assertEqual([summary_sheet.cell(3, col).value for col in (30, 31, 32)], [100.0, 60.0, 40.0])
        self.assertEqual(
            [summary_sheet.cell(41, col).value for col in (8, 9, 10)],
            ["项目编码", "科目名称", "工时分配率"],
        )
        self.assertEqual(summary_sheet.cell(43, 10).number_format, "0.00%")

    def test_material_allocation_tables_move_together_after_content(self) -> None:
        workbook = load_workbook(OUTPUT_TEMPLATE, data_only=False)
        sheet = workbook["14.材料明细账"]
        write_template_expense_source(
            sheet,
            records("FYXM28-SYS", "直接材料", 11),
            month=1,
        )

        header_row = relocate_material_allocation_tables(
            sheet,
            blank_rows=3,
            content_last_row=16,
        )

        self.assertEqual(header_row, 20)
        self.assertEqual(
            [sheet.cell(header_row, col).value for col in (1, 2, 3, 5, 6, 7, 9, 10, 11)],
            [
                "项目编码", "科目名称", "1月直接材料合计",
                "项目编码", "科目名称", "1月公摊材料合计",
                "项目编码", "科目名称", "1月材料合计",
            ],
        )

    def test_standard_fee_table_follows_arbitrary_detail_rows(self) -> None:
        workbook = load_workbook(OUTPUT_TEMPLATE, data_only=False)
        sheet = workbook["7.办公费"]
        detail_rows = records("FYXM14-SYS", "办公费", 75)
        content_last_row = fee_source_content_last_row(detail_rows)
        expected_header_row = content_last_row + 4

        actual_header_row = relocate_standard_fee_allocation_table(
            sheet,
            blank_rows=3,
            content_last_row=content_last_row,
        )
        write_template_expense_source(
            sheet,
            detail_rows,
            month=1,
            allocation_start_row=actual_header_row,
        )

        self.assertEqual(actual_header_row, expected_header_row)
        self.assertEqual(standard_fee_allocation_table(sheet)[0], expected_header_row)
        self.assertEqual(sheet.cell(content_last_row, 1).value, "合计")
        for row in range(content_last_row + 1, expected_header_row):
            self.assertTrue(all(sheet.cell(row, col).value in (None, "") for col in range(1, 10)))

    def test_freight_table_is_visible_and_amounts_balance(self) -> None:
        workbook = load_workbook(OUTPUT_TEMPLATE, data_only=False)
        sheet = workbook["1.运费"]
        detail_rows = records("FYXM02-SYS", "运费", 75)
        content_last_row = fee_source_content_last_row(detail_rows)
        expected_header_row = content_last_row + 4
        projects = OrderedDict(
            [
                ("RD141232", Project("RD141232", "80V低压驱动器研发")),
                ("RD140928", Project("RD140928", "230kW一体化电控")),
            ]
        )
        totals = {"RD141232": Decimal("50.01"), "RD140928": Decimal("49.99")}

        actual_header_row = relocate_freight_visible_table(sheet, expected_header_row)
        write_template_expense_source(
            sheet,
            detail_rows,
            month=1,
            allocation_start_row=actual_header_row,
        )
        written = write_visible_freight_allocation_table(
            sheet,
            month=1,
            projects=projects,
            project_totals=totals,
            header_row=actual_header_row,
        )

        self.assertTrue(written)
        self.assertEqual(sheet.cell(expected_header_row, 6).value, "项目编码")
        self.assertFalse(sheet.row_dimensions[expected_header_row].hidden)
        self.assertEqual(Decimal(str(sheet.cell(expected_header_row + 20, 8).value)), Decimal("100.0"))
        self.assertEqual(
            sum(Decimal(str(sheet.cell(row, 8).value or 0)) for row in range(expected_header_row + 2, expected_header_row + 20)),
            Decimal("100.00"),
        )
        for row in range(content_last_row + 1, expected_header_row):
            self.assertTrue(all(sheet.cell(row, col).value in (None, "") for col in range(1, 10)))

    def test_relocated_freight_clears_all_historical_rows_outside_active_table(self) -> None:
        workbook = load_workbook(OUTPUT_TEMPLATE, data_only=False)
        sheet = workbook["1.运费"]
        detail_rows = records("FYXM02-SYS", "运费", 3)
        expected_header_row = fee_source_content_last_row(detail_rows) + 4

        actual_header_row = relocate_freight_visible_table(sheet, expected_header_row)
        write_template_expense_source(
            sheet,
            detail_rows,
            month=7,
            allocation_start_row=actual_header_row,
        )
        projects = OrderedDict(
            [("RD141232", Project("RD141232", "80V低压驱动器研发"))]
        )
        write_visible_freight_allocation_table(
            sheet,
            month=7,
            projects=projects,
            project_totals={"RD141232": Decimal("100")},
            header_row=actual_header_row,
        )

        self.assertEqual(actual_header_row, expected_header_row)
        self.assertEqual(sheet.cell(expected_header_row, 6).value, "项目编码")
        self.assertIsNone(sheet["A12"].value)
        self.assertIsNone(sheet["G56"].value)


if __name__ == "__main__":
    unittest.main()
