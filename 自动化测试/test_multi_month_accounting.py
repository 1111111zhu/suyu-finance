"""Real fixed-template monthly isolation (synthetic source data, no production data)."""
import hashlib
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from test_rd_accounting_output import fixtures
from rd_cost_workflow_v1.engine import build_output, build_outputs, classify_input_workbooks, RuleError, OUTPUT_TEMPLATE, input_workbook_period


def month_fixtures(folder, month, salary, first_hours):
    folder.mkdir()
    original = fixtures(folder)
    result = []
    for index, path in enumerate(original):
        book = load_workbook(path)
        for sheet in book:
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.value = cell.value.replace("9月", f"{month}月").replace("2025/9/", f"2025/{month}/")
        if index == 0:
            book.active.cell(3, 4).value = first_hours
            book.active.cell(3, 5).value = 100 - first_hours
        if index == 1:
            book.active.cell(2, 4).value = salary
        if index == 3:
            book.active.cell(2, 3).value = month
        target = folder / (path.stem.replace("9月", f"{month}月") + f"_{month}月.xlsx")
        book.save(target)
        result.append(target)
    return result


class MultiMonthTest(unittest.TestCase):
    def test_two_months_independent_salary_hours_template_and_inputs(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            paths = month_fixtures(folder / "march", 3, 1000, 60) + month_fixtures(folder / "august", 8, 2000, 25)
            hashes = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in [*paths, OUTPUT_TEMPLATE]}
            results = build_outputs(paths, folder / "results", accounting_year=2025)["results"]
            self.assertEqual([r["month"] for r in results], [3, 8])
            self.assertEqual([r["cost_total"] for r in results], ["1300.00", "2300.00"])
            self.assertEqual({p["code"]: p["total"] for p in results[1]["project_costs"]},
                             {"RD141232": "725.00", "RD1260301": "1575.00"})
            template = load_workbook(OUTPUT_TEMPLATE)
            for result in results:
                book = load_workbook(result["output_path"])
                self.assertEqual(book.sheetnames, [s.replace("汇总表1月", f"汇总表{result['month']}月") for s in template.sheetnames])
                # Existing single-month writer relocates some template blocks;
                # batching must not introduce any further layout/value changes.
                inputs = classify_input_workbooks([p for p in paths if f"_{result['month']}月" in p.stem])
                baseline_path = folder / f"baseline_{result['month']}.xlsx"
                baseline = build_output(inputs["hours"], inputs["salary"], inputs["dimension_ledger"],
                                        inputs["expense_ledgers"], baseline_path, accounting_year=2025)
                self.assertEqual(baseline["project_costs"], result["project_costs"])
                baseline_book = load_workbook(baseline_path)
                for sheet, source in zip(book, baseline_book):
                    self.assertEqual(str(sheet.merged_cells), str(source.merged_cells))
                    self.assertEqual(list(sheet.values), list(source.values))
                for sheet in book:
                    self.assertFalse(any(cell.data_type == "e" for row in sheet for cell in row))
            self.assertEqual(hashes, {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in hashes})

    def test_ledger_summary_month_is_not_accounting_period(self):
        with tempfile.TemporaryDirectory() as temp:
            paths = fixtures(Path(temp))
            book = load_workbook(paths[3])
            book.active.cell(2, 6).value = "12月月结快递费，补结2024/12/31费用"
            book.save(paths[3])
            self.assertEqual(input_workbook_period(paths[3]), (9, {2025}))
            book.active.cell(2, 1).value = "2025/12/1"
            book.save(paths[3])
            with self.assertRaisesRegex(RuleError, "不一致或多个期间"):
                input_workbook_period(paths[3])

    def test_missing_month_salary_does_not_publish_partial_results(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            paths = month_fixtures(folder / "march", 3, 1000, 60) + month_fixtures(folder / "august", 8, 2000, 25)
            paths = [p for p in paths if not ("工资" in p.name and "8月" in p.name)]
            with self.assertRaisesRegex(RuleError, "8月资料不完整"):
                build_outputs(paths, folder / "results", accounting_year=2025)
            self.assertFalse((folder / "results").exists())

    def test_conflicting_filename_and_contents_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            paths = month_fixtures(folder / "march", 3, 1000, 60)
            wrong = paths[0].with_name("工时8月.xlsx")
            paths[0].rename(wrong)
            paths[0] = wrong
            with self.assertRaisesRegex(RuleError, "不一致或多个期间"):
                build_outputs(paths, folder / "results", accounting_year=2025)

    def test_multi_month_sheet_and_cross_year_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            paths = month_fixtures(folder / "march", 3, 1000, 60)
            book = load_workbook(paths[0])
            book.create_sheet("8月工时")
            book.save(paths[0])
            with self.assertRaisesRegex(RuleError, "多个期间"):
                build_outputs(paths, folder / "results", accounting_year=2025)
            book.remove(book["8月工时"])
            book.save(paths[0])
            with self.assertRaisesRegex(RuleError, "年度不一致"):
                build_outputs(paths, folder / "results", accounting_year=2026)


if __name__ == "__main__":
    unittest.main()
