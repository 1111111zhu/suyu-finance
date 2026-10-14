"""End-to-end accounting contract against the real fixed-template engine."""
import sys, json, hashlib, unittest, tempfile
from pathlib import Path
from decimal import Decimal
from openpyxl import Workbook, load_workbook
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/"public"/"py"))
from rd_cost_workflow_v1.engine import build_output, classify_input_workbooks, OUTPUT_TEMPLATE
def save(path,rows,title="Sheet1"):
    wb=Workbook();ws=wb.active;ws.title=title
    for row in rows:ws.append(row)
    wb.save(path)
def fixtures(folder):
    paths=[folder/"工时.xlsx",folder/"工资.xlsx",folder/"核算维度与科目组合表9月.xlsx",folder/"费用明细账.xlsx"]
    save(paths[0],[["序号","姓名","9月","RD141232","RD1260301","工时合计"],[None,None,None,"80V低压驱动器研发","生产自研工装",None],[1,"虚拟测试员工",100,60,40,100]],"工时")
    save(paths[1],[["序号","人员名称","9月","9月工资费用"],["001","虚拟测试员工",100,1000]])
    save(paths[2],[["费用项目",None,None,"研发项目","借方发生(原币)"],["FYXM02_SYS  运费",None,None,"RD9999  研发公摊",100],["FYXM28_SYS  直接材料",None,None,"RD141232  80V低压驱动器研发",200]])
    save(paths[3],[["日期","会计年度","期间","凭证字","凭证号","摘要","借方","研发项目","研发名称"],["2025/9/1",2025,9,"记",1,"运费测试",100,"RD9999","研发公摊"]],"FYXM02_SYS_运费")
    return paths
class AccountingTest(unittest.TestCase):
    def test_intact_source_columns_and_duplicate_amount_headers(self):
        with tempfile.TemporaryDirectory() as temp:
            folder=Path(temp);paths=fixtures(folder)
            source=load_workbook(paths[2]);sheet=source.active
            sheet.insert_cols(1)
            sheet.cell(1,1).value="备注"
            sheet.cell(1,8).value="借方发生(原币)"
            sheet.cell(2,1).value="保留这条备注"
            sheet.cell(2,8).value=999999
            sheet.cell(3,8).value=888888
            source.save(paths[2])
            hashes={p:hashlib.sha256(p.read_bytes()).hexdigest() for p in paths}
            inputs=classify_input_workbooks(paths)
            output=folder/"result.xlsx"
            result=build_output(inputs["hours"],inputs["salary"],inputs["dimension_ledger"],inputs["expense_ledgers"],output)
            self.assertEqual(result["cost_total"],"1300.00")
            formulas=load_workbook(output,data_only=False)
            copied=formulas["项目分类明细汇总表9月"]
            self.assertEqual(list(copied.values),list(sheet.values))
            self.assertEqual(formulas.sheetnames,[name.replace("汇总表1月","汇总表9月") for name in load_workbook(OUTPUT_TEMPLATE).sheetnames])
            links=[c.value for ws in formulas for c in ws._cells.values() if c.data_type=="f" and "项目分类明细汇总表9月" in c.value]
            self.assertTrue(any("!F2" in value for value in links))
            self.assertTrue(any("!F3" in value for value in links))
            self.assertFalse(any("!H2" in value or "!H3" in value for value in links))
            self.assertEqual(hashes,{p:hashlib.sha256(p.read_bytes()).hexdigest() for p in paths})

    def test_frontier_rc_projects_are_not_dropped(self):
        with tempfile.TemporaryDirectory() as temp:
            folder=Path(temp);paths=fixtures(folder)
            for path in paths:
                book=load_workbook(path)
                for sheet in book:
                    for row in sheet:
                        for cell in row:
                            if isinstance(cell.value,str):
                                cell.value=cell.value.replace("RD1260301","RC22505-01")
                book.save(path)
            classified=classify_input_workbooks(paths)
            result=build_output(classified["hours"],classified["salary"],classified["dimension_ledger"],classified["expense_ledgers"],folder/"result.xlsx")
            totals={p["code"]:p["total"] for p in result["project_costs"]}
            self.assertEqual(totals["RC22505-01"],"440.00")
            self.assertEqual(result["cost_total"],"1300.00")
    def test_workbook_and_database_payload_match(self):
        with tempfile.TemporaryDirectory() as temp:
            folder=Path(temp);paths=fixtures(folder)
            hashes={p:hashlib.sha256(p.read_bytes()).hexdigest() for p in [*paths,OUTPUT_TEMPLATE]}
            classified=classify_input_workbooks(paths)
            output=folder/"result.xlsx"
            result=build_output(classified["hours"],classified["salary"],classified["dimension_ledger"],classified["expense_ledgers"],output)
            totals={p["code"]:Decimal(p["total"]) for p in result["project_costs"]}
            self.assertEqual(totals,{"RD141232":Decimal("860.00"),"RD1260301":Decimal("440.00")})
            self.assertEqual(result["cost_total"],"1300.00")
            self.assertEqual(result["year"],2025)
            self.assertEqual(result["month"],9)
            self.assertEqual(hashes,{p:hashlib.sha256(p.read_bytes()).hexdigest() for p in hashes})
            formula_book=load_workbook(output,data_only=False)
            self.assertTrue(formula_book["研发项目主表"]["C3"].value.startswith("="))
            self.assertEqual(formula_book["研发项目主表"]["Y3"].value,"=SUM(C3:X3)")
            self.assertFalse(any("LET(" in c.value or "_xlfn" in c.value for sheet in formula_book for c in sheet._cells.values() if c.data_type == "f"))
            labor_formula=formula_book["2.直接人工"]["C12"].value
            self.assertEqual("=ROUND($G$5*'研发人员工时统计合计'!$J$43,2)",labor_formula)
            self.assertNotIn("SUMPRODUCT",labor_formula)
            wb=load_workbook(output,data_only=True);source=load_workbook(OUTPUT_TEMPLATE,data_only=False)
            self.assertEqual(wb.sheetnames,[name.replace("汇总表1月","汇总表9月") for name in source.sheetnames])
            main=wb["研发项目主表"]
            for record in result["project_costs"]:
                row=next(r for r in main.iter_rows() if r[0].value==record["code"])
                self.assertEqual(sum((Decimal(str(c.value or 0)) for c in row[2:24]),Decimal(0)),Decimal(record["total"]))
            # A JSON round-trip must preserve exact decimals for the web/API boundary.
            payload=json.loads(json.dumps(result,default=str))
            self.assertEqual(payload["project_costs"][0]["costs"][0]["amount"],"200.00")
    def test_unequal_salaries_use_total_pool_and_displayed_rates(self):
        with tempfile.TemporaryDirectory() as temp:
            folder=Path(temp);paths=fixtures(folder)
            # Different salary levels must not outweigh the visible hours rate.
            save(paths[0],[["序号","姓名","9月","RD141232","RD1260301","工时合计"],
                [None,None,None,"80V低压驱动器研发","生产自研工装",None],
                [1,"甲",203,203,0,203],[2,"乙",9797,0,9797,9797]],"工时")
            save(paths[1],[["序号","人员名称","9月","9月工资费用"],
                ["001","甲",203,10000],["002","乙",9797,636720.65]])
            classified=classify_input_workbooks(paths);output=folder/"result.xlsx"
            result=build_output(classified["hours"],classified["salary"],classified["dimension_ledger"],classified["expense_ledgers"],output)
            wb=load_workbook(output,data_only=True)
            self.assertEqual(wb["研发人员工时统计合计"]["J43"].value,0.0203)
            self.assertEqual(wb["2.直接人工"]["C12"].value,13128.43)
            self.assertEqual(wb["研发项目主表"]["D3"].value,13128.43)
            self.assertEqual(wb["2.直接人工"]["C30"].value,646720.65)
            costs={p["code"]:{c["category"]:c["amount"] for c in p["costs"]} for p in result["project_costs"]}
            self.assertIn("13128.43",costs["RD141232"].values())
    def test_project_rounding_is_summed_without_residual_adjustment(self):
        with tempfile.TemporaryDirectory() as temp:
            folder=Path(temp);paths=fixtures(folder)
            save(paths[0],[["序号","姓名","9月","RD141232","RD1260301","RD140928","工时合计"],
                [None,None,None,"80V低压驱动器研发","生产自研工装","230kW一体化电控",None],
                [1,"虚拟测试员工",3,1,1,1,3]],"工时")
            save(paths[1],[["序号","人员名称","9月","9月工资费用"],["001","虚拟测试员工",3,0.01]])
            classified=classify_input_workbooks(paths);output=folder/"result.xlsx"
            result=build_output(classified["hours"],classified["salary"],classified["dimension_ledger"],classified["expense_ledgers"],output)
            wb=load_workbook(output,data_only=True);formulas=load_workbook(output,data_only=False)
            self.assertEqual(wb["2.直接人工"]["G5"].value,0.01)
            self.assertEqual([wb["2.直接人工"].cell(r,3).value for r in range(12,15)],[0,0,0])
            self.assertEqual(wb["2.直接人工"]["C30"].value,0)
            self.assertEqual(Decimal(result["labor_total"]),0)
            for r in range(12,15):
                self.assertEqual(formulas["2.直接人工"].cell(r,3).value,
                    f"=ROUND($G$5*'研发人员工时统计合计'!$J${r+31},2)")
if __name__=="__main__":unittest.main()
