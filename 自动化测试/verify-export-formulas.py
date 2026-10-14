"""Independently evaluate exported formulas from their actual cached inputs."""
import sys, re, math, logging
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import range_boundaries,column_index_from_string
import formulas
logging.disable(logging.CRITICAL)
def expand_let(expression):
    # The independent interpreter cannot resolve dependent LET bindings.
    # Expand lexical bindings, then let it evaluate the resulting Excel math.
    while (match:=re.search(r"\bLET\(",expression,re.I)):
        start=match.end();depth=1;brace=0;arguments=[];last=start;end=start
        for end in range(start,len(expression)):
            char=expression[end]
            if char=='(':depth+=1
            elif char==')':
                depth-=1
                if depth==0:arguments.append(expression[last:end]);break
            elif char=='{':brace+=1
            elif char=='}':brace-=1
            elif char==',' and depth==1 and brace==0:arguments.append(expression[last:end]);last=end+1
        bindings={}
        def substitute(value):
            if not bindings:return value
            pattern=r'\b('+ '|'.join(re.escape(k) for k in sorted(bindings,key=len,reverse=True))+r')\b'
            return re.sub(pattern,lambda m:'('+bindings[m[1].lower()]+')',value,flags=re.I)
        for i in range(0,len(arguments)-1,2):bindings[arguments[i].lower()]=substitute(arguments[i+1])
        expression=expression[:match.start()]+'('+substitute(arguments[-1])+')'+expression[end+1:]
    return expression

path=Path(sys.argv[1]);formula_book=load_workbook(path,data_only=False);values=load_workbook(path,data_only=True)
errors=[];checked=0
for ws in formula_book:
 for row in ws:
  for cell in row:
   expected=values[ws.title][cell.coordinate].value
   if cell.data_type!='f' or not isinstance(expected,(int,float)):continue
   parts=[]
   try:
    for token in Tokenizer(cell.value).items:
     text=token.value
     if token.type=='OPERAND' and token.subtype=='RANGE':
      sheet=ws.title;address=text
      if '!' in text:
       sheet,address=text.rsplit('!',1);sheet=sheet.strip("'").replace("''", "'")
      match=re.fullmatch(r'\$?([A-Z]+)\$?\d+(?::\$?[A-Z]+\$?\d+)?',address,re.I)
      if match and len(match[1])<=3 and column_index_from_string(match[1])<=16384:
       x1,y1,x2,y2=range_boundaries(address.replace('$',''));rows=[]
       for cells in values[sheet].iter_rows(min_row=y1,max_row=y2,min_col=x1,max_col=x2):
        row_values=[]
        for item in cells:
         value=item.value
         if value is None:value=0
         if not isinstance(value,(int,float)):raise ValueError(f'non-numeric input {sheet}!{item.coordinate}')
         row_values.append(str(value))
        rows.append(','.join(row_values))
       text=rows[0] if x1==x2 and y1==y2 else '{'+';'.join(rows)+'}'
     parts.append(text)
    expression=expand_let('='+''.join(parts))
    actual=formulas.Parser().ast(expression)[1].compile()()
    number=float(actual.reshape(-1)[0]) if hasattr(actual,'reshape') else float(actual)
    if not math.isfinite(number) or abs(number-expected)>1e-7:errors.append((ws.title,cell.coordinate,expected,number))
    checked+=1
   except Exception as error:errors.append((ws.title,cell.coordinate,type(error).__name__,str(error)[:100]))
print('independent formulas checked:',checked,'mismatches:',len(errors));print(errors[:15]);assert not errors
