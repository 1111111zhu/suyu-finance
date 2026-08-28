#!/usr/bin/env python3
"""研发人员RD工时识别与按人员项目比例分摊工具。

本程序不依赖固定行列号。它先定位“人员 × 项目”矩阵，再执行规则文件中
定义的工时计算，并输出可追溯的 JSON 报告；在输入工作簿副本右侧新增
项目公摊矩阵，并新增“RD工时加公摊”和识别日志工作表。
"""

from __future__ import annotations

import argparse
import json
import math
import os
import posixpath
import re
import shutil
import sys
import tempfile
import unicodedata
import xml.etree.ElementTree as ET
import zipfile
from copy import copy
from dataclasses import asdict, dataclass, field
from decimal import Decimal, InvalidOperation, getcontext
from pathlib import Path
from typing import Any, Iterable, Optional

# 通用离线包把共享依赖放在固定壳根目录的 vendor 中；开发环境仍可使用
# 规则包旁的 vendor，二者都不存在时再回退到系统 Python 环境。
PACK_ROOT = Path(__file__).resolve().parent
for vendor_dir in (PACK_ROOT / "vendor", PACK_ROOT.parent.parent / "vendor"):
    if vendor_dir.is_dir():
        sys.path.insert(0, str(vendor_dir))
        break

try:
    import yaml
except ImportError as exc:  # pragma: no cover - dependency guard
    raise SystemExit("缺少 PyYAML。请先执行：python3 -m pip install PyYAML openpyxl") from exc

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import column_index_from_string, get_column_letter
    from openpyxl.workbook.properties import CalcProperties
except ImportError as exc:  # pragma: no cover - dependency guard
    raise SystemExit("缺少 openpyxl。请先执行：python3 -m pip install PyYAML openpyxl") from exc


getcontext().prec = 28


class RuleError(RuntimeError):
    """可定位、不可静默处理的业务或版式异常。"""


@dataclass
class ProjectColumn:
    column: int
    column_letter: str
    project_code: str
    project_name: str
    header_cell: str
    master_matched: bool


@dataclass
class EmployeeRow:
    row: int
    employee_id: str
    employee_name: str
    source_cell: str


@dataclass
class Layout:
    sheet: str
    project_header_row: int
    project_name_row: Optional[int]
    employee_id_column: Optional[int]
    employee_name_column: int
    month_reference_column: int
    source_rd_column: int
    source_public_rd_column: int
    source_hours_total_column: int
    employee_start_row: int
    employee_end_row: int
    total_row: Optional[int]
    project_columns: list[ProjectColumn]
    employees: list[EmployeeRow]
    summary_rows: dict[str, int]
    confidence: float
    evidence: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value)).strip()
    text = re.sub(r"[‐‑‒–—―−]", "-", text)
    return re.sub(r"\s+", " ", text)


def normalize_key(value: Any) -> str:
    return re.sub(r"[\s_()（）:：]+", "", normalize_text(value)).lower()


def canonical_project_code(value: Any) -> str:
    return normalize_text(value).replace(" ", "").replace("_", "-").upper()


def aliases(rules: dict[str, Any], key: str) -> list[str]:
    return [normalize_key(v) for v in rules["identification"]["aliases"].get(key, [])]


def is_alias(value: Any, candidates: Iterable[str]) -> bool:
    key = normalize_key(value)
    return bool(key) and any(key == item or item in key for item in candidates)


def parse_column_override(value: Optional[str]) -> Optional[int]:
    if not value:
        return None
    value = value.strip()
    if value.isdigit():
        return int(value)
    return column_index_from_string(value.upper())


def to_decimal(value: Any, location: str, *, blank_zero: bool = False) -> Decimal:
    if value is None or (isinstance(value, str) and not value.strip()):
        if blank_zero:
            return Decimal("0")
        raise RuleError(f"{location} 缺少数值")
    if isinstance(value, bool):
        raise RuleError(f"{location} 不能使用布尔值作为工时")
    if isinstance(value, str) and value.startswith("#"):
        raise RuleError(f"{location} 含 Excel 错误值：{value}")
    try:
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            raise InvalidOperation
        return Decimal(str(value).replace(",", ""))
    except (InvalidOperation, ValueError) as exc:
        raise RuleError(f"{location} 不是有效工时数值：{value!r}") from exc


def cell_effective_value(ws_formula: Any, ws_values: Any, row: int, col: int) -> Any:
    formula_value = ws_formula.cell(row, col).value
    if isinstance(formula_value, str) and formula_value.startswith("="):
        cached = ws_values.cell(row, col).value
        if cached is None:
            raise RuleError(
                f"{ws_formula.title}!{get_column_letter(col)}{row} 是公式但没有缓存值，"
                "请先用 Excel/WPS 打开并保存，使公式完成重算"
            )
        return cached
    return formula_value


def load_rules(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as handle:
        data = yaml.safe_load(handle)
    if not isinstance(data, dict) or "identification" not in data or "calculation" not in data:
        raise RuleError("规则文件缺少 identification 或 calculation 顶层配置")
    return data


def project_pattern(rules: dict[str, Any]) -> re.Pattern[str]:
    return re.compile(rules["identification"]["project_code_pattern"], re.IGNORECASE)


def scan_header_pair(
    workbook: Any,
    left_aliases: list[str],
    right_aliases: list[str],
    max_rows: int,
    max_cols: int,
) -> list[tuple[Any, int, int, int]]:
    found: list[tuple[Any, int, int, int]] = []
    for ws in workbook.worksheets:
        for row in range(1, min(ws.max_row, max_rows) + 1):
            left_col = right_col = None
            for col in range(1, min(ws.max_column, max_cols) + 1):
                value = ws.cell(row, col).value
                if is_alias(value, left_aliases):
                    left_col = col
                if is_alias(value, right_aliases):
                    right_col = col
            if left_col and right_col and left_col != right_col:
                found.append((ws, row, left_col, right_col))
    return found


def discover_project_master(workbook: Any, rules: dict[str, Any]) -> tuple[dict[str, str], list[str]]:
    pattern = project_pattern(rules)
    candidates = scan_header_pair(
        workbook,
        aliases(rules, "project_code_header"),
        aliases(rules, "project_name_header"),
        rules["identification"]["scan_limits"]["master_header_rows"],
        rules["identification"]["scan_limits"]["master_header_columns"],
    )
    mapping: dict[str, str] = {}
    sources: list[str] = []
    for ws, header_row, code_col, name_col in candidates:
        blank_streak = 0
        local_count = 0
        for row in range(header_row + 1, min(ws.max_row, header_row + 500) + 1):
            raw_code = ws.cell(row, code_col).value
            raw_name = ws.cell(row, name_col).value
            if raw_code in (None, "") and raw_name in (None, ""):
                blank_streak += 1
                if blank_streak >= 3:
                    break
                continue
            blank_streak = 0
            code = canonical_project_code(raw_code)
            if not pattern.fullmatch(code):
                continue
            name = normalize_text(raw_name)
            if not name:
                continue
            if code in mapping and normalize_key(mapping[code]) != normalize_key(name):
                raise RuleError(f"项目主数据中编码 {code} 对应多个不同名称")
            mapping[code] = name
            local_count += 1
        if local_count:
            sources.append(f"{ws.title}!{get_column_letter(code_col)}{header_row}:{get_column_letter(name_col)}{header_row}")
    return mapping, sources


def row_numeric_density(ws_formula: Any, ws_values: Any, row_start: int, columns: list[int]) -> float:
    if not columns:
        return 0.0
    observed = numeric = 0
    for row in range(row_start, min(ws_formula.max_row, row_start + 12) + 1):
        for col in columns:
            try:
                value = cell_effective_value(ws_formula, ws_values, row, col)
            except RuleError:
                value = None
            if value in (None, ""):
                continue
            observed += 1
            try:
                Decimal(str(value).replace(",", ""))
                numeric += 1
            except (InvalidOperation, ValueError):
                pass
    return numeric / max(observed, 1)


def select_hours_sheet(
    wb_formula: Any,
    wb_values: Any,
    rules: dict[str, Any],
    project_master: dict[str, str],
    sheet_override: Optional[str],
    header_row_override: Optional[int],
) -> tuple[Any, int, list[tuple[int, str]]]:
    pattern = project_pattern(rules)
    min_projects = int(rules["identification"]["minimum_project_columns"])
    scan_rows = int(rules["identification"]["scan_limits"]["hours_header_rows"])
    scan_cols = int(rules["identification"]["scan_limits"]["hours_header_columns"])
    preferred_names = [normalize_key(x) for x in rules["identification"]["preferred_hours_sheet_names"]]

    if sheet_override:
        if sheet_override not in wb_formula.sheetnames:
            raise RuleError(f"找不到指定工作表：{sheet_override}")
        sheet_names = [sheet_override]
    else:
        sheet_names = wb_formula.sheetnames

    candidates: list[tuple[float, Any, int, list[tuple[int, str]]]] = []
    for sheet_name in sheet_names:
        ws = wb_formula[sheet_name]
        ws_values = wb_values[sheet_name]
        rows = [header_row_override] if header_row_override else range(1, min(ws.max_row, scan_rows) + 1)
        for row in rows:
            if row is None or row < 1 or row > ws.max_row:
                continue
            all_matches: list[tuple[int, str]] = []
            for col in range(1, min(ws.max_column, scan_cols) + 1):
                code = canonical_project_code(ws.cell(row, col).value)
                if pattern.fullmatch(code):
                    all_matches.append((col, code))
            # 结果表会在源表右侧复写一组项目编码。只把连续项目列视为同一矩阵，
            # 避免把源RD工时和已生成的公摊RD工时表拼成重复项目。
            match_groups: list[list[tuple[int, str]]] = []
            for match in all_matches:
                if not match_groups or match[0] != match_groups[-1][-1][0] + 1:
                    match_groups.append([match])
                else:
                    match_groups[-1].append(match)
            for matches in match_groups:
                if len(matches) < min_projects:
                    continue
                density = row_numeric_density(ws, ws_values, row + 1, [col for col, _ in matches])
                master_ratio = (
                    sum(1 for _, code in matches if code in project_master) / len(matches)
                    if project_master
                    else 0.0
                )
                sheet_bonus = 20 if any(key in normalize_key(sheet_name) for key in preferred_names) else 0
                score = len(matches) * 8 + density * 25 + master_ratio * 30 + sheet_bonus
                candidates.append((score, ws, row, matches))

    if not candidates:
        raise RuleError("未找到横向排列的项目编码表头；请检查项目编码格式或使用 --project-header-row 指定")
    candidates.sort(key=lambda item: item[0], reverse=True)
    best = candidates[0]
    if len(candidates) > 1:
        second = candidates[1]
        if second[1].title != best[1].title and second[0] >= best[0] * 0.95:
            raise RuleError(
                "存在两个接近的工时矩阵候选："
                f"{best[1].title}!{best[2]} 与 {second[1].title}!{second[2]}，请用 --sheet 指定"
            )
    return best[1], best[2], best[3]


def find_alias_column(
    ws: Any,
    header_rows: Iterable[int],
    candidate_aliases: list[str],
    col_start: int,
    col_end: int,
) -> Optional[int]:
    valid_rows = [row for row in header_rows if 1 <= row <= ws.max_row]
    # 先找完整表头。这样标题“研发人员工时统计合计”不会抢在
    # 明确的“人员姓名”表头前被误判为姓名列。
    for exact_only in (True, False):
        for row in valid_rows:
            for col in range(max(1, col_start), min(ws.max_column, col_end) + 1):
                key = normalize_key(ws.cell(row, col).value)
                if not key:
                    continue
                if exact_only and key in candidate_aliases:
                    return col
                if not exact_only and any(alias in key for alias in candidate_aliases):
                    return col
    return None


def locate_layout(
    wb_formula: Any,
    wb_values: Any,
    rules: dict[str, Any],
    args: argparse.Namespace,
) -> Layout:
    project_master, master_sources = discover_project_master(wb_formula, rules)
    ws, code_row, code_matches = select_hours_sheet(
        wb_formula,
        wb_values,
        rules,
        project_master,
        args.sheet,
        args.project_header_row,
    )
    ws_values = wb_values[ws.title]
    first_project_col = min(col for col, _ in code_matches)
    last_project_col = max(col for col, _ in code_matches)

    name_row = None
    best_name_count = 0
    for row in range(code_row + 1, min(ws.max_row, code_row + 3) + 1):
        count = sum(
            1
            for col, _ in code_matches
            if normalize_text(ws.cell(row, col).value)
            and not isinstance(ws.cell(row, col).value, (int, float))
        )
        if count > best_name_count:
            best_name_count = count
            name_row = row
    if best_name_count < max(2, len(code_matches) // 3):
        name_row = None
    if name_row is None:
        raise RuleError("未找到模板的项目名称表头行；项目编码下一行必须填写项目名称")
    header_rows = sorted(set([code_row, *( [name_row] if name_row else [] )]))
    search_header_rows = range(max(1, code_row - 2), min(ws.max_row, (name_row or code_row) + 1) + 1)

    name_col = parse_column_override(args.name_column) or find_alias_column(
        ws,
        search_header_rows,
        aliases(rules, "employee_name_header"),
        1,
        first_project_col - 1,
    )
    if not name_col:
        raise RuleError("未能高置信度定位人员姓名列；请使用 --name-column 指定")
    id_col = parse_column_override(args.employee_id_column) or find_alias_column(
        ws,
        search_header_rows,
        aliases(rules, "employee_id_header"),
        1,
        first_project_col - 1,
    )
    month_col = find_alias_column(
        ws,
        search_header_rows,
        aliases(rules, "month_reference_header"),
        1,
        max(first_project_col - 1, 1),
    )
    if not month_col:
        raise RuleError("未定位到模板月份列（如“1月”或“一月”）；该列仅用于展示")

    stats_end = min(ws.max_column, last_project_col + int(rules["identification"]["scan_limits"]["statistics_columns_right"]))
    rd_col = find_alias_column(
        ws, search_header_rows, aliases(rules, "actual_hours_header"), last_project_col + 1, stats_end
    )
    public_rd_col = find_alias_column(
        ws, search_header_rows, aliases(rules, "allocated_hours_header"), last_project_col + 1, stats_end
    )
    hours_total_col = find_alias_column(
        ws, search_header_rows, aliases(rules, "total_hours_header"), last_project_col + 1, stats_end
    )
    missing_template_fields = [
        label
        for label, col in (("工时汇总", hours_total_col), ("RD工时", rd_col), ("公摊RD工时", public_rd_col))
        if col is None
    ]
    if missing_template_fields:
        raise RuleError(f"未定位到模板右侧字段：{'、'.join(missing_template_fields)}")
    if not (hours_total_col < rd_col < public_rd_col):
        raise RuleError("模板右侧字段顺序必须为：工时汇总、RD工时、公摊RD工时")

    projects: list[ProjectColumn] = []
    seen_codes: set[str] = set()
    master_matches = 0
    for col, code in code_matches:
        if code in seen_codes:
            raise RuleError(f"工时矩阵项目编码重复：{code}")
        seen_codes.add(code)
        header_name = normalize_text(ws.cell(name_row, col).value) if name_row else ""
        master_name = project_master.get(code, "")
        if master_name:
            master_matches += 1
            if header_name and normalize_key(header_name) != normalize_key(master_name):
                raise RuleError(
                    f"{ws.title}!{get_column_letter(col)}{name_row} 项目名称与主数据不一致："
                    f"{header_name!r} / {master_name!r}"
                )
        project_name = master_name or header_name
        if not project_name:
            raise RuleError(f"项目 {code} 缺少项目名称")
        projects.append(
            ProjectColumn(
                column=col,
                column_letter=get_column_letter(col),
                project_code=code,
                project_name=project_name,
                header_cell=f"{get_column_letter(col)}{code_row}",
                master_matched=bool(master_name),
            )
        )

    excluded = aliases(rules, "summary_row_labels")
    data_start = (name_row or code_row) + 1
    employees: list[EmployeeRow] = []
    blank_streak = 0
    for row in range(data_start, ws.max_row + 1):
        raw_name = ws.cell(row, name_col).value
        raw_id = ws.cell(row, id_col).value if id_col else None
        left_values = [ws.cell(row, col).value for col in range(1, first_project_col)]
        if any(is_alias(value, excluded) for value in left_values if value not in (None, "")):
            if employees:
                break
        name = normalize_text(raw_name)
        employee_id = normalize_text(raw_id)
        has_numeric_evidence = False
        for project in projects:
            value = cell_effective_value(ws, ws_values, row, project.column)
            if value not in (None, ""):
                try:
                    Decimal(str(value).replace(",", ""))
                    has_numeric_evidence = True
                except (InvalidOperation, ValueError):
                    has_numeric_evidence = True
        for col in (rd_col, public_rd_col, hours_total_col):
            value = cell_effective_value(ws, ws_values, row, col)
            if value not in (None, ""):
                has_numeric_evidence = True

        if not name:
            if employee_id or has_numeric_evidence:
                raise RuleError(f"{ws.title}!{get_column_letter(name_col)}{row} 姓名为空但该行存在人员/工时数据")
            blank_streak += 1
            if employees and blank_streak >= 2:
                break
            continue
        blank_streak = 0
        if is_alias(name, excluded):
            break
        if not has_numeric_evidence:
            continue
        employees.append(
            EmployeeRow(
                row=row,
                employee_id=employee_id,
                employee_name=name,
                source_cell=f"{get_column_letter(name_col)}{row}",
            )
        )

    if len(employees) < int(rules["identification"]["minimum_employee_rows"]):
        raise RuleError("有效人员行数量不足，无法确认人员 × 项目矩阵")
    ids = [person.employee_id for person in employees if person.employee_id]
    if len(ids) != len(set(ids)):
        raise RuleError("人员编号存在重复，无法唯一识别人员")
    if not id_col:
        names = [person.employee_name for person in employees]
        if len(names) != len(set(names)):
            raise RuleError("存在同名人员且没有人员编号，无法唯一识别")

    total_row = None
    for row in range(employees[-1].row + 1, min(ws.max_row, employees[-1].row + 15) + 1):
        row_values = [ws.cell(row, col).value for col in range(1, first_project_col)]
        if any(is_alias(value, aliases(rules, "grand_total_row")) for value in row_values if value not in (None, "")):
            total_row = row
            break

    summary_rows: dict[str, int] = {}
    label_groups = {
        "project_actual": aliases(rules, "project_actual_summary"),
        "project_allocated": aliases(rules, "project_allocated_summary"),
        "project_final": aliases(rules, "project_final_summary"),
    }
    for row in range(employees[-1].row + 1, min(ws.max_row, employees[-1].row + 15) + 1):
        row_values = [ws.cell(row, col).value for col in range(1, first_project_col)]
        for key, candidate_aliases in label_groups.items():
            if any(is_alias(value, candidate_aliases) for value in row_values if value not in (None, "")):
                summary_rows[key] = row

    master_ratio = master_matches / len(projects) if projects else 0
    confidence = Decimal("0.35")
    evidence = [f"横向识别 {len(projects)} 个项目列", f"纵向识别 {len(employees)} 名人员"]
    warnings: list[str] = []
    if any(normalize_key(item) in normalize_key(ws.title) for item in rules["identification"]["preferred_hours_sheet_names"]):
        confidence += Decimal("0.10")
        evidence.append("工作表名称命中工时表候选")
    if master_ratio >= Decimal("0.8"):
        confidence += Decimal("0.20")
        evidence.append(f"{master_matches}/{len(projects)} 个项目精确命中项目主数据")
    elif not project_master:
        evidence.append("项目由编码、名称、两行表头和数字矩阵定位（无需辅助项目表）")
    else:
        warnings.append(f"仅 {master_matches}/{len(projects)} 个项目命中项目主数据")
    if id_col:
        confidence += Decimal("0.08")
        evidence.append("存在明确人员编号列")
    confidence += Decimal("0.08")  # explicit name header is required above
    confidence += Decimal("0.06")  # template month is present, but only for display
    confidence += Decimal("0.15")  # 工时汇总、RD工时、公摊RD工时 are required above
    evidence.append("模板右侧工时汇总、RD工时、公摊RD工时字段齐全")
    if name_row:
        confidence += Decimal("0.08")
        evidence.append("存在项目名称表头行")
    if len(summary_rows) == 3:
        confidence += Decimal("0.06")
        evidence.append("底部三类项目汇总行全部定位")
    confidence = min(confidence, Decimal("1"))

    layout = Layout(
        sheet=ws.title,
        project_header_row=code_row,
        project_name_row=name_row,
        employee_id_column=id_col,
        employee_name_column=name_col,
        month_reference_column=month_col,
        source_rd_column=rd_col,
        source_public_rd_column=public_rd_col,
        source_hours_total_column=hours_total_col,
        employee_start_row=employees[0].row,
        employee_end_row=employees[-1].row,
        total_row=total_row,
        project_columns=projects,
        employees=employees,
        summary_rows=summary_rows,
        confidence=float(confidence),
        evidence=evidence + ([f"项目主数据来源：{', '.join(master_sources)}"] if master_sources else []),
        warnings=warnings,
    )
    threshold = float(rules["identification"]["auto_confirm_threshold"])
    if layout.confidence < threshold and not args.accept_low_confidence:
        raise RuleError(
            f"识别置信度 {layout.confidence:.1%} 低于自动处理阈值 {threshold:.1%}；"
            "请补充项目/人员主数据或用参数明确指定版式"
        )
    return layout



def calculate(
    wb_formula: Any,
    wb_values: Any,
    layout: Layout,
    rules: dict[str, Any],
) -> dict[str, Any]:
    """按每个人自己的RD项目占比，分配该人的公摊RD工时。"""
    ws = wb_formula[layout.sheet]
    ws_values = wb_values[layout.sheet]
    tolerance = Decimal(str(rules["calculation"]["reconciliation_tolerance"]))
    errors: list[str] = []
    warnings = list(layout.warnings)
    people: list[dict[str, Any]] = []
    project_actual = {project.project_code: Decimal("0") for project in layout.project_columns}
    project_allocated = {project.project_code: Decimal("0") for project in layout.project_columns}
    personal_checks: list[dict[str, Any]] = []

    for person in layout.employees:
        by_project: dict[str, Decimal] = {}
        for project in layout.project_columns:
            location = f"{layout.sheet}!{project.column_letter}{person.row}"
            hours = to_decimal(
                cell_effective_value(ws, ws_values, person.row, project.column),
                location,
                blank_zero=True,
            )
            if hours < 0:
                errors.append(f"{location} 出现负RD工时 {hours}")
            by_project[project.project_code] = hours
            project_actual[project.project_code] += hours

        matrix_rd_hours = sum(by_project.values(), Decimal("0"))
        rd_location = f"{layout.sheet}!{get_column_letter(layout.source_rd_column)}{person.row}"
        public_location = f"{layout.sheet}!{get_column_letter(layout.source_public_rd_column)}{person.row}"
        total_location = f"{layout.sheet}!{get_column_letter(layout.source_hours_total_column)}{person.row}"
        source_rd_hours = to_decimal(
            cell_effective_value(ws, ws_values, person.row, layout.source_rd_column),
            rd_location,
        )
        source_public_rd_hours = to_decimal(
            cell_effective_value(ws, ws_values, person.row, layout.source_public_rd_column),
            public_location,
        )
        source_hours_total = to_decimal(
            cell_effective_value(ws, ws_values, person.row, layout.source_hours_total_column),
            total_location,
        )
        derived_public_rd_hours = source_hours_total - source_rd_hours
        if derived_public_rd_hours < 0:
            errors.append(
                f"{person.employee_name} 的工时汇总 {source_hours_total} 小于RD工时 {source_rd_hours}，"
                "公摊RD工时不能为负数"
            )
        if abs(source_rd_hours - matrix_rd_hours) > tolerance:
            errors.append(
                f"{rd_location} 源表RD工时={source_rd_hours}，项目RD工时行合计={matrix_rd_hours}"
            )
        if abs(source_public_rd_hours - derived_public_rd_hours) > tolerance:
            errors.append(
                f"{public_location} 源表公摊RD工时={source_public_rd_hours}，"
                f"按工时汇总-RD工时计算={derived_public_rd_hours}"
            )
        if source_rd_hours == 0:
            if derived_public_rd_hours != 0:
                errors.append(
                    f"{person.employee_name} 的RD工时为0但公摊RD工时为{derived_public_rd_hours}，"
                    "无法按个人项目占比分摊"
                )
            ratio = Decimal("0")
        else:
            ratio = derived_public_rd_hours / source_rd_hours
        allocated_by_project = {code: hours * ratio for code, hours in by_project.items()}
        final_by_project = {
            code: by_project[code] + allocated_by_project[code] for code in by_project
        }
        for code, allocated in allocated_by_project.items():
            project_allocated[code] += allocated
        allocated_sum = sum(allocated_by_project.values(), Decimal("0"))
        final_sum = sum(final_by_project.values(), Decimal("0"))
        allocation_difference = allocated_sum - derived_public_rd_hours
        final_difference = final_sum - source_hours_total
        if abs(allocation_difference) > tolerance:
            errors.append(
                f"{person.employee_name} 项目公摊合计={allocated_sum}，"
                f"个人公摊RD工时={derived_public_rd_hours}"
            )
        if abs(final_difference) > tolerance:
            errors.append(
                f"{person.employee_name} 项目最终工时合计={final_sum}，工时汇总={source_hours_total}"
            )
        personal_checks.append(
            {
                "employee_name": person.employee_name,
                "allocated_status": "PASS" if abs(allocation_difference) <= tolerance else "FAIL",
                "allocated_difference": allocation_difference,
                "final_status": "PASS" if abs(final_difference) <= tolerance else "FAIL",
                "final_difference": final_difference,
            }
        )
        people.append(
            {
                "row": person.row,
                "employee_id": person.employee_id,
                "employee_name": person.employee_name,
                "by_project": by_project,
                "allocated_by_project": allocated_by_project,
                "final_by_project": final_by_project,
                "actual_hours": source_rd_hours,
                "allocated_hours": derived_public_rd_hours,
                "source_public_rd_hours": source_public_rd_hours,
                "final_hours": source_hours_total,
                "allocation_ratio": ratio,
            }
        )

    all_actual = sum(project_actual.values(), Decimal("0"))
    all_person_actual = sum((person["actual_hours"] for person in people), Decimal("0"))
    all_allocated = sum((person["allocated_hours"] for person in people), Decimal("0"))
    all_hours_total = sum((person["final_hours"] for person in people), Decimal("0"))
    if abs(all_actual - all_person_actual) > tolerance:
        errors.append(f"矩阵横纵RD工时不一致：项目={all_actual}，人员={all_person_actual}")

    project_final = {
        code: project_actual[code] + project_allocated[code] for code in project_actual
    }
    total_project_allocated = sum(project_allocated.values(), Decimal("0"))
    total_project_final = sum(project_final.values(), Decimal("0"))
    all_person_project_allocated = sum(
        (sum(person["allocated_by_project"].values(), Decimal("0")) for person in people),
        Decimal("0"),
    )
    checks = {
        "matrix_bidirectional_total": {
            "status": "PASS" if abs(all_actual - all_person_actual) <= tolerance else "FAIL",
            "project_total": all_actual,
            "person_total": all_person_actual,
            "difference": all_actual - all_person_actual,
        },
        "personal_allocation_conservation": {
            "status": "PASS" if all(
                check["allocated_status"] == "PASS" and check["final_status"] == "PASS"
                for check in personal_checks
            ) else "FAIL",
            "project_total": all_person_project_allocated,
            "person_total": all_allocated,
            "difference": all_person_project_allocated - all_allocated,
            "people": personal_checks,
        },
        "allocated_hours_conservation": {
            "status": "PASS" if abs(total_project_allocated - all_allocated) <= tolerance else "FAIL",
            "project_total": total_project_allocated,
            "person_total": all_allocated,
            "difference": total_project_allocated - all_allocated,
        },
        "final_hours_conservation": {
            "status": "PASS" if abs(total_project_final - all_hours_total) <= tolerance else "FAIL",
            "project_total": total_project_final,
            "person_total": all_hours_total,
            "difference": total_project_final - all_hours_total,
        },
    }
    for check_name, detail in checks.items():
        if detail["status"] != "PASS":
            errors.append(f"勾稽检查未通过：{check_name}")

    projects = []
    for project in layout.project_columns:
        code = project.project_code
        projects.append(
            {
                "project_code": code,
                "project_name": project.project_name,
                "source_column": project.column_letter,
                "actual_hours": project_actual[code],
                "actual_ratio": project_actual[code] / all_actual if all_actual else Decimal("0"),
                "allocated_hours": project_allocated[code],
                "final_hours": project_final[code],
            }
        )
    return {
        "status": "PASS" if not errors else "FAIL",
        "errors": errors,
        "warnings": warnings,
        "totals": {
            "actual_hours": all_actual,
            "allocated_hours": all_allocated,
            "final_hours": total_project_final,
            "hours_total": all_hours_total,
        },
        "checks": checks,
        "people": people,
        "projects": projects,
    }


def json_safe(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, dict):
        return {key: json_safe(item) for key, item in value.items()}
    if isinstance(value, list):
        return [json_safe(item) for item in value]
    return value


def layout_to_dict(layout: Layout) -> dict[str, Any]:
    data = asdict(layout)
    for key in (
        "employee_id_column",
        "employee_name_column",
        "month_reference_column",
        "source_rd_column",
        "source_public_rd_column",
        "source_hours_total_column",
    ):
        col = data[key]
        data[key] = get_column_letter(col) if col else None
    return data


def copy_cell_format(source: Any, target: Any) -> None:
    """Copy only presentation attributes; values/formulas are written explicitly."""
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)
    target.number_format = source.number_format


def quoted_sheet_name(title: str) -> str:
    return "'" + title.replace("'", "''") + "'"


def collect_formula_caches(path: Path) -> dict[str, dict[str, Any]]:
    """Read cached results already present in the uploaded workbook."""
    formula_book = load_workbook(path, data_only=False, read_only=False)
    value_book = load_workbook(path, data_only=True, read_only=False)
    caches: dict[str, dict[str, Any]] = {}
    for sheet in formula_book.worksheets:
        value_sheet = value_book[sheet.title]
        sheet_cache: dict[str, Any] = {}
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type != "f":
                    continue
                cached = value_sheet[cell.coordinate].value
                if isinstance(cached, (str, int, float, bool, Decimal)):
                    sheet_cache[cell.coordinate] = cached
        if sheet_cache:
            caches[sheet.title] = sheet_cache
    return caches


def formula_cache_text(value: Any) -> tuple[str, Optional[str]]:
    if isinstance(value, bool):
        return ("1" if value else "0", "b")
    if isinstance(value, Decimal):
        return (format(value, "f"), None)
    if isinstance(value, int):
        return (str(value), None)
    if isinstance(value, float):
        if not math.isfinite(value):
            raise RuleError("公式缓存值必须是有限数值")
        return (repr(value), None)
    return (str(value), "str")


def inject_formula_caches(path: Path, caches: dict[str, dict[str, Any]]) -> None:
    """Restore cached formula values so previews show results before Excel recalculates."""
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    ET.register_namespace("", main_ns)
    ET.register_namespace("r", rel_ns)
    with zipfile.ZipFile(path, "r") as source_zip:
        infos = source_zip.infolist()
        entries = {info.filename: source_zip.read(info.filename) for info in infos}
    workbook_xml = ET.fromstring(entries["xl/workbook.xml"])
    rels_xml = ET.fromstring(entries["xl/_rels/workbook.xml.rels"])
    relationships = {
        item.attrib["Id"]: item.attrib["Target"]
        for item in rels_xml.findall(f"{{{package_rel_ns}}}Relationship")
    }
    sheet_paths: dict[str, str] = {}
    sheets = workbook_xml.find(f"{{{main_ns}}}sheets")
    if sheets is not None:
        for sheet in sheets:
            title = sheet.attrib.get("name", "")
            rel_id = sheet.attrib.get(f"{{{rel_ns}}}id", "")
            target = relationships.get(rel_id, "")
            if not target:
                continue
            sheet_paths[title] = (
                target.lstrip("/")
                if target.startswith("/")
                else posixpath.normpath(posixpath.join("xl", target))
            )
    for title, cell_values in caches.items():
        sheet_path = sheet_paths.get(title)
        if not sheet_path or sheet_path not in entries:
            continue
        root = ET.fromstring(entries[sheet_path])
        for cell in root.findall(f".//{{{main_ns}}}c"):
            coordinate = cell.attrib.get("r", "")
            if coordinate not in cell_values or cell.find(f"{{{main_ns}}}f") is None:
                continue
            text, cell_type = formula_cache_text(cell_values[coordinate])
            value_node = cell.find(f"{{{main_ns}}}v")
            if value_node is None:
                value_node = ET.SubElement(cell, f"{{{main_ns}}}v")
            value_node.text = text
            if cell_type:
                cell.attrib["t"] = cell_type
            else:
                cell.attrib.pop("t", None)
        entries[sheet_path] = ET.tostring(root, encoding="utf-8", xml_declaration=True)

    temp_fd, temp_name = tempfile.mkstemp(prefix="rd-hours-cache-", suffix=".xlsx", dir=path.parent)
    os.close(temp_fd)
    temporary = Path(temp_name)
    try:
        with zipfile.ZipFile(temporary, "w") as target_zip:
            for info in infos:
                target_zip.writestr(info, entries[info.filename])
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def write_result_workbook(
    input_path: Path,
    output_path: Path,
    layout: Layout,
    result: dict[str, Any],
) -> None:
    """Write the allocation matrix at source right and a separate final matrix sheet."""
    if input_path.resolve() == output_path.resolve():
        raise RuleError("输出文件不能覆盖输入文件")
    if result["status"] != "PASS":
        raise RuleError("规则勾稽未通过，不生成正式结果")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(input_path, output_path)
    formula_caches = collect_formula_caches(input_path)
    workbook = load_workbook(output_path, data_only=False)
    for sheet_name in ("智能体测算_工时", "智能体识别日志", "RD工时加公摊"):
        if sheet_name in workbook.sheetnames:
            del workbook[sheet_name]

    source = workbook[layout.sheet]
    final_sheet = workbook.copy_worksheet(source)
    final_sheet.title = "RD工时加公摊"
    formula_caches[final_sheet.title] = dict(formula_caches.get(layout.sheet, {}))
    formula_caches.setdefault(layout.sheet, {})
    source_ref = quoted_sheet_name(layout.sheet)
    project_totals = {item["project_code"]: item for item in result["projects"]}
    right_start_col = layout.source_public_rd_column + 1
    right_end_col = right_start_col + len(layout.project_columns) - 1
    target_rows = [layout.project_header_row]
    if layout.project_name_row:
        target_rows.append(layout.project_name_row)
    target_rows.extend(person.row for person in layout.employees)
    if layout.total_row:
        target_rows.append(layout.total_row)

    existing_headers = [
        canonical_project_code(source.cell(layout.project_header_row, right_start_col + index).value)
        for index in range(len(layout.project_columns))
    ]
    expected_headers = [project.project_code for project in layout.project_columns]
    target_has_content = any(
        source.cell(row, col).value not in (None, "")
        for row in target_rows
        for col in range(right_start_col, right_end_col + 1)
    )
    if target_has_content and existing_headers != expected_headers:
        start = get_column_letter(right_start_col)
        end = get_column_letter(right_end_col)
        raise RuleError(f"源表右侧目标区域 {start}:{end} 已有其他内容，不能安全写入公摊RD工时表")

    for index, project in enumerate(layout.project_columns):
        destination_col = right_start_col + index
        source_letter = get_column_letter(project.column)
        destination_letter = get_column_letter(destination_col)
        source.column_dimensions[destination_letter].width = source.column_dimensions[source_letter].width
        for row in target_rows:
            copy_cell_format(source.cell(row, project.column), source.cell(row, destination_col))
        source.cell(layout.project_header_row, destination_col, project.project_code)
        if layout.project_name_row:
            source.cell(layout.project_name_row, destination_col, project.project_name)
        for person in result["people"]:
            row = person["row"]
            project_cell = f"{source_letter}{row}"
            rd_cell = f"{get_column_letter(layout.source_rd_column)}{row}"
            public_cell = f"{get_column_letter(layout.source_public_rd_column)}{row}"
            source.cell(
                row,
                destination_col,
                f'=IF({project_cell}="","",IF({rd_cell}=0,0,{project_cell}*{public_cell}/{rd_cell}))',
            )
            source.cell(row, destination_col).number_format = "0.00"
            cache_value = (
                ""
                if source.cell(row, project.column).value in (None, "")
                else person["allocated_by_project"][project.project_code]
            )
            formula_caches[layout.sheet][f"{destination_letter}{row}"] = cache_value
        if layout.total_row:
            source.cell(
                layout.total_row,
                destination_col,
                f"=SUM({destination_letter}{layout.employee_start_row}:{destination_letter}{layout.employee_end_row})",
            )
            source.cell(layout.total_row, destination_col).number_format = "0.00"
            formula_caches[layout.sheet][f"{destination_letter}{layout.total_row}"] = (
                project_totals[project.project_code]["allocated_hours"]
            )

    for project_index, project in enumerate(layout.project_columns):
        source_letter = get_column_letter(project.column)
        allocation_letter = get_column_letter(right_start_col + project_index)
        for person in result["people"]:
            row = person["row"]
            final_sheet.cell(
                row,
                project.column,
                f'=IF({source_ref}!{source_letter}{row}="","",'
                f"{source_ref}!{source_letter}{row}+{source_ref}!{allocation_letter}{row})",
            )
            final_sheet.cell(row, project.column).number_format = "0.00"
            cache_value = (
                ""
                if source.cell(row, project.column).value in (None, "")
                else person["final_by_project"][project.project_code]
            )
            formula_caches[final_sheet.title][f"{source_letter}{row}"] = cache_value
        if layout.total_row:
            final_sheet.cell(
                layout.total_row,
                project.column,
                f"=SUM({source_letter}{layout.employee_start_row}:{source_letter}{layout.employee_end_row})",
            )
            final_sheet.cell(layout.total_row, project.column).number_format = "0.00"
            formula_caches[final_sheet.title][f"{source_letter}{layout.total_row}"] = (
                project_totals[project.project_code]["final_hours"]
            )

    for person in result["people"]:
        row = person["row"]
        for col in (
            layout.source_hours_total_column,
            layout.source_rd_column,
            layout.source_public_rd_column,
        ):
            letter = get_column_letter(col)
            final_sheet.cell(row, col, f"={source_ref}!{letter}{row}")
            final_sheet.cell(row, col).number_format = "0.00"
        formula_caches[final_sheet.title][f"{get_column_letter(layout.source_hours_total_column)}{row}"] = person["final_hours"]
        formula_caches[final_sheet.title][f"{get_column_letter(layout.source_rd_column)}{row}"] = person["actual_hours"]
        formula_caches[final_sheet.title][f"{get_column_letter(layout.source_public_rd_column)}{row}"] = person["allocated_hours"]
    if layout.total_row:
        for col in (
            layout.source_hours_total_column,
            layout.source_rd_column,
            layout.source_public_rd_column,
        ):
            letter = get_column_letter(col)
            final_sheet.cell(layout.total_row, col, f"={source_ref}!{letter}{layout.total_row}")
            final_sheet.cell(layout.total_row, col).number_format = "0.00"
        formula_caches[final_sheet.title][f"{get_column_letter(layout.source_hours_total_column)}{layout.total_row}"] = result["totals"]["hours_total"]
        formula_caches[final_sheet.title][f"{get_column_letter(layout.source_rd_column)}{layout.total_row}"] = result["totals"]["actual_hours"]
        formula_caches[final_sheet.title][f"{get_column_letter(layout.source_public_rd_column)}{layout.total_row}"] = result["totals"]["allocated_hours"]

    log = workbook.create_sheet("智能体识别日志")
    dark_blue = "1F4E78"
    thin = Side(style="thin", color="B7C9D6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    log.append(["识别项目", "结果"])
    log_rows = [
        ("来源工作表", layout.sheet),
        ("项目编码表头行", layout.project_header_row),
        ("项目名称行", layout.project_name_row),
        ("人员编号列", get_column_letter(layout.employee_id_column) if layout.employee_id_column else "未提供"),
        ("姓名列", get_column_letter(layout.employee_name_column)),
        ("月份列", get_column_letter(layout.month_reference_column)),
        ("工时汇总列", get_column_letter(layout.source_hours_total_column)),
        ("RD工时列", get_column_letter(layout.source_rd_column)),
        ("公摊RD工时列", get_column_letter(layout.source_public_rd_column)),
        ("人员数据范围", f"{layout.employee_start_row}:{layout.employee_end_row}"),
        ("合计行", layout.total_row or "未提供"),
        ("右侧公摊表范围", f"{get_column_letter(right_start_col)}:{get_column_letter(right_end_col)}"),
        ("最终工时工作表", final_sheet.title),
        ("项目数", len(layout.project_columns)),
        ("人员数", len(layout.employees)),
        ("识别置信度", layout.confidence),
        ("状态", result["status"]),
        ("公式", "个人项目公摊RD工时=个人项目RD工时÷个人RD工时×个人公摊RD工时"),
        ("公摊来源", "公摊RD工时=工时汇总-RD工时；源公摊列用于勾稽"),
        ("舍入", "底层不舍入、不补尾差；单元格仅显示两位小数"),
        ("证据", "；".join(layout.evidence)),
        ("警告", "；".join(result["warnings"]) or "无"),
        ("错误", "；".join(result["errors"]) or "无"),
    ]
    for item in log_rows:
        log.append(item)
    for cell in log[1]:
        cell.fill = PatternFill("solid", fgColor=dark_blue)
        cell.font = Font(color="FFFFFF", bold=True)
    for row in log.iter_rows(min_row=1, max_row=log.max_row, min_col=1, max_col=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in range(2, log.max_row + 1):
        if log.cell(row, 1).value == "识别置信度":
            log.cell(row, 2).number_format = "0.0%"
    log.column_dimensions["A"].width = 22
    log.column_dimensions["B"].width = 100

    if workbook.calculation is None:
        workbook.calculation = CalcProperties()
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(output_path)
    inject_formula_caches(output_path, formula_caches)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="自动识别并测算研发人员项目工时")
    parser.add_argument("input", type=Path, help="待测算的 .xlsx 文件")
    parser.add_argument(
        "--rules",
        type=Path,
        default=Path(__file__).with_name("rules.yaml"),
        help="YAML 规则文件",
    )
    parser.add_argument("--report-json", type=Path, help="保存机器可读验证报告")
    parser.add_argument("--output", type=Path, help="输出含测算结果的新工作簿；不会覆盖源文件")
    parser.add_argument("--sheet", help="明确指定工时工作表")
    parser.add_argument("--project-header-row", type=int, help="明确指定项目编码表头行")
    parser.add_argument("--name-column", help="明确指定姓名列，例如 B 或 2")
    parser.add_argument("--employee-id-column", help="明确指定人员编号列")
    parser.add_argument(
        "--accept-low-confidence",
        action="store_true",
        help="经人工确认后允许低于自动阈值的定位结果继续计算",
    )
    return parser


def main(argv: Optional[list[str]] = None) -> int:
    args = build_parser().parse_args(argv)
    input_path = args.input.expanduser().resolve()
    if not input_path.exists():
        print(f"错误：找不到输入文件 {input_path}", file=sys.stderr)
        return 2
    if input_path.suffix.lower() != ".xlsx":
        print("错误：首版只支持 .xlsx；含宏文件请先另存为不含宏的副本", file=sys.stderr)
        return 2
    try:
        rules = load_rules(args.rules.expanduser().resolve())
        wb_formula = load_workbook(input_path, data_only=False, read_only=False)
        wb_values = load_workbook(input_path, data_only=True, read_only=False)
        layout = locate_layout(wb_formula, wb_values, rules, args)
        result = calculate(wb_formula, wb_values, layout, rules)
        report = {
            "rule_set": rules.get("rule_set", {}),
            "input_file": str(input_path),
            "layout": layout_to_dict(layout),
            "result": result,
        }
        report_safe = json_safe(report)
        if args.report_json:
            report_path = args.report_json.expanduser().resolve()
            report_path.parent.mkdir(parents=True, exist_ok=True)
            report_path.write_text(
                json.dumps(report_safe, ensure_ascii=False, indent=2), encoding="utf-8"
            )
        if args.output and result["status"] == "PASS":
            write_result_workbook(input_path, args.output.expanduser().resolve(), layout, result)
        print(json.dumps({
            "status": result["status"],
            "sheet": layout.sheet,
            "confidence": layout.confidence,
            "projects": len(layout.project_columns),
            "employees": len(layout.employees),
            "totals": json_safe(result["totals"]),
            "errors": result["errors"],
        }, ensure_ascii=False, indent=2))
        return 0 if result["status"] == "PASS" else 2
    except RuleError as exc:
        print(f"规则处理停止：{exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
