"""统一研发成本工作流：工时分配率 -> 直接人工 -> 运费。"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
import unicodedata
from collections import OrderedDict, defaultdict
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from types import SimpleNamespace
from typing import Any, Iterable

PACK_ROOT = Path(__file__).resolve().parent
RULE_PACKAGES = PACK_ROOT.parent
VENDOR_DIR = RULE_PACKAGES.parent / "vendor"
for import_path in (RULE_PACKAGES, VENDOR_DIR):
    if import_path.is_dir():
        sys.path.insert(0, str(import_path))

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from rd_hours_v1 import engine as hours_engine
import yaml


PACK_INFO = {"id": "rd_cost_workflow_v1", "name": "一键综合核算（工时＋直接人工＋其他费用）", "version": "2.2"}
MONEY = Decimal("0.01")
DISPLAY_RATE = Decimal("0.0001")
TOLERANCE = Decimal("0.000001")
PROJECT_RE = re.compile(r"^RD[0-9A-Z]+(?:-[0-9A-Z]+)*$", re.I)
MONTH_RE = re.compile(r"(?:1[0-2]|[1-9]|一|二|三|四|五|六|七|八|九|十|十一|十二)月")
CHINESE_MONTHS = {
    "一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6,
    "七": 7, "八": 8, "九": 9, "十": 10, "十一": 11, "十二": 12,
}
SUMMARY_WORDS = {"期初余额", "期末余额", "本期合计", "本年累计", "结转研发支出"}
OUTPUT_TEMPLATE = PACK_ROOT / "templates" / "研发成本输出模板.xlsx"
TEMPLATE_MONTH_SHEET = "项目分类明细汇总表1月"
COST_RULES_FILE = PACK_ROOT / "rules.yaml"
DIRECT_MATERIAL_CODE = "FYXM28-SYS"
PUBLIC_PROJECT_CODE = "RD9999"

MAIN_HEADERS = [
    "项目编码", "科目名称", "直接材料", "人员人工（含社保+公积金）", "运费", "折旧",
    "业务招待费", "差旅费", "交通费", "办公费", "会展费", "租赁费", "福利费",
    "长期费用摊销", "维修费", "无形资产摊销", "招聘费", "服务费", "专利服务",
    "水电费", "咨询/审计费", "测试费", "职工教育经费", "劳动保护费", "项目成本合计",
]



class RuleError(RuntimeError):
    pass


@dataclass
class Project:
    code: str
    name: str
    final_hours: Decimal = Decimal("0")
    rate: Decimal = Decimal("0")


@dataclass
class SalaryRow:
    source_row: int
    employee_id: str
    name: str
    salary: Decimal


@dataclass
class LedgerRecord:
    source_path: Path
    sheet: str
    fee_code: str
    fee_name: str
    project_code: str
    project_name: str
    posting_date: Any
    fiscal_year: Any
    period: Any
    voucher_word: str
    voucher_no: Any
    summary: str
    debit: Decimal
    credit: Decimal
    source_kind: str = "ledger"
    source_row: int = 0

    @property
    def amount(self) -> Decimal:
        # 已接入的“其他”规则只取借方；贷方不参与本次费用分配。
        return self.debit


@dataclass(frozen=True)
class LedgerSourceSheet:
    """一个需要原样保留、并在末尾追加分配表的原始明细账工作表。"""

    source_path: Path
    sheet: str
    fee_code: str
    fee_name: str

    @property
    def key(self) -> tuple[Path, str]:
        return self.source_path, self.sheet


def norm(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).strip()
    text = re.sub(r"[‐‑‒–—―−]", "-", text)
    return re.sub(r"\s+", " ", text)


def name_key(value: Any) -> str:
    return re.sub(r"[\s_()（）:：]+", "", norm(value)).lower()


def code_key(value: Any) -> str:
    return norm(value).replace(" ", "").replace("_", "-").upper()


def load_cost_rules() -> dict[str, Any]:
    if not COST_RULES_FILE.is_file():
        raise RuleError(f"缺少网页端成本规则包：{COST_RULES_FILE.name}")
    with COST_RULES_FILE.open("r", encoding="utf-8") as handle:
        rules = yaml.safe_load(handle) or {}
    items = rules.get("fee_items") or []
    if not items:
        raise RuleError("网页端成本规则包未定义费用映射")
    for item in items:
        required = ("code", "name", "detail_sheet", "master_header")
        if any(not norm(item.get(key)) for key in required):
            raise RuleError(f"网页端成本规则存在不完整费用映射：{item!r}")
        item["code"] = code_key(item["code"])
        item["name"] = norm(item["name"])
        item["detail_sheet"] = norm(item["detail_sheet"])
        item["master_header"] = norm(item["master_header"])
        item["aliases"] = [norm(alias) for alias in item.get("aliases", []) if norm(alias)]
    cleanup = rules.get("combination_cleanup") or {}
    required_cleanup_keys = ("source_headers", "primary_amount_headers", "amount_header", "output_mode")
    if any(not cleanup.get(key) for key in required_cleanup_keys):
        raise RuleError("网页端成本规则缺少核算维度与科目组合整理配置")
    cleanup["source_headers"] = [norm(value) for value in cleanup["source_headers"] if norm(value)]
    cleanup["primary_amount_headers"] = [norm(value) for value in cleanup["primary_amount_headers"] if norm(value)]
    # 输出表头需要保留用户指定的中文全角括号，不能用 ``norm`` 的 NFKC
    # 规范化结果替换显示文本。
    cleanup["amount_header"] = str(cleanup["amount_header"]).strip()
    if len(cleanup["source_headers"]) != 2 or not cleanup["primary_amount_headers"]:
        raise RuleError("网页端成本规则的核算维度与科目组合整理字段不完整")
    rules["fee_items"] = items
    rules["combination_cleanup"] = cleanup
    return rules


COST_RULES = load_cost_rules()
COMBINATION_CLEANUP_RULES = COST_RULES["combination_cleanup"]
FEE_RULES_BY_CODE = {item["code"]: item for item in COST_RULES["fee_items"]}
DIRECT_MATERIAL_RULE = {
    "code": DIRECT_MATERIAL_CODE,
    "name": "直接材料",
    "detail_sheet": "14.材料明细账",
    "master_header": "直接材料",
    "aliases": ["直接投入(材料、燃料)", "直接投入（材料、燃料）", "材料费", "材料明细账"],
}
FEE_RULES_BY_CODE[DIRECT_MATERIAL_CODE] = DIRECT_MATERIAL_RULE
EXPENSE_SHEETS = [(item["detail_sheet"], item["name"]) for item in COST_RULES["fee_items"]]


def fee_rule_for_text(*values: Any) -> dict[str, Any] | None:
    text = " ".join(norm(value) for value in values if norm(value)).lower()
    code_match = re.search(r"(FYXM\d+_SYS)", text, re.I)
    if code_match and code_key(code_match.group(1)) in FEE_RULES_BY_CODE:
        return FEE_RULES_BY_CODE[code_key(code_match.group(1))]
    if any(alias.lower() in text for alias in DIRECT_MATERIAL_RULE["aliases"]):
        return DIRECT_MATERIAL_RULE
    for item in COST_RULES["fee_items"]:
        candidates = [item["name"], *item.get("aliases", [])]
        if any(candidate.lower() in text for candidate in candidates):
            return item
    return None


def decimal_value(value: Any, location: str, *, blank_zero: bool = False, signed: bool = True) -> Decimal:
    if value is None or (isinstance(value, str) and not value.strip()):
        if blank_zero:
            return Decimal("0")
        raise RuleError(f"{location} 缺少数值")
    if isinstance(value, bool) or (isinstance(value, str) and value.startswith("#")):
        raise RuleError(f"{location} 不是有效数值")
    try:
        result = Decimal(str(value).replace(",", ""))
    except (InvalidOperation, ValueError) as exc:
        raise RuleError(f"{location} 不是有效数值：{value!r}") from exc
    if not signed and result < 0:
        raise RuleError(f"{location} 不能为负数")
    return result


def money(value: Decimal) -> Decimal:
    return value.quantize(MONEY, rounding=ROUND_HALF_UP)


def displayed_rate(value: Decimal) -> Decimal:
    """Return the exact rate represented by a percentage shown with two decimals."""
    return value.quantize(DISPLAY_RATE, rounding=ROUND_HALF_UP)


def allocate_by_displayed_rates(
    total: Decimal,
    keys: list[str],
    rates: dict[str, Decimal],
) -> dict[str, Decimal]:
    """Allocate from displayed 2-decimal percentages; put the money residual last."""
    positive = [key for key in keys if displayed_rate(rates.get(key, Decimal("0"))) > 0]
    if not positive:
        return {}
    output: dict[str, Decimal] = {}
    for key in positive[:-1]:
        output[key] = money(total * displayed_rate(rates[key]))
    output[positive[-1]] = money(total) - sum(output.values(), Decimal("0"))
    return output


def allocate_amount(total: Decimal, keys: list[str], weights: dict[str, Decimal]) -> dict[str, Decimal]:
    positive = [key for key in keys if weights.get(key, Decimal("0")) > 0]
    denominator = sum((weights[key] for key in positive), Decimal("0"))
    if not positive or denominator <= 0:
        return {}
    unrounded = {key: total * weights[key] / denominator for key in positive}
    output = {key: money(value) for key, value in unrounded.items()}
    residual = money(total) - sum(output.values(), Decimal("0"))
    if residual:
        # “其他”规则要求尾差进入未取整金额最大的项目；同额按项目编码排序。
        target = sorted(positive, key=lambda key: (-abs(unrounded[key]), key))[0]
        output[target] += residual
    return output


def cell_texts(ws: Any, row: int, max_col: int = 100) -> dict[str, int]:
    found: dict[str, int] = {}
    for col in range(1, min(ws.max_column, max_col) + 1):
        value = norm(ws.cell(row, col).value)
        if value and value not in found:
            found[value] = col
    return found


def copy_sheet_values(source: Any, target: Any) -> None:
    for merged in source.merged_cells.ranges:
        target.merge_cells(str(merged))
    for row in source.iter_rows():
        for source_cell in row:
            target_cell = target.cell(source_cell.row, source_cell.column, source_cell.value)
            if source_cell.has_style:
                target_cell._style = copy(source_cell._style)
            if source_cell.number_format:
                target_cell.number_format = source_cell.number_format
            if source_cell.alignment:
                target_cell.alignment = copy(source_cell.alignment)
            if source_cell.comment:
                target_cell.comment = copy(source_cell.comment)
            if source_cell.hyperlink:
                target_cell._hyperlink = copy(source_cell.hyperlink)
            if source_cell.protection:
                target_cell.protection = copy(source_cell.protection)
    for key, dimension in source.column_dimensions.items():
        target.column_dimensions[key].width = dimension.width
        target.column_dimensions[key].hidden = dimension.hidden
    for key, dimension in source.row_dimensions.items():
        target.row_dimensions[key].height = dimension.height
        target.row_dimensions[key].hidden = dimension.hidden
    target.freeze_panes = source.freeze_panes
    target.sheet_view.showGridLines = source.sheet_view.showGridLines
    target.sheet_properties = copy(source.sheet_properties)
    target.sheet_format = copy(source.sheet_format)


def calculate_legacy_hours(path: Path, project_names: dict[str, Project]) -> tuple[Any, dict[str, Any], OrderedDict[str, Project]]:
    """执行被锁定的简版工时口径，绝不根据人员或聊天上下文猜测。

    简版只接受一个明确结构：姓名列、紧邻 RD 连续项目块左侧的月份列（该列的
    数值就是个人工时汇总）、以及至少两个连续 RD 项目列。若存在多个候选结构或
    月份列不紧邻项目块，停止处理而不是选一个“看起来像”的表。
    """
    wb = load_workbook(path, data_only=True, read_only=False)
    candidates: list[tuple[Any, int, int, int, int, list[tuple[int, str]], str]] = []
    for ws in wb.worksheets:
        for header_row in range(1, min(ws.max_row, 40) + 1):
            headers = cell_texts(ws, header_row, 200)
            name_col = next((headers[item] for item in ("姓名", "人员姓名", "员工姓名") if item in headers), 0)
            if not name_col:
                continue
            matches = [
                (col, code_key(ws.cell(header_row, col).value))
                for col in range(1, min(ws.max_column, 300) + 1)
                if PROJECT_RE.fullmatch(code_key(ws.cell(header_row, col).value))
            ]
            groups: list[list[tuple[int, str]]] = []
            for match in matches:
                if not groups or match[0] != groups[-1][-1][0] + 1:
                    groups.append([match])
                else:
                    groups[-1].append(match)
            for project_cols in groups:
                if len(project_cols) < 2:
                    continue
                first_project_col = project_cols[0][0]
                explicit_total_col = next(
                    (headers[item] for item in ("工时汇总", "总工时", "工时总计") if item in headers),
                    0,
                )
                month_total_col = first_project_col - 1
                month_label = norm(ws.cell(header_row, month_total_col).value) if month_total_col > 0 else ""
                if explicit_total_col and explicit_total_col < first_project_col:
                    total_col, total_basis = explicit_total_col, "明确的工时汇总列"
                elif MONTH_RE.fullmatch(month_label):
                    total_col, total_basis = month_total_col, f"紧邻项目块的{month_label}总工时列"
                else:
                    continue
                id_col = next((headers[item] for item in ("序号", "人员编号", "员工编号", "工号") if item in headers), 0)
                candidates.append((ws, header_row, id_col, name_col, total_col, project_cols, total_basis))
    if not candidates:
        raise RuleError(
            "未找到锁定的简版工时结构：需有姓名列、紧邻 RD 项目块左侧的月份总工时列，"
            "或明确的工时汇总列"
        )
    if len(candidates) > 1:
        locations = "、".join(f"{ws.title}!{header_row}" for ws, header_row, *_ in candidates)
        raise RuleError(f"发现多个简版工时矩阵候选（{locations}），不能依靠上下文猜测，请保留唯一源表")
    ws, header_row, id_col, name_col, total_col, project_cols, total_basis = candidates[0]
    project_actual = {code: Decimal("0") for _, code in project_cols}
    project_allocated = {code: Decimal("0") for _, code in project_cols}
    people: list[dict[str, Any]] = []
    blanks = 0
    for row in range(header_row + 1, min(ws.max_row, header_row + 2000) + 1):
        name = norm(ws.cell(row, name_col).value)
        employee_id = norm(ws.cell(row, id_col).value) if id_col else ""
        if not name and not employee_id:
            blanks += 1
            if people and blanks >= 5:
                break
            continue
        blanks = 0
        if not name or name in {"合计", "总计", "小计"}:
            continue
        by_project = {code: decimal_value(ws.cell(row, col).value, f"{ws.title}!{get_column_letter(col)}{row}", blank_zero=True, signed=False) for col, code in project_cols}
        rd_hours = sum(by_project.values(), Decimal("0"))
        hours_total = decimal_value(ws.cell(row, total_col).value, f"{ws.title}!{get_column_letter(total_col)}{row}", signed=False)
        public_hours = hours_total - rd_hours
        if public_hours < 0:
            raise RuleError(f"{name} 的工时汇总小于RD项目工时合计")
        if rd_hours == 0 and public_hours != 0:
            raise RuleError(f"{name} 的RD工时为0但公摊RD工时为{public_hours}，不能猜测分配")
        ratio = public_hours / rd_hours if rd_hours else Decimal("0")
        allocated = {code: value * ratio for code, value in by_project.items()}
        final = {code: by_project[code] + allocated[code] for code in by_project}
        for code in by_project:
            project_actual[code] += by_project[code]
            project_allocated[code] += allocated[code]
        people.append({
            "row": row,
            "employee_id": employee_id,
            "employee_name": name,
            "by_project": by_project,
            "allocated_by_project": allocated,
            "final_by_project": final,
            "actual_hours": rd_hours,
            "allocated_hours": public_hours,
            "source_public_rd_hours": public_hours,
            "final_hours": hours_total,
            "allocation_ratio": ratio,
        })
    if not people:
        raise RuleError("工时表未找到人员明细")
    total_final = sum((person["final_hours"] for person in people), Decimal("0"))
    project_items: list[dict[str, Any]] = []
    projects: OrderedDict[str, Project] = OrderedDict()
    for _, code in project_cols:
        final_hours = project_actual[code] + project_allocated[code]
        project_name = project_names.get(code, Project(code, code)).name
        project_items.append({
            "project_code": code,
            "project_name": project_name,
            "actual_hours": project_actual[code],
            "allocated_hours": project_allocated[code],
            "final_hours": final_hours,
        })
        projects[code] = Project(code, project_name, final_hours, final_hours / total_final if total_final else Decimal("0"))
    result = {
        "status": "PASS",
        "errors": [],
        "warnings": [
            "输入为锁定的简版工时模板；"
            f"个人工时汇总取自{total_basis}；"
            "所有人员统一按 项目源工时 × (工时汇总 - 项目源工时合计) ÷ 项目源工时合计 计算，"
            "不为任何人员设置例外"
        ],
        "totals": {
            "actual_hours": sum(project_actual.values(), Decimal("0")),
            "allocated_hours": sum(project_allocated.values(), Decimal("0")),
            "final_hours": total_final,
            "hours_total": total_final,
        },
        "people": people,
        "projects": project_items,
    }
    return SimpleNamespace(sheet=ws.title), result, projects


def calculate_hours(path: Path, project_names: dict[str, Project]) -> tuple[Any, dict[str, Any], OrderedDict[str, Project]]:
    rules = hours_engine.load_rules(RULE_PACKAGES / "rd_hours_v1" / "rules.yaml")
    formula_wb = load_workbook(path, data_only=False, read_only=False)
    values_wb = load_workbook(path, data_only=True, read_only=False)
    args = SimpleNamespace(
        sheet=None,
        project_header_row=None,
        name_column=None,
        employee_id_column=None,
        accept_low_confidence=False,
    )
    try:
        layout = hours_engine.locate_layout(formula_wb, values_wb, rules, args)
        result = hours_engine.calculate(formula_wb, values_wb, layout, rules)
    except hours_engine.RuleError as exc:
        if (
            "项目名称表头行" in str(exc)
            or "未找到“公摊RD工时”" in str(exc)
            or "未定位到模板右侧字段" in str(exc)
            or "未找到模板" in str(exc)
        ):
            return calculate_legacy_hours(path, project_names)
        raise RuleError(str(exc)) from exc
    if result["status"] != "PASS":
        raise RuleError("工时规则未通过：" + "；".join(result["errors"]))
    total = result["totals"]["final_hours"]
    if total <= 0:
        raise RuleError("项目最终工时合计必须大于0")
    projects: OrderedDict[str, Project] = OrderedDict()
    for item in result["projects"]:
        projects[item["project_code"]] = Project(
            item["project_code"], item["project_name"], item["final_hours"], item["final_hours"] / total
        )
    rate_sum = sum((project.rate for project in projects.values()), Decimal("0"))
    if abs(rate_sum - Decimal("1")) > TOLERANCE:
        raise RuleError(f"工时分配率合计不是100%：{rate_sum}")
    return layout, result, projects


def hours_workbook_month(path: Path, sheet_name: str) -> int:
    """读取工时表头中独立的月份标签，仅用于期间一致性提示。"""
    wb = load_workbook(path, data_only=True, read_only=False)
    ws = safe_sheet(wb, sheet_name)
    for row in range(1, min(ws.max_row, 5) + 1):
        for col in range(1, min(ws.max_column, 40) + 1):
            label = norm(ws.cell(row, col).value)
            match = MONTH_RE.fullmatch(label)
            if not match:
                continue
            month_label = match.group(0)[:-1]
            return CHINESE_MONTHS[month_label] if month_label in CHINESE_MONTHS else int(month_label)
    return 0


def salary_header(path: Path) -> tuple[Any, Any, int, dict[str, int], int]:
    wb = load_workbook(path, data_only=True, read_only=False)
    name_aliases = {"人员名称", "姓名", "员工姓名", "研发人员"}
    id_aliases = {"序号", "人员编号", "员工编号", "工号", "人员编码"}
    for ws in wb.worksheets:
        for row in range(1, min(ws.max_row, 40) + 1):
            headers = cell_texts(ws, row, 120)
            name_col = next((headers[item] for item in name_aliases if item in headers), None)
            salary_col = next((col for text, col in headers.items() if text in {"工资费用", "工资", "应发工资", "实发工资", "工资合计"} or "月工资费用" in text), None)
            if not name_col or not salary_col:
                continue
            id_col = next((headers[item] for item in id_aliases if item in headers), 0)
            return wb, ws, row, {"name": name_col, "id": id_col, "salary": salary_col}, salary_col
    raise RuleError("工资表未找到“人员名称/姓名”和“工资费用/某月工资费用”表头")


def read_salary(path: Path) -> tuple[list[SalaryRow], str, int]:
    _, ws, header_row, headers, _ = salary_header(path)
    rows: list[SalaryRow] = []
    seen_names: set[str] = set()
    blanks = 0
    month = 0
    salary_header_text = norm(ws.cell(header_row, headers["salary"]).value)
    match = MONTH_RE.search(salary_header_text)
    if match:
        month_label = match.group(0)[:-1]
        month = CHINESE_MONTHS[month_label] if month_label in CHINESE_MONTHS else int(month_label)
    for row in range(header_row + 1, min(ws.max_row, header_row + 2000) + 1):
        name = norm(ws.cell(row, headers["name"]).value)
        employee_id = norm(ws.cell(row, headers["id"]).value) if headers["id"] else ""
        if name in {"合计", "总计", "小计"}:
            break
        if not name:
            blanks += 1
            if rows and blanks >= 5:
                break
            continue
        blanks = 0
        salary = decimal_value(ws.cell(row, headers["salary"]).value, f"{ws.title}!{row}行工资费用", blank_zero=True)
        key = name_key(name)
        if key in seen_names:
            raise RuleError(f"工资表人员姓名重复，无法唯一匹配：{name}")
        seen_names.add(key)
        rows.append(SalaryRow(row, employee_id, name, salary))
    if not rows:
        raise RuleError("工资表未找到人员明细")
    return rows, f"{ws.title}!{get_column_letter(headers['name'])}{header_row}:{get_column_letter(headers['salary'])}{header_row + len(rows) + 1}", month


def read_salary_project_names(path: Path) -> dict[str, str]:
    _, ws, header_row, headers, _ = salary_header(path)
    names: dict[str, str] = {}
    for col in range(headers["salary"] + 1, min(ws.max_column, headers["salary"] + 120) + 1):
        code = code_key(ws.cell(header_row, col).value)
        if not PROJECT_RE.fullmatch(code):
            continue
        project_name = norm(ws.cell(header_row + 1, col).value)
        if project_name:
            names.setdefault(code, project_name)
    return names


def allocate_salary(
    salary_rows: list[SalaryRow],
    hours_result: dict[str, Any],
    project_codes: list[str],
    warnings: list[str],
) -> tuple[list[dict[str, Any]], dict[str, Decimal], Decimal, Decimal, Decimal]:
    hours_people = hours_result["people"]
    by_name: dict[str, dict[str, Any]] = {}
    for person in hours_people:
        key = name_key(person["employee_name"])
        if key in by_name:
            raise RuleError(f"工时表姓名重复，无法唯一匹配：{person['employee_name']}")
        by_name[key] = person

    output: list[dict[str, Any]] = []
    project_totals = {code: Decimal("0") for code in project_codes}
    salary_total = sum((row.salary for row in salary_rows), Decimal("0"))
    matched_salary = Decimal("0")
    allocated_total = Decimal("0")
    unmatched_people: list[str] = []
    for salary_row in salary_rows:
        matched = by_name.get(name_key(salary_row.name))
        match_type = "姓名" if matched is not None else ""
        row_result = {"salary": salary_row, "matched": matched, "match_type": match_type, "amounts": {}}
        if matched is None:
            unmatched_people.append(salary_row.name)
            output.append(row_result)
            continue
        denominator = matched["final_hours"]
        weights = {code: matched["final_by_project"].get(code, Decimal("0")) for code in project_codes}
        if denominator <= 0 or abs(sum(weights.values(), Decimal("0")) - denominator) > TOLERANCE:
            raise RuleError(f"{salary_row.name} 的项目最终工时与工时汇总不一致")
        amounts = allocate_amount(salary_row.salary, project_codes, weights)
        matched_salary += salary_row.salary
        for code, amount in amounts.items():
            project_totals[code] += amount
            allocated_total += amount
        row_result["amounts"] = amounts
        output.append(row_result)
    if matched_salary != allocated_total:
        raise RuleError(f"直接人工未勾稽：可分配工资={matched_salary}，已分配={allocated_total}")
    if unmatched_people:
        warnings.append("按姓名未匹配人员（工资保留未分配）：" + "、".join(unmatched_people))
    return output, project_totals, salary_total, matched_salary, salary_total - matched_salary


def parse_dimension_text(values: Iterable[Any]) -> tuple[str, str, str, str]:
    text = " / ".join(norm(value) for value in values if norm(value))
    fee = re.search(r"费用项目[^\[]*\[([A-Z0-9_]+)\]\s*([^/]+)", text, re.I)
    project = re.search(r"研发项目[^\[]*\[(RD[0-9A-Z-]+)\]\s*([^/]+)", text, re.I)
    return (
        fee.group(1).upper() if fee else "",
        norm(fee.group(2)) if fee else "",
        code_key(project.group(1)) if project else "",
        norm(project.group(2)) if project else "",
    )


def month_from_summary_title(title: str) -> int:
    match = re.search(r"项目分类明细汇总表\s*([1-9]|1[0-2])月", norm(title))
    return int(match.group(1)) if match else 0


def parse_summary_sheet(path: Path, ws: Any, warnings: list[str]) -> tuple[list[LedgerRecord] | None, int]:
    """读取核算维度的项目分类汇总表。

    该表是本工作流的合法来源，不能再被误判为“没有摘要和借方列”而静默跳过。
    它同时会在导出时逐格映射到模板中的同构汇总页。
    """
    headers = {norm(label): col for label, col in cell_texts(ws, 1, 40).items()}
    required = [norm(item) for item in COST_RULES.get("summary_input", {}).get("required_headers", [])]
    if not required or any(item not in headers for item in required):
        return None, 0
    fee_col, project_col, amount_col = (headers[required[1]], headers[required[2]], headers[required[3]])
    inferred_month = month_from_summary_title(ws.title)
    records: list[LedgerRecord] = []
    current_fee = ""
    for row in range(2, ws.max_row + 1):
        fee_text = norm(ws.cell(row, fee_col).value)
        project_text = norm(ws.cell(row, project_col).value)
        amount_value = ws.cell(row, amount_col).value
        if fee_text:
            current_fee = fee_text
        source_fee = fee_text or current_fee
        if not source_fee:
            continue
        if source_fee.endswith("小计") or source_fee in {"五险合计", "总计", "合计"}:
            continue
        rule = fee_rule_for_text(source_fee)
        if rule is None:
            continue
        parsed = re.search(r"(RD[0-9A-Z]+(?:-[0-9A-Z]+)*)\s*(.*)", project_text, re.I)
        if not parsed:
            if not project_text and COST_RULES.get("summary_input", {}).get("blank_project_policy") == "rd9999_public_allocation":
                project_code, project_name = "RD9999", "研发公摊"
                warnings.append(f"{ws.title}!{get_column_letter(project_col)}{row} 未填写研发项目，已按规则作为RD9999研发公摊分配")
            else:
                warnings.append(f"{ws.title}!{get_column_letter(project_col)}{row} 未识别研发项目，未进入费用分配")
                continue
        else:
            project_code, project_name = code_key(parsed.group(1)), norm(parsed.group(2))
        if not PROJECT_RE.fullmatch(project_code):
            continue
        amount = decimal_value(amount_value, f"{ws.title}!{get_column_letter(amount_col)}{row}", blank_zero=True)
        records.append(LedgerRecord(
            source_path=path,
            sheet=ws.title,
            fee_code=rule["code"],
            fee_name=rule["name"],
            project_code=project_code,
            project_name=project_name,
            posting_date=None,
            fiscal_year=date.today().year,
            period=inferred_month or None,
            voucher_word="",
            voucher_no=None,
            summary=source_fee,
            debit=amount,
            credit=Decimal("0"),
            source_kind="summary",
            source_row=row,
        ))
    return records, inferred_month


def locate_combination_summary_sheet(path: Path) -> tuple[Path, str, int] | None:
    """定位“核算维度与科目组合”导出的交叉汇总结构。

    该结构的表头同一行包含“费用项目”“研发项目”和位于研发
    项目右侧的金额列。既接受原始导出的“借方发生(原币)”，也接受
    操作台已整理过的“费用金额（元）”。它是核算维度金额来源，不是项目清单。
    """
    wb = load_workbook(path, data_only=True, read_only=False)
    for ws in wb.worksheets:
        for row in range(1, min(ws.max_row, 100) + 1):
            headers = cell_texts(ws, row, 100)
            if "费用项目" not in headers or "研发项目" not in headers:
                continue
            project_col = headers["研发项目"]
            amount_headers = {
                *COMBINATION_CLEANUP_RULES["primary_amount_headers"],
                norm(COMBINATION_CLEANUP_RULES["amount_header"]),
            }
            debit_cols = [
                col for label, col in headers.items()
                if label in amount_headers and col > project_col
            ]
            if debit_cols:
                return path, ws.title, row
    return None


def combination_filename_matches(path: Path) -> bool:
    """按包含词识别组合表文件名提示，不锁定完整文件名。"""
    groups = (
        COST_RULES.get("input_recognition", {})
        .get("combination_filename_keyword_groups", [])
    )
    if not groups:
        return False
    filename = re.sub(r"[\s_.-]+", "", norm(path.stem)).lower()
    return all(
        any(re.sub(r"[\s_.-]+", "", norm(keyword)).lower() in filename for keyword in group)
        for group in groups
        if group
    )


def remove_top_rows_preserving_merges(ws: Any, row_count: int) -> None:
    """删除来源报表的前置表头，同时保持保留下来的合并单元格位置正确。"""
    if row_count <= 0:
        return
    merged_ranges = list(ws.merged_cells.ranges)
    crossing = [str(item) for item in merged_ranges if item.min_row <= row_count < item.max_row]
    if crossing:
        raise RuleError(
            f"{ws.title} 的前置表头与正文合并单元格交叉（{'、'.join(crossing[:5])}），不能安全删除"
        )
    for merged in merged_ranges:
        ws.unmerge_cells(str(merged))
    ws.delete_rows(1, row_count)
    for merged in merged_ranges:
        if merged.min_row <= row_count:
            continue
        ws.merge_cells(
            start_row=merged.min_row - row_count,
            start_column=merged.min_col,
            end_row=merged.max_row - row_count,
            end_column=merged.max_col,
        )


def combination_amount_columns(ws: Any, header_row: int, project_col: int) -> list[int]:
    """定位研发项目右侧所有借/贷方金额列，首个借方列是唯一保留金额列。"""
    related: list[int] = []
    for col in range(project_col + 1, ws.max_column + 1):
        label = norm(ws.cell(header_row, col).value)
        if (
            "借方发生" in label
            or "贷方发生" in label
            or label == norm(COMBINATION_CLEANUP_RULES["amount_header"])
        ):
            related.append(col)
    return related


def normalize_combination_workbook(source_path: Path, output_path: Path) -> dict[str, Any]:
    """生成“核算维度与科目组合”整理副本。

    规则严格对应用户确认的版式：保留费用项目、研发项目和首个借方金额列；
    删除这三个字段之上的报表表头；将保留金额列命名为“费用金额（元）”；
    清空其余借/贷方重复金额列。该函数只操作来源文件的副本。
    """
    if source_path.resolve() == output_path.resolve():
        raise RuleError("核算维度与科目组合原文件只读，整理结果必须另存为新工作簿")
    if not source_path.is_file():
        raise RuleError(f"未找到核算维度与科目组合文件：{source_path.name}")
    if COMBINATION_CLEANUP_RULES.get("output_mode") != "copied_workbook":
        raise RuleError("当前规则不允许以来源副本方式整理核算维度与科目组合")

    located = locate_combination_summary_sheet(source_path)
    if located is None:
        raise RuleError("未找到同一行包含“费用项目”“研发项目”和借方发生额的核算维度与科目组合表头")
    _, sheet_name, header_row = located
    shutil.copy2(source_path, output_path)
    wb = load_workbook(output_path, data_only=False, read_only=False)
    ws = wb[sheet_name]

    headers = cell_texts(ws, header_row, 100)
    fee_header, project_header = COMBINATION_CLEANUP_RULES["source_headers"]
    if fee_header not in headers or project_header not in headers:
        raise RuleError(f"{ws.title}!{header_row} 未找到费用项目或研发项目表头")
    project_col = headers[project_header]
    amount_columns = combination_amount_columns(ws, header_row, project_col)
    primary_amount_headers = {
        *COMBINATION_CLEANUP_RULES["primary_amount_headers"],
        norm(COMBINATION_CLEANUP_RULES["amount_header"]),
    }
    primary_amount_columns = [
        col for col in amount_columns
        if norm(ws.cell(header_row, col).value) in primary_amount_headers
    ]
    if not primary_amount_columns:
        raise RuleError(f"{ws.title}!{header_row} 未找到研发项目右侧的借方发生额列")
    amount_col = primary_amount_columns[0]

    removed_rows = header_row - 1 if COMBINATION_CLEANUP_RULES.get("remove_rows_before_header") else 0
    remove_top_rows_preserving_merges(ws, removed_rows)
    header_row -= removed_rows
    ws.cell(header_row, amount_col).value = COMBINATION_CLEANUP_RULES["amount_header"]

    cleared_columns: list[str] = []
    if COMBINATION_CLEANUP_RULES.get("clear_redundant_amount_columns"):
        for col in amount_columns:
            if col == amount_col:
                continue
            for row in range(header_row, ws.max_row + 1):
                cell = ws.cell(row, col)
                if not isinstance(cell, MergedCell):
                    cell.value = None
            cleared_columns.append(get_column_letter(col))

    retained_rows = sum(
        1
        for row in range(header_row + 1, ws.max_row + 1)
        if any(ws.cell(row, col).value not in (None, "") for col in (headers[fee_header], project_col, amount_col))
    )
    wb.save(output_path)
    return {
        "sheet": sheet_name,
        "header_row": header_row,
        "retained_rows": retained_rows,
        "primary_amount_column": get_column_letter(amount_col),
        "cleared_amount_columns": cleared_columns,
        "output_path": str(output_path),
    }


def parse_combination_summary_sheet(
    path: Path,
    ws: Any,
    header_row: int,
    warnings: list[str],
) -> tuple[list[LedgerRecord], int]:
    """把“核算维度与科目组合”交叉表规范化为分类汇总记录。"""
    headers = cell_texts(ws, header_row, 100)
    fee_col = headers["费用项目"]
    project_col = headers["研发项目"]
    primary_amount_headers = {
        *COMBINATION_CLEANUP_RULES["primary_amount_headers"],
        norm(COMBINATION_CLEANUP_RULES["amount_header"]),
    }
    debit_cols = [
        col for label, col in headers.items()
        if label in primary_amount_headers and col > project_col
    ]
    if not debit_cols:
        raise RuleError(f"{ws.title}!{header_row} 未找到研发项目右侧的借方发生或费用金额列")
    amount_col = min(debit_cols)
    inferred_month = month_from_summary_title(ws.title)
    records: list[LedgerRecord] = []
    current_fee_text = ""
    current_fee_code = ""
    current_fee_name = ""
    for row in range(header_row + 1, ws.max_row + 1):
        fee_text = norm(ws.cell(row, fee_col).value)
        project_text = norm(ws.cell(row, project_col).value)
        if fee_text in {"总计", "合计"}:
            break
        if fee_text and not fee_text.endswith("小计"):
            fee_match = re.search(r"(FYXM\d+[_-]SYS)\s*(.*)", fee_text, re.I)
            if fee_match:
                current_fee_text = fee_text
                current_fee_code = code_key(fee_match.group(1))
                current_fee_name = norm(fee_match.group(2))
                known_rule = FEE_RULES_BY_CODE.get(current_fee_code)
                if known_rule is not None:
                    current_fee_name = known_rule["name"]
        if not current_fee_code or not project_text or fee_text.endswith("小计"):
            continue
        project_match = re.search(r"(RD\s*[0-9A-Z]+(?:\s*-\s*[0-9A-Z]+)*)\s*(.*)", project_text, re.I)
        if project_match is None:
            warnings.append(f"{ws.title}!{get_column_letter(project_col)}{row} 未识别研发项目，未进入费用分配")
            continue
        project_code = code_key(project_match.group(1))
        project_name = norm(project_match.group(2))
        if not PROJECT_RE.fullmatch(project_code):
            continue
        amount = decimal_value(
            ws.cell(row, amount_col).value,
            f"{ws.title}!{get_column_letter(amount_col)}{row}",
            blank_zero=True,
        )
        records.append(LedgerRecord(
            source_path=path,
            sheet=ws.title,
            fee_code=current_fee_code,
            fee_name=current_fee_name or current_fee_text,
            project_code=project_code,
            project_name=project_name,
            posting_date=None,
            fiscal_year=date.today().year,
            period=inferred_month or None,
            voucher_word="",
            voucher_no=None,
            summary=current_fee_text,
            debit=amount,
            credit=Decimal("0"),
            source_kind="combination_summary",
            source_row=row,
        ))
    return records, inferred_month


def parse_ledger_file(path: Path, warnings: list[str]) -> tuple[list[LedgerRecord], int, list[LedgerSourceSheet]]:
    wb = load_workbook(path, data_only=True, read_only=False)
    records: list[LedgerRecord] = []
    source_sheets: list[LedgerSourceSheet] = []
    inferred_month = 0
    for ws in wb.worksheets:
        summary_records, summary_month = parse_summary_sheet(path, ws, warnings)
        if summary_records is not None:
            records.extend(summary_records)
            inferred_month = inferred_month or summary_month
            continue
        header_row = 0
        columns: dict[str, int] = {}
        for row in range(1, min(ws.max_row, 30) + 1):
            headers = cell_texts(ws, row, 80)
            if "摘要" in headers and ("借方" in headers or "贷方" in headers):
                header_row = row
                columns = headers
                break
        if not header_row:
            continue
        preface = [ws.cell(row, col).value for row in range(1, header_row) for col in range(1, min(ws.max_column, 20) + 1)]
        fee_code, fee_name, fixed_project_code, fixed_project_name = parse_dimension_text(preface + [ws.title])
        rule = FEE_RULES_BY_CODE.get(code_key(fee_code)) if fee_code else fee_rule_for_text(path.stem, ws.title, *preface)
        fee_col = next((columns[item] for item in ("费用项目", "费用类别", "费用科目") if item in columns), 0)
        if rule is not None:
            fee_code, fee_name = rule["code"], rule["name"]
        if (not fee_name or rule is None) and not fee_col:
            if fee_name:
                warnings.append(f"{ws.title} 的{fee_name}不属于固定其他费用明细页，已跳过")
            else:
                warnings.append(f"{ws.title} 未识别费用项目，已跳过")
            continue
        if rule is not None:
            source_sheets.append(LedgerSourceSheet(path, ws.title, fee_code, fee_name))
        for value in preface:
            match = re.search(r"期间.*?(\d{4})年(\d{1,2})月", norm(value))
            if match:
                inferred_month = inferred_month or int(match.group(2))
        project_col = next((columns[item] for item in ("研发项目", "项目编码") if item in columns), 0)
        project_name_col = next((columns[item] for item in ("研发名称", "项目名称") if item in columns), 0)
        for row in range(header_row + 1, ws.max_row + 1):
            summary = norm(ws.cell(row, columns["摘要"]).value)
            if not summary or any(word in summary for word in SUMMARY_WORDS):
                continue
            # 部分旧版工作簿会在凭证明细下方（运费页甚至是隐藏行）
            # 再嵌入“项目编号/工时合计/费用”或“项目编码/科目名称/
            # 费用合计”表。这是本页的分配结果而不是凭证，遇到时必须结束
            # 明细解析，否则“工时合计”会被误当成借方数值。
            embedded_labels = [
                norm(ws.cell(row, col).value)
                for col in range(columns["摘要"], min(ws.max_column, columns["摘要"] + 4) + 1)
            ]
            if summary in {"项目编号", "项目编码"} and (
                "工时合计" in embedded_labels
                or "科目名称" in embedded_labels
                or any("费用合计" in label for label in embedded_labels)
            ):
                break
            row_rule = rule
            if fee_col:
                row_rule = fee_rule_for_text(ws.cell(row, fee_col).value) or row_rule
            if row_rule is None:
                continue
            row_fee_code, row_fee_name = row_rule["code"], row_rule["name"]
            debit = decimal_value(ws.cell(row, columns.get("借方", 0)).value if columns.get("借方") else None, f"{ws.title}!{row}借方", blank_zero=True)
            credit = decimal_value(ws.cell(row, columns.get("贷方", 0)).value if columns.get("贷方") else None, f"{ws.title}!{row}贷方", blank_zero=True)
            if debit == 0 and credit == 0:
                continue
            posting_date = ws.cell(row, columns.get("日期", 0)).value if columns.get("日期") else None
            voucher_no = ws.cell(row, columns.get("凭证号", 0)).value if columns.get("凭证号") else None
            if (
                posting_date in (None, "")
                and voucher_no in (None, "")
                and re.search(r"RD\s*[0-9A-Z]+", summary, re.I) is None
            ):
                continue
            project_code = code_key(ws.cell(row, project_col).value) if project_col else fixed_project_code
            project_name = norm(ws.cell(row, project_name_col).value) if project_name_col else fixed_project_name
            if not project_code:
                parsed = re.search(r"(RD[0-9A-Z]+(?:-[0-9A-Z]+)*)\s+([^,，;；]+)", summary, re.I)
                if parsed:
                    project_code, project_name = code_key(parsed.group(1)), norm(parsed.group(2))
            records.append(LedgerRecord(
                source_path=path,
                sheet=ws.title,
                fee_code=row_fee_code,
                fee_name=row_fee_name,
                project_code=project_code,
                project_name=project_name,
                posting_date=posting_date,
                fiscal_year=ws.cell(row, columns.get("会计年度", 0)).value if columns.get("会计年度") else None,
                period=ws.cell(row, columns.get("期间", 0)).value if columns.get("期间") else None,
                voucher_word=norm(ws.cell(row, columns.get("凭证字", 0)).value) if columns.get("凭证字") else "",
                voucher_no=voucher_no,
                summary=summary,
                debit=debit,
                credit=credit,
                source_kind="ledger",
                source_row=row,
            ))
    return records, inferred_month, source_sheets


def category_sheet(record: LedgerRecord) -> str:
    return category_sheet_for_fee(record.fee_code, record.fee_name)


def category_sheet_for_fee(fee_code: str, fee_name: str) -> str:
    rule = FEE_RULES_BY_CODE.get(code_key(fee_code)) or fee_rule_for_text(fee_name)
    return rule["detail_sheet"] if rule else ""


def master_header_for_fee(fee_code: str, fee_name: str) -> str:
    rule = FEE_RULES_BY_CODE.get(code_key(fee_code)) or fee_rule_for_text(fee_name)
    return rule["master_header"] if rule else ""


def allocate_expense_records(
    records: list[LedgerRecord],
    projects: OrderedDict[str, Project],
    warnings: list[str],
    expense_name: str,
    *,
    use_displayed_rates: bool = False,
) -> tuple[dict[str, Decimal], Decimal, Decimal, Decimal, dict[str, Decimal]]:
    """按明细账内的具体项目/RD9999标签生成一张项目分配表。

    这个函数既用于运费，也用于其他明细科目；这样每一张原始明细账
    下方的分配表都遵循同一套工时分配率和尾差规则。
    """
    pool = sum((record.amount for record in records if record.project_code == "RD9999"), Decimal("0"))
    direct_records = [record for record in records if record.project_code != "RD9999"]
    direct_total = sum((record.amount for record in direct_records), Decimal("0"))
    totals = {code: Decimal("0") for code in projects}
    direct_by_source: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
    for record in direct_records:
        direct_by_source[record.project_code] += record.amount
        if record.project_code in projects:
            totals[record.project_code] += money(record.amount)
            continue
        descendants = [code for code in projects if code.startswith(record.project_code + "-")]
        descendant_weights = {code: projects[code].rate for code in descendants}
        allocated = allocate_amount(record.amount, descendants, descendant_weights)
        if allocated:
            for code, amount in allocated.items():
                totals[code] += amount
            warnings.append(f"{expense_name} {record.project_code} 为母项目，未进入全项目公摊，仅在 {len(allocated)} 个子孙项目内分配")
            continue
        raise RuleError(
            f"{expense_name}具体项目 {record.project_code} 不在已确认的项目工时/模板槽位中，"
            "不能新建项目行或按名称猜测归属"
        )
    rate_codes = [code for code, project in projects.items() if project.rate > 0]
    pool_rates = {code: projects[code].rate for code in rate_codes}
    pool_allocated = (
        allocate_by_displayed_rates(pool, rate_codes, pool_rates)
        if use_displayed_rates
        else allocate_amount(pool, rate_codes, pool_rates)
    )
    for code, amount in pool_allocated.items():
        totals[code] += amount
    result_total = sum(totals.values(), Decimal("0"))
    source_total = money(pool + direct_total)
    if result_total != source_total:
        raise RuleError(f"{expense_name}未勾稽：有效明细={source_total}，结果={result_total}")
    return totals, money(pool), money(direct_total), result_total, dict(direct_by_source)


def allocate_freight(
    records: list[LedgerRecord],
    projects: OrderedDict[str, Project],
    warnings: list[str],
) -> tuple[dict[str, Decimal], Decimal, Decimal, Decimal, dict[str, Decimal]]:
    freight = [record for record in records if category_sheet(record) == "1.运费"]
    return allocate_expense_records(
        freight,
        projects,
        warnings,
        "运费",
        use_displayed_rates=True,
    )


def style_header(ws: Any, row: int, start: int, end: int) -> None:
    fill = PatternFill("solid", fgColor="5B9BD5")
    white = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="7F7F7F")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col in range(start, end + 1):
        cell = ws.cell(row, col)
        cell.fill, cell.font, cell.border = fill, white, border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_table(ws: Any, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    thin = Side(style="thin", color="7F7F7F")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def write_hours_sheet(wb: Any, hours_result: dict[str, Any], projects: OrderedDict[str, Project]) -> None:
    ws = wb.create_sheet("RD工时加公摊")
    project_codes = [item["project_code"] for item in hours_result["projects"]]
    project_map = {item["project_code"]: item for item in hours_result["projects"]}
    headers = ["序号", "姓名", "工时汇总"] + project_codes
    ws.append(headers)
    ws.append(["", ""] + [""] + [project_map[code]["project_name"] for code in project_codes])
    for person in hours_result["people"]:
        ws.append([person.get("employee_id"), person["employee_name"], float(person["final_hours"])] + [float(person["final_by_project"][code]) for code in project_codes])
    total_row = ws.max_row + 1
    ws.append(["合计", "", float(hours_result["totals"]["final_hours"])] + [float(project_map[code]["final_hours"]) for code in project_codes])
    summary_col = len(headers) + 2
    for offset, title in enumerate(("项目编码", "项目最终工时", "工时分配率")):
        ws.cell(1, summary_col + offset, title)
    for index, code in enumerate(project_codes, start=2):
        ws.cell(index, summary_col, code)
        ws.cell(index, summary_col + 1, float(projects[code].final_hours))
        ws.cell(index, summary_col + 2, float(projects[code].rate))
        ws.cell(index, summary_col + 2).number_format = "0.00%"
    summary_total = 2 + len(project_codes)
    ws.cell(summary_total, summary_col, "合计")
    ws.cell(summary_total, summary_col + 1, float(sum((projects[code].final_hours for code in project_codes), Decimal("0"))))
    ws.cell(summary_total, summary_col + 2, float(sum((projects[code].rate for code in project_codes), Decimal("0"))))
    ws.cell(summary_total, summary_col + 2).number_format = "0.00%"
    style_header(ws, 1, 1, len(headers))
    style_header(ws, 1, summary_col, summary_col + 2)
    style_table(ws, 1, total_row, 1, len(headers))
    style_table(ws, 1, summary_total, summary_col, summary_col + 2)
    for row in range(3, total_row + 1):
        for col in range(3, len(headers) + 1):
            ws.cell(row, col).number_format = "0.00"
    ws.freeze_panes = "D3"
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 16


def write_salary_sheet(
    wb: Any,
    rows: list[dict[str, Any]],
    project_order: list[str],
    projects: OrderedDict[str, Project],
    month: int,
) -> None:
    ws = wb.create_sheet("研发人员分配率（直接人工）")
    month_label = f"{month or 1}月"
    headers = ["序号", "人员名称", month_label, f"{month_label}工资费用", "单位成本"] + project_order
    ws.append(headers)
    ws.append(["", "", "", "", ""] + [projects[code].name for code in project_order])
    for item in rows:
        salary_row, person = item["salary"], item["matched"]
        total_hours = person["final_hours"] if person else None
        unit_cost = salary_row.salary / total_hours if person and total_hours else None
        ws.append([
            salary_row.employee_id, salary_row.name, float(total_hours) if total_hours is not None else None,
            float(salary_row.salary), float(unit_cost) if unit_cost is not None else None,
        ] + [float(item["amounts"].get(code, Decimal("0"))) if item["amounts"] else None for code in project_order])
    total_row = ws.max_row + 1
    total_salary = sum((item["salary"].salary for item in rows), Decimal("0"))
    total_hours = sum((item["matched"]["final_hours"] for item in rows if item["matched"]), Decimal("0"))
    project_totals = {code: sum((item["amounts"].get(code, Decimal("0")) for item in rows), Decimal("0")) for code in project_order}
    ws.append(["合计", "", float(total_hours), float(total_salary), float(total_salary / total_hours) if total_hours else None] + [float(project_totals[code]) for code in project_order])
    style_header(ws, 1, 1, len(headers))
    style_table(ws, 1, total_row, 1, len(headers))
    for row in range(3, total_row + 1):
        for col in range(3, len(headers) + 1):
            ws.cell(row, col).number_format = "#,##0.00"
    ws.freeze_panes = "F3"
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 15 if col != 2 else 18


def write_expense_sheet(ws: Any, records: list[LedgerRecord], month: int, project_totals: dict[str, Decimal] | None, projects: OrderedDict[str, Project]) -> None:
    year = records[0].fiscal_year if records else ""
    year_prefix = f"{year}年" if year else ""
    title = f"{year_prefix}研发费用审核明细表-{ws.title.split('.', 1)[-1]}"
    ws.merge_cells("A1:I1")
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")
    detail_headers = ["日期", "会计年度", "期间", "凭证字", "凭证号", "摘要", "借方", "研发项目", "研发名称"]
    for col, header in enumerate(detail_headers, 1):
        ws.cell(4, col, header)
    style_header(ws, 4, 1, 9)
    row = 5
    for record in records:
        values = [record.posting_date, record.fiscal_year, record.period, record.voucher_word, record.voucher_no, record.summary, float(record.amount), record.project_code, record.project_name]
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
        ws.cell(row, 7).number_format = "#,##0.00"
        row += 1
    if row == 5:
        row += 1
    style_table(ws, 4, row - 1, 1, 9)
    summary_start = row + 1
    for col, header in enumerate(("项目编码", "科目名称", f"{month or 1}月费用合计"), 1):
        ws.cell(summary_start, col, header)
    style_header(ws, summary_start, 1, 3)
    if project_totals is not None:
        data_row = summary_start + 1
        for code, amount in project_totals.items():
            if amount == 0 and projects[code].rate == 0:
                continue
            ws.cell(data_row, 1, code)
            ws.cell(data_row, 2, projects[code].name)
            ws.cell(data_row, 3, float(amount))
            ws.cell(data_row, 3).number_format = "#,##0.00"
            data_row += 1
        ws.cell(data_row, 1, "合计")
        ws.cell(data_row, 3, float(sum(project_totals.values(), Decimal("0"))))
        ws.cell(data_row, 3).number_format = "#,##0.00"
        style_table(ws, summary_start, data_row, 1, 3)
    ws.freeze_panes = "A5"
    widths = [13, 12, 10, 10, 12, 52, 16, 18, 34]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width


def write_main_sheet(
    wb: Any,
    projects: OrderedDict[str, Project],
    labor: dict[str, Decimal],
    freight: dict[str, Decimal],
) -> None:
    ws = wb.create_sheet("研发项目主表")
    ws.append(MAIN_HEADERS)
    style_header(ws, 1, 1, len(MAIN_HEADERS))
    for project in projects.values():
        values: list[Any] = [project.code, project.name, None, float(labor.get(project.code, Decimal("0"))), float(freight.get(project.code, Decimal("0")))]
        values.extend([None] * (len(MAIN_HEADERS) - 6))
        values.append(float(labor.get(project.code, Decimal("0")) + freight.get(project.code, Decimal("0"))))
        ws.append(values)
    style_table(ws, 1, ws.max_row, 1, len(MAIN_HEADERS))
    for row in range(2, ws.max_row + 1):
        for col in range(3, len(MAIN_HEADERS) + 1):
            ws.cell(row, col).number_format = "#,##0.00"
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:Y{ws.max_row}"
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 38
    for col in range(3, 26):
        ws.column_dimensions[get_column_letter(col)].width = 16


def write_combo_summary(wb: Any, records: list[LedgerRecord], month: int) -> None:
    ws = wb.create_sheet(f"项目分类明细汇总表{month or 1}月")
    ws.append(["费用项目", "研发项目", "费用金额(元)", "来源工作表"])
    style_header(ws, 1, 1, 4)
    grouped: OrderedDict[tuple[str, str, str], Decimal] = OrderedDict()
    sources: dict[tuple[str, str, str], set[str]] = defaultdict(set)
    for record in records:
        key = (f"{record.fee_code} {record.fee_name}".strip(), f"{record.project_code} {record.project_name}".strip(), category_sheet(record))
        grouped[key] = grouped.get(key, Decimal("0")) + record.amount
        sources[key].add(record.sheet)
    for key, amount in grouped.items():
        ws.append([key[0], key[1], float(amount), "；".join(sorted(sources[key]))])
    if ws.max_row == 1:
        ws.append(["", "", 0.0, ""])
    style_table(ws, 1, ws.max_row, 1, 4)
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 3).number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 38


def write_detail_sheet(wb: Any, projects: OrderedDict[str, Project], labor: dict[str, Decimal], freight: dict[str, Decimal]) -> None:
    ws = wb.create_sheet("研发明细表")
    ws.append(["项目编码", "项目名称", "费用类别", "金额"])
    style_header(ws, 1, 1, 4)
    for code, project in projects.items():
        if labor.get(code, Decimal("0")) != 0:
            ws.append([code, project.name, "直接人工（工资）", float(labor[code])])
        if freight.get(code, Decimal("0")) != 0:
            ws.append([code, project.name, "运费", float(freight[code])])
    if ws.max_row == 1:
        ws.append(["", "", "", 0.0])
    style_table(ws, 1, ws.max_row, 1, 4)
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 4).number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["C"].width = 24
    ws.column_dimensions["D"].width = 18


def clear_sheet_values(ws: Any) -> None:
    """只清除模板单元格的值和公式，保留原有格式、合并单元格与列宽。"""
    for row in ws.iter_rows():
        for cell in row:
            if not isinstance(cell, MergedCell):
                cell.value = None


def copy_row_style(ws: Any, source_row: int, target_row: int, max_col: int) -> None:
    """模板预留行不足时，沿用同一表格内的数据行格式。"""
    if target_row <= source_row:
        return
    for col in range(1, max_col + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)
        if source.has_style:
            target._style = copy(source._style)
        if source.alignment:
            target.alignment = copy(source.alignment)
        if source.number_format:
            target.number_format = source.number_format
    if ws.row_dimensions[source_row].height:
        ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def prepare_rows(ws: Any, first_data_row: int, data_count: int, template_row: int, max_col: int) -> None:
    for row in range(first_data_row, first_data_row + data_count):
        if row > ws.max_row or row > template_row:
            copy_row_style(ws, template_row, row, max_col)


def write_values(ws: Any, row: int, values: list[Any], start_col: int = 1) -> None:
    for offset, value in enumerate(values):
        cell = ws.cell(row, start_col + offset)
        if not isinstance(cell, MergedCell):
            cell.value = value


def safe_sheet(wb: Any, name: str) -> Any:
    if name not in wb.sheetnames:
        raise RuleError(f"输出模板缺少必需工作表：{name}")
    return wb[name]


def output_year(records: list[LedgerRecord]) -> str:
    for record in records:
        value = norm(record.fiscal_year)
        if value:
            return value
        if isinstance(record.posting_date, (date, datetime)):
            return str(record.posting_date.year)
    return str(date.today().year)


def write_template_main_sheet(
    wb: Any,
    projects: OrderedDict[str, Project],
    master_costs: dict[str, dict[str, Decimal]],
) -> None:
    ws = safe_sheet(wb, "研发项目主表")
    header_columns = {norm(ws.cell(1, col).value): col for col in range(1, ws.max_column + 1)}
    for header in MAIN_HEADERS:
        if norm(header) not in header_columns:
            raise RuleError(f"固定模板主表缺少费用列：{header}")
    slot_rows = list(range(3, 21))
    if len(projects) > len(slot_rows):
        raise RuleError(f"研发项目共{len(projects)}个，超出固定模板的18个项目预留槽位")
    for row in range(3, 21):
        for col in range(1, 26):
            cell = ws.cell(row, col)
            if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                cell.value = None

    for row, project in zip(slot_rows, projects.values()):
        values_by_header = {norm(header): costs.get(project.code, Decimal("0")) for header, costs in master_costs.items()}
        ws.cell(row, 1).value = project.code
        ws.cell(row, 2).value = project.name
        total = Decimal("0")
        for header in MAIN_HEADERS[2:-1]:
            value = money(values_by_header.get(norm(header), Decimal("0")))
            cell = ws.cell(row, header_columns[norm(header)])
            cell.value = float(value)
            cell.number_format = "#,##0.00"
            total += value
        ws.cell(row, header_columns[norm("项目成本合计")]).value = float(money(total))
        ws.cell(row, header_columns[norm("项目成本合计")]).number_format = "#,##0.00"


def copy_summary_to_template(
    wb: Any,
    summary_source: tuple[Path, str, str],
    records: list[LedgerRecord],
    month: int,
    warnings: list[str] | None = None,
) -> str:
    old_name = TEMPLATE_MONTH_SHEET
    new_name = f"项目分类明细汇总表{month or 1}月"
    ws = safe_sheet(wb, old_name)
    if new_name != old_name:
        ws.title = new_name
        # 费用页的预置公式引用月份汇总页；仅随允许的月份页更名同步更新引用。
        for formula_ws in wb.worksheets:
            for row in formula_ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("=") and old_name in cell.value:
                        cell.value = cell.value.replace(old_name, new_name)
    source_path, source_sheet, source_kind = summary_source
    if source_kind == "generated":
        source_rows = [
            [
                f"{record.fee_code.replace('-', '_')}  {record.fee_name}",
                f"{record.project_code}  {record.project_name}".strip(),
                record.amount,
            ]
            for record in records
        ]
        if len(source_rows) > ws.max_row - 1:
            raise RuleError(
                f"费用明细账汇总后有{len(source_rows)}行，"
                f"超出固定模板项目分类明细汇总表的{ws.max_row - 1}行预留区域"
            )
        for row in range(1, ws.max_row + 1):
            for col in range(1, 5):
                cell = ws.cell(row, col)
                if not isinstance(cell, MergedCell):
                    cell.value = None
        write_values(ws, 1, ["序号", "费用项目", "研发项目", COMBINATION_CLEANUP_RULES["amount_header"]])
        for sequence, (fee_value, project_value, amount_value) in enumerate(source_rows, start=1):
            target_row = sequence + 1
            write_values(ws, target_row, [sequence, fee_value, project_value, amount_value])
            ws.cell(target_row, 4).number_format = "#,##0.00"
        return new_name

    if source_kind == "combination":
        source_book = load_workbook(source_path, data_only=True, read_only=False)
        source = source_book[source_sheet]
        located = locate_combination_summary_sheet(source_path)
        if located is None:
            raise RuleError("核算维度与科目组合表在输出时未找到可整理的表头")
        header_row = located[2]
        headers = cell_texts(source, header_row, 100)
        fee_col = headers[COMBINATION_CLEANUP_RULES["source_headers"][0]]
        project_col = headers[COMBINATION_CLEANUP_RULES["source_headers"][1]]
        amount_headers = {
            *COMBINATION_CLEANUP_RULES["primary_amount_headers"],
            norm(COMBINATION_CLEANUP_RULES["amount_header"]),
        }
        amount_cols = [
            col for label, col in headers.items()
            if label in amount_headers and col > project_col
        ]
        if not amount_cols:
            raise RuleError("核算维度与科目组合表在输出时未找到首个借方发生额")
        amount_col = min(amount_cols)
        last_content_row = max(
            (
                row for row in range(header_row + 1, source.max_row + 1)
                if any(source.cell(row, col).value not in (None, "") for col in (fee_col, project_col, amount_col))
            ),
            default=header_row,
        )
        source_rows = [
            [
                source.cell(row, fee_col).value,
                source.cell(row, project_col).value,
                source.cell(row, amount_col).value,
            ]
            for row in range(header_row + 1, last_content_row + 1)
        ]
        capacity = ws.max_row - 1
        if len(source_rows) > capacity:
            filtered_rows: list[list[Any]] = []
            skipped_fees: list[str] = []
            current_fee_supported = False
            for fee_value, project_value, amount_value in source_rows:
                fee_text = norm(fee_value)
                if fee_text in {"总计", "合计"}:
                    continue
                if fee_text and not fee_text.endswith("小计"):
                    match = re.search(r"(FYXM\d+[_-]SYS)", fee_text, re.I)
                    fee_code = code_key(match.group(1)) if match else ""
                    current_fee_supported = fee_code in FEE_RULES_BY_CODE
                    if fee_code and not current_fee_supported:
                        skipped_fees.append(fee_text)
                if current_fee_supported:
                    filtered_rows.append([fee_value, project_value, amount_value])
            if len(filtered_rows) <= capacity:
                source_rows = filtered_rows
                if warnings is not None and skipped_fees:
                    labels = "、".join(dict.fromkeys(skipped_fees))
                    warnings.append(
                        "核算维度与科目组合表包含未进入固定费用页的分类，"
                        f"已从项目分类明细汇总表中剔除：{labels}；直接人工仍按工资表计算"
                    )
        if len(source_rows) > capacity:
            allocatable_rows = [
                [
                    f"{record.fee_code.replace('-', '_')}  {record.fee_name}",
                    f"{record.project_code}  {record.project_name}".strip(),
                    record.amount,
                ]
                for record in records
                if record.fee_code in FEE_RULES_BY_CODE
            ]
            if len(allocatable_rows) <= capacity:
                source_rows = allocatable_rows
                if warnings is not None:
                    warnings.append("项目分类明细汇总表预留区不足，已改按可测算费用明细行整理，不复制小计行")
        if len(source_rows) > capacity:
            raise RuleError(
                f"核算维度与科目组合整理后有{len(source_rows)}行，"
                f"超出固定模板项目分类明细汇总表的{capacity}行预留区域"
            )
        for row in range(1, ws.max_row + 1):
            for col in range(1, 5):
                cell = ws.cell(row, col)
                if not isinstance(cell, MergedCell):
                    cell.value = None
        write_values(ws, 1, ["序号", "费用项目", "研发项目", COMBINATION_CLEANUP_RULES["amount_header"]])
        sequence = 0
        for target_row, (fee_value, project_value, amount_value) in enumerate(source_rows, start=2):
            fee_text = norm(fee_value)
            if fee_text and not fee_text.endswith("小计") and fee_text not in {"五险合计", "总计", "合计"}:
                sequence += 1
                sequence_value: int | None = sequence
            else:
                sequence_value = None
            write_values(ws, target_row, [sequence_value, fee_value, project_value, amount_value])
            if isinstance(amount_value, (int, float, Decimal)):
                ws.cell(target_row, 4).number_format = "#,##0.00"
        return new_name

    source_book = load_workbook(source_path, data_only=True, read_only=False)
    source = source_book[source_sheet]
    if source.max_row > ws.max_row or source.max_column > ws.max_column:
        raise RuleError("核算维度汇总表超出固定模板预留区域，不能扩展模板布局")
    for row in range(1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            target = ws.cell(row, col)
            if isinstance(target, MergedCell) or (isinstance(target.value, str) and target.value.startswith("=")):
                continue
            target.value = None
    for row in range(1, source.max_row + 1):
        for col in range(1, source.max_column + 1):
            target = ws.cell(row, col)
            if isinstance(target, MergedCell) or (isinstance(target.value, str) and target.value.startswith("=")):
                continue
            target.value = source.cell(row, col).value
            if col == 4 and isinstance(target.value, (int, float, Decimal)):
                target.number_format = "#,##0.00"
    return new_name


def write_template_material_sheet(
    wb: Any,
    material_records: list[LedgerRecord],
    month: int,
) -> None:
    """将直接材料的核算维度原始归属写入模板已有材料台账区域。"""
    ws = safe_sheet(wb, "研发费用材料项目台账")
    data_rows = list(range(2, 19))
    if len(material_records) > len(data_rows):
        raise RuleError("直接材料项目记录超出固定模板材料台账的17个预留槽位")
    month_col = 3 + (month or 1)
    if month_col < 4 or month_col > 14:
        raise RuleError(f"固定模板材料台账不支持{month}月")
    for row in data_rows:
        for col in range(1, 15):
            cell = ws.cell(row, col)
            if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                cell.value = None
    for row, record in zip(data_rows, material_records):
        ws.cell(row, 1).value = record.project_code
        ws.cell(row, 2).value = record.project_name
        ws.cell(row, 3).value = "材料费"
        ws.cell(row, month_col).value = float(money(record.amount))
        ws.cell(row, month_col).number_format = "#,##0.00"


def write_template_material_allocation(
    wb: Any,
    projects: OrderedDict[str, Project],
    direct_totals: dict[str, Decimal],
    public_totals: dict[str, Decimal],
    combined_totals: dict[str, Decimal],
    month: int,
) -> None:
    """用引擎已核定的直接/公摊材料结果覆盖模板旧帮助公式。"""
    ws = safe_sheet(wb, DIRECT_MATERIAL_RULE["detail_sheet"])
    slots = list(range(50, 68))
    if len(projects) > len(slots):
        raise RuleError("固定模板材料分配表仅预留18个项目槽位")
    for col, title in ((1, "项目编码"), (2, "科目名称"), (3, f"{month or 1}月直接材料合计"),
                       (5, "项目编码"), (6, "科目名称"), (7, f"{month or 1}月公摊材料合计"),
                       (9, "项目编码"), (10, "科目名称"), (11, f"{month or 1}月材料合计")):
        ws.cell(48, col).value = title
    for row in slots:
        for col in range(1, 12):
            cell = ws.cell(row, col)
            if not isinstance(cell, MergedCell):
                cell.value = None
    ordered_projects = list(projects.values())
    for row, project in zip(slots, ordered_projects):
        for start_col, totals in ((1, direct_totals), (5, public_totals), (9, combined_totals)):
            write_values(
                ws,
                row,
                [project.code, project.name, float(money(totals.get(project.code, Decimal("0"))))],
                start_col=start_col,
            )
            ws.cell(row, start_col + 2).number_format = "#,##0.00"
        # D/H 是模板旧公式的帮助列；本规则不再从费用明细重复汇总材料。
        ws.cell(row, 4).value = None
        ws.cell(row, 8).value = None
    ws.cell(68, 1).value = "合计"
    ws.cell(68, 3).value = float(money(sum(direct_totals.values(), Decimal("0"))))
    ws.cell(68, 5).value = "合计"
    ws.cell(68, 7).value = float(money(sum(public_totals.values(), Decimal("0"))))
    ws.cell(68, 9).value = "合计"
    ws.cell(68, 11).value = float(money(sum(combined_totals.values(), Decimal("0"))))
    for col in (3, 7, 11):
        ws.cell(68, col).number_format = "#,##0.00"
    ws.cell(52, 4).value = None


def write_template_detail_sheet(
    wb: Any,
    year: str,
) -> None:
    ws = safe_sheet(wb, "研发明细表")
    # 该页的审核数和审定数本来就引用主表；保留正常公式，避免把模板变成
    # 无法追溯的静态报表。
    write_values(ws, 1, [f"{year}年研发费用审核明细表"])
    write_values(ws, 2, ["编制单位："])


def copy_input_values_to_template(source_ws: Any, target_ws: Any) -> None:
    clear_sheet_values(target_ws)
    for row in source_ws.iter_rows():
        for source_cell in row:
            target_cell = target_ws.cell(source_cell.row, source_cell.column)
            if not isinstance(target_cell, MergedCell):
                target_cell.value = source_cell.value


def write_template_hours_sheets(
    wb: Any,
    hours_path: Path,
    hours_layout: Any,
    hours_result: dict[str, Any],
    projects: OrderedDict[str, Project],
    month: int,
) -> None:
    project_codes = [item["project_code"] for item in hours_result["projects"]]
    if len(project_codes) > 18:
        raise RuleError("工时项目超过固定模板的18个项目预留槽位")
    project_map = {item["project_code"]: item for item in hours_result["projects"]}

    # 固定模板重构后曾漏掉“研发人员工时统计”源表的回填。
    # 只处理上方预留矩阵 A:AF36，下方项目分配率区单独重写。
    source_book = load_workbook(hours_path, data_only=True, read_only=False)
    source_ws = safe_sheet(source_book, hours_layout.sheet)
    raw_ws = safe_sheet(wb, "研发人员工时统计")
    for row in range(1, 37):
        for col in range(1, 33):
            cell = raw_ws.cell(row, col)
            if not isinstance(cell, MergedCell):
                cell.value = None
    for row in range(1, min(source_ws.max_row, 36) + 1):
        for col in range(1, min(source_ws.max_column, 30) + 1):
            target = raw_ws.cell(row, col)
            if not isinstance(target, MergedCell):
                target.value = source_ws.cell(row, col).value
    raw_ws.cell(1, 30).value = "工时汇总"
    raw_ws.cell(1, 31).value = "RD工时"
    raw_ws.cell(1, 32).value = "公摊RD工时"
    raw_ws.cell(1, 3).value = f"{month or 1}月"
    for person in hours_result["people"]:
        row = int(person["row"])
        if row > 35:
            raise RuleError("人员数超出固定模板工时源矩阵预留行")
        raw_ws.cell(row, 1).value = person.get("employee_id")
        raw_ws.cell(row, 2).value = person["employee_name"]
        raw_ws.cell(row, 3).value = float(person["final_hours"])
        raw_ws.cell(row, 30).value = float(person["final_hours"])
        raw_ws.cell(row, 31).value = float(person["actual_hours"])
        raw_ws.cell(row, 32).value = float(person["allocated_hours"])
        for col in (3, 30, 31, 32):
            raw_ws.cell(row, col).number_format = "0.00"
    raw_total_row = 36
    raw_ws.cell(raw_total_row, 1).value = "合计"
    raw_ws.cell(raw_total_row, 3).value = float(hours_result["totals"]["final_hours"])
    raw_ws.cell(raw_total_row, 30).value = float(hours_result["totals"]["final_hours"])
    raw_ws.cell(raw_total_row, 31).value = float(hours_result["totals"]["actual_hours"])
    raw_ws.cell(raw_total_row, 32).value = float(hours_result["totals"]["allocated_hours"])
    for col in (3, 30, 31, 32):
        raw_ws.cell(raw_total_row, col).number_format = "0.00"

    for row in range(45, 64):
        for col in range(2, 5):
            cell = raw_ws.cell(row, col)
            if not isinstance(cell, MergedCell):
                cell.value = None
    raw_ws.cell(43, 2).value = "项目编码"
    raw_ws.cell(43, 3).value = "科目名称"
    raw_ws.cell(43, 4).value = "工时分配率"
    for row, project in zip(range(45, 63), projects.values()):
        write_values(raw_ws, row, [project.code, project.name, float(displayed_rate(project.rate))], start_col=2)
        raw_ws.cell(row, 4).number_format = "0.00%"
    raw_ws.cell(63, 2).value = "合计"
    raw_ws.cell(63, 4).value = float(
        sum((displayed_rate(project.rate) for project in projects.values()), Decimal("0"))
    )
    raw_ws.cell(63, 4).number_format = "0.00%"

    ws = safe_sheet(wb, "研发人员工时统计合计")
    headers = ["序号", "姓名", f"{month or 1}月"] + project_codes
    max_input_col = 3 + len(project_codes)
    matrix_last_row = 3 + len(hours_result["people"])
    if matrix_last_row > 42:
        raise RuleError("人员数超出固定模板工时矩阵预留行，不能扩展布局")
    # A:AH全量清理，避免R:AC残留模板历史非RD数据。
    for row in range(1, 43):
        for col in range(1, 35):
            cell = ws.cell(row, col)
            if not isinstance(cell, MergedCell):
                cell.value = None
    write_values(ws, 1, headers)
    write_values(ws, 1, ["工时汇总", "RD工时", "公摊RD工时"], start_col=30)
    write_values(ws, 2, [None, None, None] + [project_map[code]["project_name"] for code in project_codes])
    for row, person in enumerate(hours_result["people"], start=3):
        write_values(ws, row, [person.get("employee_id"), person["employee_name"], float(person["final_hours"])] + [float(person["final_by_project"][code]) for code in project_codes])
        write_values(ws, row, [float(person["final_hours"]), float(person["actual_hours"]), float(person["allocated_hours"])], start_col=30)
    total_row = 3 + len(hours_result["people"])
    write_values(ws, total_row, ["合计", None, float(hours_result["totals"]["final_hours"])] + [float(project_map[code]["final_hours"]) for code in project_codes])
    write_values(ws, total_row, [float(hours_result["totals"]["final_hours"]), float(hours_result["totals"]["actual_hours"]), float(hours_result["totals"]["allocated_hours"])], start_col=30)
    for row in range(3, total_row + 1):
        for col in list(range(3, max_input_col + 1)) + [30, 31, 32]:
            ws.cell(row, col).number_format = "0.00"

    # H:J43:60 固定预留18个项目槽位，合计行仍在第61行。
    for row in range(43, 62):
        for col in range(8, 11):
            cell = ws.cell(row, col)
            if not isinstance(cell, MergedCell):
                cell.value = None
    # 模板把 H:I:J 的 41-42 行合并为表头区；表头必须写到合并左上角 41 行，
    # 否则会因 42 行是 MergedCell 而静默丢失，导致分配率表没有表头。
    write_values(ws, 41, ["项目编码", "科目名称", "工时分配率"], start_col=8)
    for row, project in zip(range(43, 61), projects.values()):
        write_values(ws, row, [project.code, project.name, float(displayed_rate(project.rate))], start_col=8)
        ws.cell(row, 10).number_format = "0.00%"
    write_values(
        ws,
        61,
        [
            "合计",
            None,
            float(sum((displayed_rate(project.rate) for project in projects.values()), Decimal("0"))),
        ],
        start_col=8,
    )
    ws.cell(61, 10).number_format = "0.00%"


def standard_fee_allocation_table(ws: Any) -> tuple[int, list[int], int] | None:
    """定位模板中既有的“项目编码、科目名称、N月费用合计”预留表。"""
    for row in range(1, ws.max_row + 1):
        if norm(ws.cell(row, 1).value) != "项目编码" or norm(ws.cell(row, 2).value) != "科目名称":
            continue
        if "费用合计" not in norm(ws.cell(row, 3).value):
            continue
        data_rows: list[int] = []
        total_row = 0
        for candidate in range(row + 1, ws.max_row + 1):
            label = norm(ws.cell(candidate, 1).value)
            if label == "合计":
                total_row = candidate
                break
            if PROJECT_RE.fullmatch(code_key(label)):
                data_rows.append(candidate)
        if data_rows and total_row:
            return row, data_rows, total_row
    return None


def relocate_standard_fee_allocation_table(
    ws: Any,
    blank_rows: int,
    content_last_row: int | None = None,
) -> int:
    """把其他费用页的既有汇总表移到上方实际内容后。

    模板原表块的合并、样式、公式和项目槽位全部原样迁移；项目行的
    跨表公式不做相对平移，以继续引用固定的项目分配率行。
    """
    table = standard_fee_allocation_table(ws)
    if table is None:
        return 0
    header_row, data_rows, total_row = table
    last_content_row = content_last_row
    if last_content_row is None:
        last_content_row = max(
            (
                row
                for row in range(1, header_row)
                if any(
                    ws.cell(row, col).value not in (None, "")
                    for col in range(1, min(9, ws.max_column) + 1)
                )
            ),
            default=4,
        )
    target_header_row = last_content_row + max(int(blank_rows), 0) + 1
    if target_header_row == header_row:
        return header_row

    block_height = total_row - header_row + 1
    target_total_row = target_header_row + block_height - 1
    for row in range(target_header_row, target_total_row + 1):
        # 旧表头之前的 A:I 是模板历史明细账区，本次生成时会
        # 被完整清空并回填，因此不能将其误判为布局冲突。仅当
        # 新表需要扩展到原汇总表之后时，才保护下方其他内容。
        if row <= total_row:
            continue
        if any(ws.cell(row, col).value not in (None, "") for col in range(1, 4)):
            raise RuleError(f"{ws.title} 的费用汇总表新位置与模板其他内容重叠")
    snapshots: list[list[dict[str, Any]]] = []
    for row in range(header_row, total_row + 1):
        snapshot_row: list[dict[str, Any]] = []
        for col in range(1, 4):
            cell = ws.cell(row, col)
            snapshot_row.append({
                "value": cell.value,
                "style": copy(cell._style),
                "number_format": cell.number_format,
                "protection": copy(cell.protection),
                "alignment": copy(cell.alignment),
            })
        snapshots.append(snapshot_row)
    row_heights = [ws.row_dimensions[row].height for row in range(header_row, total_row + 1)]
    source_merges = []
    for merged in list(ws.merged_cells.ranges):
        if (
            merged.min_row >= header_row
            and merged.max_row <= total_row
            and merged.min_col >= 1
            and merged.max_col <= 3
        ):
            source_merges.append((merged.min_row, merged.max_row, merged.min_col, merged.max_col))
            ws.unmerge_cells(str(merged))
    # 新位置可能与模板中空白装饰性合并格重叠；表块会在完成写入后
    # 按原表块的相对合并关系重建。
    for merged in list(ws.merged_cells.ranges):
        intersects_target = not (
            merged.max_row < target_header_row
            or merged.min_row > target_total_row
            or merged.max_col < 1
            or merged.min_col > 3
        )
        if intersects_target:
            ws.unmerge_cells(str(merged))

    raw_style_row = 5 if ws.max_row >= 5 else 1
    raw_height = ws.row_dimensions[raw_style_row].height
    for row in range(header_row, total_row + 1):
        ws.row_dimensions[row].height = raw_height
        for col in range(1, 4):
            cell = ws.cell(row, col)
            cell.value = None
            cell._style = copy(ws.cell(raw_style_row, col)._style)

    for row_offset, snapshot_row in enumerate(snapshots):
        target_row = target_header_row + row_offset
        ws.row_dimensions[target_row].height = row_heights[row_offset]
        for col_offset, snapshot in enumerate(snapshot_row, start=1):
            cell = ws.cell(target_row, col_offset)
            cell.value = snapshot["value"]
            cell._style = copy(snapshot["style"])
            cell.number_format = snapshot["number_format"]
            cell.protection = copy(snapshot["protection"])
            cell.alignment = copy(snapshot["alignment"])
    row_shift = target_header_row - header_row
    for min_row, max_row, min_col, max_col in source_merges:
        ws.merge_cells(
            start_row=min_row + row_shift,
            end_row=max_row + row_shift,
            start_column=min_col,
            end_column=max_col,
        )

    # 只重建块内合计公式；项目行的跨表公式必须保持原引用。
    data_start_offset = data_rows[0] - header_row
    data_end_offset = data_rows[-1] - header_row
    new_total_row = target_header_row + block_height - 1
    ws.cell(new_total_row, 3).value = (
        f"=SUM(C{target_header_row + data_start_offset}:C{target_header_row + data_end_offset})"
    )
    return target_header_row


def relocate_material_allocation_tables(
    ws: Any,
    blank_rows: int = 3,
    content_last_row: int | None = None,
) -> int:
    """Move the three parallel material allocation blocks as one template section."""
    header_row = 0
    total_row = 0
    for candidate in range(1, ws.max_row + 1):
        if (
            norm(ws.cell(candidate, 1).value) == "项目编码"
            and norm(ws.cell(candidate, 2).value) == "科目名称"
            and "材料合计" in norm(ws.cell(candidate, 3).value)
        ):
            header_row = candidate
            break
    if not header_row:
        return 0
    for candidate in range(header_row + 1, ws.max_row + 1):
        if norm(ws.cell(candidate, 1).value) == "合计":
            total_row = candidate
            break
    if not total_row:
        return 0
    last_content_row = content_last_row
    if last_content_row is None:
        last_content_row = max(
            (
                row
                for row in range(1, header_row)
                if any(
                    ws.cell(row, col).value not in (None, "")
                    and not (isinstance(ws.cell(row, col).value, str) and ws.cell(row, col).value.startswith("="))
                    for col in range(1, 12)
                )
            ),
            default=4,
        )
    target_header_row = last_content_row + max(int(blank_rows), 0) + 1
    if target_header_row == header_row:
        return header_row

    side_snapshots: dict[int, list[list[dict[str, Any]]]] = {}
    side_merges: dict[int, list[tuple[int, int, int, int]]] = {}
    for start_col in (5, 9):
        side_snapshots[start_col] = []
        for row in range(header_row, total_row + 1):
            snapshot_row: list[dict[str, Any]] = []
            for col in range(start_col, start_col + 3):
                cell = ws.cell(row, col)
                snapshot_row.append({
                    "value": cell.value,
                    "style": copy(cell._style),
                    "number_format": cell.number_format,
                    "protection": copy(cell.protection),
                    "alignment": copy(cell.alignment),
                })
            side_snapshots[start_col].append(snapshot_row)
        side_merges[start_col] = []
        for merged in list(ws.merged_cells.ranges):
            if (
                merged.min_row >= header_row
                and merged.max_row <= total_row
                and merged.min_col >= start_col
                and merged.max_col <= start_col + 2
            ):
                side_merges[start_col].append(
                    (merged.min_row, merged.max_row, merged.min_col, merged.max_col)
                )
                ws.unmerge_cells(str(merged))

    # A:C follows the same 21-row visual structure as a standard allocation table,
    # but its third header says “材料合计”, so move it here explicitly.
    primary_snapshots: list[list[dict[str, Any]]] = []
    primary_merges: list[tuple[int, int, int, int]] = []
    for row in range(header_row, total_row + 1):
        snapshot_row: list[dict[str, Any]] = []
        for col in range(1, 4):
            cell = ws.cell(row, col)
            snapshot_row.append({
                "value": cell.value,
                "style": copy(cell._style),
                "number_format": cell.number_format,
                "protection": copy(cell.protection),
                "alignment": copy(cell.alignment),
            })
        primary_snapshots.append(snapshot_row)
    for merged in list(ws.merged_cells.ranges):
        if (
            merged.min_row >= header_row
            and merged.max_row <= total_row
            and merged.min_col >= 1
            and merged.max_col <= 3
        ):
            primary_merges.append((merged.min_row, merged.max_row, merged.min_col, merged.max_col))
            ws.unmerge_cells(str(merged))

    raw_style_row = 5 if ws.max_row >= 5 else 1
    raw_height = ws.row_dimensions[raw_style_row].height
    block_height = total_row - header_row + 1
    for merged in list(ws.merged_cells.ranges):
        intersects_target = not (
            merged.max_row < target_header_row
            or merged.min_row > target_header_row + block_height - 1
            or merged.max_col < 1
            or merged.min_col > 3
        )
        if intersects_target:
            ws.unmerge_cells(str(merged))
    for row in range(header_row, total_row + 1):
        ws.row_dimensions[row].height = raw_height
        for col in range(1, 4):
            cell = ws.cell(row, col)
            cell.value = None
            cell._style = copy(ws.cell(raw_style_row, col)._style)
    for row_offset, snapshot_row in enumerate(primary_snapshots):
        target_row = target_header_row + row_offset
        for col_offset, snapshot in enumerate(snapshot_row, start=1):
            cell = ws.cell(target_row, col_offset)
            cell.value = snapshot["value"]
            cell._style = copy(snapshot["style"])
            cell.number_format = snapshot["number_format"]
            cell.protection = copy(snapshot["protection"])
            cell.alignment = copy(snapshot["alignment"])
    row_shift = target_header_row - header_row
    for min_row, max_row, min_col, max_col in primary_merges:
        ws.merge_cells(
            start_row=min_row + row_shift,
            end_row=max_row + row_shift,
            start_column=min_col,
            end_column=max_col,
        )
    ws.cell(target_header_row + block_height - 1, 3).value = (
        f"=SUM(C{target_header_row + 2}:C{target_header_row + 19})"
    )
    for start_col in (5, 9):
        for merged in list(ws.merged_cells.ranges):
            intersects_target = not (
                merged.max_row < target_header_row
                or merged.min_row > target_header_row + block_height - 1
                or merged.max_col < start_col
                or merged.min_col > start_col + 2
            )
            if intersects_target:
                ws.unmerge_cells(str(merged))
        for row in range(header_row, total_row + 1):
            ws.row_dimensions[row].height = raw_height
            for col in range(start_col, start_col + 3):
                cell = ws.cell(row, col)
                cell.value = None
                cell._style = copy(ws.cell(raw_style_row, col)._style)
        for row_offset, snapshot_row in enumerate(side_snapshots[start_col]):
            target_row = target_header_row + row_offset
            for col_offset, snapshot in enumerate(snapshot_row, start=start_col):
                cell = ws.cell(target_row, col_offset)
                cell.value = snapshot["value"]
                cell._style = copy(snapshot["style"])
                cell.number_format = snapshot["number_format"]
                cell.protection = copy(snapshot["protection"])
                cell.alignment = copy(snapshot["alignment"])
        row_shift = target_header_row - header_row
        for min_row, max_row, min_col, max_col in side_merges[start_col]:
            ws.merge_cells(
                start_row=min_row + row_shift,
                end_row=max_row + row_shift,
                start_column=min_col,
                end_column=max_col,
            )
        ws.cell(target_header_row + block_height - 1, start_col + 2).value = (
            f"=SUM({get_column_letter(start_col + 2)}{target_header_row + 2}:"
            f"{get_column_letter(start_col + 2)}{target_header_row + 19})"
        )
    return target_header_row


def write_static_fee_allocation_table(
    ws: Any,
    fee_name: str,
    month: int,
    projects: OrderedDict[str, Project],
    project_totals: dict[str, Decimal],
) -> bool:
    """向模板已预留但没有公式的费用分配表写入本次计算结果。

    已有正确公式的费用页保持公式不变；没有公式的费用页才回填既有的金额
    单元格。这样既覆盖所有“其他”科目，也不改变固定模板的公式结构。
    """
    table = standard_fee_allocation_table(ws)
    if table is None:
        return False
    header_row, data_rows, total_row = table
    ordered_projects = list(projects.values())
    if len(ordered_projects) > len(data_rows):
        raise RuleError(f"固定模板的{ws.title}仅预留{len(data_rows)}个项目槽位")
    ws.cell(header_row, 3).value = f"{month or 1}月费用合计"
    for index, row in enumerate(data_rows):
        project = ordered_projects[index] if index < len(ordered_projects) else None
        ws.cell(row, 1).value = project.code if project is not None else None
        ws.cell(row, 2).value = project.name if project is not None else None
    for index, row in enumerate(data_rows):
        project = ordered_projects[index] if index < len(ordered_projects) else None
        amount = money(project_totals.get(project.code, Decimal("0"))) if project is not None else Decimal("0")
        cell = ws.cell(row, 3)
        cell.value = float(amount)
        cell.number_format = "#,##0.00"
    ws.cell(total_row, 3).number_format = "#,##0.00"
    ws.cell(total_row, 3).value = float(money(sum(project_totals.values(), Decimal("0"))))
    return True


def write_freight_allocation_table(
    ws: Any,
    projects: OrderedDict[str, Project],
    project_totals: dict[str, Decimal],
) -> bool:
    """回填运费页原有的“项目编号/工时/分配率/费用”预留表。"""
    header_row = 0
    start_col = 0
    for row in range(1, ws.max_row + 1):
        for col in range(1, ws.max_column - 3):
            labels = [norm(ws.cell(row, col + offset).value) for offset in range(4)]
            if labels == ["项目编号", "工时合计", "工时分配率", "费用"]:
                header_row, start_col = row, col
                break
        if header_row:
            break
    if not header_row:
        return False
    formula_col = start_col + 4
    rows = [
        row for row in range(header_row + 1, ws.max_row + 1)
        if isinstance(ws.cell(row, formula_col).value, str) and ws.cell(row, formula_col).value.startswith("=")
    ]
    if not rows:
        return False
    total = money(sum(project_totals.values(), Decimal("0")))
    ordered_projects = list(projects.values())
    for index, row in enumerate(rows):
        project = ordered_projects[index] if index < len(ordered_projects) else None
        for col in range(start_col, start_col + 4):
            ws.cell(row, col).value = None
        if project is not None:
            write_values(
                ws,
                row,
                [project.code, float(project.final_hours), float(displayed_rate(project.rate)), float(total)],
                start_col=start_col,
            )
            ws.cell(row, start_col + 1).number_format = "0.00"
            ws.cell(row, start_col + 2).number_format = "0.00%"
            ws.cell(row, start_col + 3).number_format = "#,##0.00"
        ws.cell(row, formula_col).number_format = "#,##0.00"
    return True


def fee_source_content_last_row(records: list[LedgerRecord]) -> int:
    """明细从第5行开始；有明细时在末尾再写1行合计。"""
    return 4 if not records else 5 + len(records)


def relocate_freight_visible_table(ws: Any, target_header_row: int) -> int:
    """迁移运费页已预留的可见 F82:H102 结果表块。

    运费也使用“上方最后一条内容 + 3个空白行”的定位规则；
    旧的隐藏表 F60:J81 只是模板历史遗留，不再作为生成目标。
    """
    source_header_row, source_total_row = 82, 102
    if ws.max_row < source_total_row or ws.max_column < 8:
        raise RuleError("固定模板的1.运费缺少可见结果表样式区 F82:H102")
    block_height = source_total_row - source_header_row + 1
    target_total_row = target_header_row + block_height - 1
    snapshots: list[list[dict[str, Any]]] = []
    for row in range(source_header_row, source_total_row + 1):
        snapshot_row: list[dict[str, Any]] = []
        for col in range(6, 9):
            cell = ws.cell(row, col)
            snapshot_row.append({
                "value": cell.value,
                "style": copy(cell._style),
                "number_format": cell.number_format,
                "protection": copy(cell.protection),
                "alignment": copy(cell.alignment),
            })
        snapshots.append(snapshot_row)
    row_heights = [ws.row_dimensions[row].height for row in range(source_header_row, source_total_row + 1)]
    source_merges: list[tuple[int, int, int, int]] = []
    for merged in list(ws.merged_cells.ranges):
        if (
            merged.min_row >= source_header_row
            and merged.max_row <= source_total_row
            and merged.min_col >= 6
            and merged.max_col <= 8
        ):
            source_merges.append((merged.min_row, merged.max_row, merged.min_col, merged.max_col))
            ws.unmerge_cells(str(merged))
    # 解除新目标区和旧隐藏区中的历史合并，后面按可见表块重建。
    for merged in list(ws.merged_cells.ranges):
        intersects_target = not (
            merged.max_row < target_header_row
            or merged.min_row > target_total_row
            or merged.max_col < 6
            or merged.min_col > 8
        )
        inside_legacy = (
            merged.min_row >= 60 and merged.max_row <= 81
            and merged.min_col >= 6 and merged.max_col <= 10
        )
        if intersects_target or inside_legacy:
            ws.unmerge_cells(str(merged))

    raw_style_row = 5
    raw_height = ws.row_dimensions[raw_style_row].height
    for row in range(source_header_row, source_total_row + 1):
        ws.row_dimensions[row].height = raw_height
        for col in range(6, 9):
            cell = ws.cell(row, col)
            cell.value = None
            cell._style = copy(ws.cell(raw_style_row, col)._style)
    for row in range(60, 82):
        for col in range(6, 11):
            cell = ws.cell(row, col)
            if not isinstance(cell, MergedCell):
                cell.value = None

    for row_offset, snapshot_row in enumerate(snapshots):
        target_row = target_header_row + row_offset
        ws.row_dimensions[target_row].height = row_heights[row_offset]
        ws.row_dimensions[target_row].hidden = False
        for col_offset, snapshot in enumerate(snapshot_row, start=6):
            cell = ws.cell(target_row, col_offset)
            cell.value = snapshot["value"]
            cell._style = copy(snapshot["style"])
            cell.number_format = snapshot["number_format"]
            cell.protection = copy(snapshot["protection"])
            cell.alignment = copy(snapshot["alignment"])
    row_shift = target_header_row - source_header_row
    for min_row, max_row, min_col, max_col in source_merges:
        ws.merge_cells(
            start_row=min_row + row_shift,
            end_row=max_row + row_shift,
            start_column=min_col,
            end_column=max_col,
        )
    return target_header_row


def write_visible_freight_allocation_table(
    ws: Any,
    month: int,
    projects: OrderedDict[str, Project],
    project_totals: dict[str, Decimal],
    header_row: int,
) -> bool:
    """把运费直接写入已迁移到明细下方的可见结果表。

    旧路径写到被隐藏的60:81行，且“总额×分配率”会重新产生0.01元
    尾差。这里直接使用已完成四舍五入及尾差守恒的项目金额。
    """
    if not header_row or ws.max_column < 8:
        return False
    first_data_row, last_data_row, total_row = header_row + 2, header_row + 19, header_row + 20
    slots = list(range(first_data_row, last_data_row + 1))
    if len(projects) > len(slots):
        raise RuleError("固定模板的1.运费可见结果区仅预留18个项目槽位")
    ws.cell(header_row, 6).value = "项目编码"
    ws.cell(header_row, 7).value = "科目名称"
    ws.cell(header_row, 8).value = f"{month or 1}月费用合计"
    for row in slots:
        for col in range(6, 9):
            cell = ws.cell(row, col)
            if not isinstance(cell, MergedCell):
                cell.value = None
    for row, project in zip(slots, projects.values()):
        amount = money(project_totals.get(project.code, Decimal("0")))
        write_values(ws, row, [project.code, project.name, float(amount)], start_col=6)
        ws.cell(row, 8).number_format = "#,##0.00"
    ws.cell(total_row, 6).value = "合计"
    ws.cell(total_row, 8).value = float(money(sum(project_totals.values(), Decimal("0"))))
    ws.cell(total_row, 8).number_format = "#,##0.00"
    return True


def write_template_other_fee_allocations(
    wb: Any,
    projects: OrderedDict[str, Project],
    master_costs: dict[str, dict[str, Decimal]],
    month: int,
    allocation_header_rows: dict[str, int],
) -> None:
    """用核算维度汇总结果回填所有20类其他费用页的既有分配表。"""
    for rule in COST_RULES["fee_items"]:
        ws = safe_sheet(wb, rule["detail_sheet"])
        totals = master_costs[norm(rule["master_header"])]
        expected_header_row = allocation_header_rows.get(rule["code"], 0)
        if rule["code"] == code_key("FYXM02_SYS"):
            written = write_visible_freight_allocation_table(
                ws, month, projects, totals, expected_header_row
            )
        else:
            table = standard_fee_allocation_table(ws)
            if table is None or table[0] != expected_header_row:
                raise RuleError(f"{ws.title} 的费用汇总表未保持在固定模板预留位置")
            written = write_static_fee_allocation_table(ws, rule["name"], month, projects, totals)
        if not written:
            raise RuleError(f"固定模板的{ws.title}缺少费用分配结果预留区域")


def project_amounts_from_records(records: list[LedgerRecord]) -> dict[str, Decimal]:
    result: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
    for record in records:
        result[record.project_code] += record.amount
    return dict(result)


def write_template_expense_sheet(
    ws: Any,
    records: list[LedgerRecord],
    month: int,
    projects: OrderedDict[str, Project],
    project_totals: dict[str, Decimal] | None,
) -> None:
    clear_sheet_values(ws)
    label = ws.title.split(".", 1)[-1]
    year = output_year(records)
    write_values(ws, 1, [f"{year}年研发费用审核明细表-{label}"])
    write_values(ws, 3, ["单位：元", None, None, "费用类别：", None, float(money(sum((record.amount for record in records), Decimal("0"))))])
    ws.cell(3, 6).number_format = "#,##0.00"
    write_values(ws, 4, ["日期", "会计年度", "期间", "凭证字", "凭证号", "摘要", "借方", "研发项目", "研发名称"])
    prepare_rows(ws, 5, len(records), 5, 9)
    for row, record in enumerate(records, start=5):
        write_values(ws, row, [record.posting_date, record.fiscal_year, record.period, record.voucher_word, record.voucher_no, record.summary, float(money(record.amount)), record.project_code, record.project_name])
        ws.cell(row, 7).number_format = "#,##0.00"
    total_row = 5 + len(records)
    copy_row_style(ws, 5, total_row, 9)
    source_total = money(sum((record.amount for record in records), Decimal("0")))
    write_values(ws, total_row, ["合计", None, None, None, None, None, float(source_total)])
    ws.cell(total_row, 7).number_format = "#,##0.00"

    summary_items = project_totals if project_totals is not None else project_amounts_from_records(records)
    summary_start = total_row + 3
    copy_row_style(ws, min(9, ws.max_row), summary_start, 3)
    write_values(ws, summary_start, ["项目编码", "科目名称", f"{month or 1}月费用合计"])
    ordered_codes = list(projects) if project_totals is not None else list(summary_items)
    prepare_rows(ws, summary_start + 1, len(ordered_codes), min(9, ws.max_row), 3)
    for row, code in enumerate(ordered_codes, start=summary_start + 1):
        project = projects.get(code, Project(code, code))
        write_values(ws, row, [code, project.name, float(money(summary_items.get(code, Decimal("0"))))])
        ws.cell(row, 3).number_format = "#,##0.00"
    summary_total_row = summary_start + len(ordered_codes) + 1
    copy_row_style(ws, min(9, ws.max_row), summary_total_row, 3)
    write_values(ws, summary_total_row, ["合计", None, float(money(sum(summary_items.values(), Decimal("0"))))])
    ws.cell(summary_total_row, 3).number_format = "#,##0.00"


def clear_broken_formula_cells(wb: Any, warnings: list[str]) -> None:
    """模板已存在的不可审计错误仅清内容、保留单元格格式和布局。"""
    tokens = ("#REF!", "#NAME?", "#N/A", "#VALUE!", "#DIV/0!")
    cleared: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and any(token in value.upper() for token in tokens):
                    cell.value = None
                    cleared.append(f"{ws.title}!{cell.coordinate}")
    if cleared:
        warnings.append(f"已清除固定模板中{len(cleared)}个既有公式错误内容，保留原有格式和布局")


def quote_internal_sheet_references(wb: Any) -> None:
    """统一为跨表公式的工作表名加单引号。

    Excel 对部分中文工作表名能容忍无引号写法，但其他标准解析器会将
    它们误当作名称并计算为 ``#NAME?``。单引号写法是兼容 Excel 的标准
    内部引用，不改变公式含义。
    """
    replacements = [
        (name, name.replace("'", "''"))
        for name in sorted(wb.sheetnames, key=len, reverse=True)
    ]
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if not (isinstance(value, str) and value.startswith("=")):
                    continue
                for name, escaped in replacements:
                    value = value.replace(f"{name}!", f"'{escaped}'!")
                cell.value = value


def clear_non_formula_cells(ws: Any, start_row: int, end_row: int, start_col: int, end_col: int) -> None:
    for row in range(start_row, end_row + 1):
        for col in range(start_col, end_col + 1):
            cell = ws.cell(row, col)
            if isinstance(cell, MergedCell) or (isinstance(cell.value, str) and cell.value.startswith("=")):
                continue
            cell.value = None


def write_template_expense_source(
    ws: Any,
    records: list[LedgerRecord],
    month: int,
    allocation_start_row: int | None = None,
) -> None:
    """把上传的费用明细账写进该费用页的既有原始账区。

    分配表由下方的固定模板预留区域承载，不能因为清空原始账区而被覆盖。
    """
    protected_cells: set[tuple[int, int]] = set()
    if allocation_start_row is None:
        allocation_starts: list[int] = []
        standard_table = standard_fee_allocation_table(ws)
        if standard_table is not None:
            allocation_starts.append(standard_table[0])
        formula_rows = [
            cell.row for row in ws.iter_rows(min_row=5) for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
        boundaries = allocation_starts + formula_rows
        raw_end = min(boundaries) - 1 if boundaries else ws.max_row
    else:
        raw_end = allocation_start_row - 1
        if ws.title == "1.运费":
            # 运费使用 F:H 的21行可见分配表。
            protected_cells.update(
                (row, col)
                for row in range(allocation_start_row, allocation_start_row + 21)
                for col in range(6, 9)
            )
        else:
            table = standard_fee_allocation_table(ws)
            if table is None or table[0] != allocation_start_row:
                raise RuleError(f"固定模板的{ws.title}未找到当前费用分配表")
            protected_cells.update(
                (row, col)
                for row in range(table[0], table[2] + 1)
                for col in range(1, 4)
            )
    if raw_end < 5:
        return
    required_last_row = fee_source_content_last_row(records)
    if required_last_row > raw_end - 3:
        raise RuleError(f"{ws.title} 的项目汇总表未保留3个空白行")
    prepare_rows(ws, 5, len(records) + (1 if records else 0), 5, min(9, ws.max_column))
    # 这一区只承载本次上传的明细；模板历史值和历史公式都必须清除。
    # 扫描整个 A:I 历史账区，同时精确保护固定位置的分配表。
    clear_through = ws.max_row if allocation_start_row is not None else raw_end
    for row in range(5, clear_through + 1):
        for col in range(1, min(9, ws.max_column) + 1):
            if (row, col) in protected_cells:
                continue
            cell = ws.cell(row, col)
            if not isinstance(cell, MergedCell):
                cell.value = None
    for row, record in enumerate(records, start=5):
        write_values(ws, row, [
            record.posting_date, record.fiscal_year, record.period, record.voucher_word,
            record.voucher_no, record.summary, float(money(record.amount)),
            record.project_code, record.project_name,
        ])
        ws.cell(row, 7).number_format = "#,##0.00"
    if records:
        total_row = 5 + len(records)
        write_values(
            ws,
            total_row,
            ["合计", None, None, None, None, None, float(money(sum((record.amount for record in records), Decimal("0"))))],
        )
        ws.cell(total_row, 7).number_format = "#,##0.00"


def write_template_labor_sheet(wb: Any, projects: OrderedDict[str, Project], labor_totals: dict[str, Decimal], month: int) -> None:
    ws = safe_sheet(wb, "2.直接人工")
    clear_non_formula_cells(ws, 5, 5, 1, min(9, ws.max_column))
    ws.cell(5, 7).value = float(money(sum(labor_totals.values(), Decimal("0"))))
    ws.cell(5, 7).number_format = "#,##0.00"
    slot_rows = list(range(9, 27))
    if len(projects) > len(slot_rows):
        raise RuleError("固定模板直接人工页仅预留18个项目槽位")
    for row in range(9, 27):
        for col in range(1, 4):
            if not (isinstance(ws.cell(row, col).value, str) and ws.cell(row, col).value.startswith("=")):
                ws.cell(row, col).value = None
    for row, project in zip(slot_rows, projects.values()):
        ws.cell(row, 1).value = project.code
        ws.cell(row, 2).value = project.name
        ws.cell(row, 3).value = float(money(labor_totals.get(project.code, Decimal("0"))))
        ws.cell(row, 3).number_format = "#,##0.00"
    ws.cell(27, 3).number_format = "#,##0.00"


def assert_formula_errors_cleared(wb: Any) -> None:
    tokens = ("#REF!", "#NAME?", "#N/A", "#VALUE!", "#DIV/0!")
    errors = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and any(token in cell.value.upper() for token in tokens):
                    errors.append(f"{ws.title}!{cell.coordinate}")
    if errors:
        raise RuleError("输出模板仍存在公式错误：" + "、".join(errors[:10]))


def sheet_layout_signature(ws: Any) -> dict[str, Any]:
    """单个工作表的布局签名：合并单元格、行高、列宽。"""
    row_dimensions = {
        index: (dimension.height, bool(dimension.hidden), dimension.outlineLevel, bool(dimension.collapsed))
        for index, dimension in ws.row_dimensions.items()
        if dimension.height is not None or dimension.hidden or dimension.outlineLevel or dimension.collapsed
    }
    column_dimensions = {
        key: (dimension.width, bool(dimension.hidden), dimension.outlineLevel, bool(dimension.collapsed))
        for key, dimension in ws.column_dimensions.items()
        if dimension.width is not None or dimension.hidden or dimension.outlineLevel or dimension.collapsed
    }
    return {
        "merges": tuple(sorted(str(item) for item in ws.merged_cells.ranges)),
        "rows": row_dimensions,
        "columns": column_dimensions,
    }


def template_layout_signature(wb: Any) -> dict[str, dict[str, Any]]:
    """记录不应因回填而变化的模板布局属性。"""
    return {ws.title: sheet_layout_signature(ws) for ws in wb.worksheets}


def assert_template_layout_preserved(
    wb: Any,
    original_signature: dict[str, dict[str, Any]],
    month: int,
) -> None:
    for original_name, expected_signature in original_signature.items():
        output_name = f"项目分类明细汇总表{month or 1}月" if original_name == TEMPLATE_MONTH_SHEET else original_name
        if output_name not in wb.sheetnames:
            raise RuleError(f"输出模板缺少固定工作表：{output_name}")
        worksheet = wb[output_name]
        actual_signature = sheet_layout_signature(worksheet)
        if actual_signature != expected_signature:
            raise RuleError(f"{output_name} 的合并单元格、行高或列宽偏离固定模板")


def write_template_output(
    hours_path: Path,
    hours_layout: Any,
    hours_result: dict[str, Any],
    projects: OrderedDict[str, Project],
    master_costs: dict[str, dict[str, Decimal]],
    labor_totals: dict[str, Decimal],
    material_direct_totals: dict[str, Decimal],
    material_public_totals: dict[str, Decimal],
    effective_records: list[LedgerRecord],
    expense_detail_records: list[LedgerRecord],
    summary_source: tuple[Path, str, str] | None,
    month: int,
    output_path: Path,
    warnings: list[str],
) -> list[str]:
    if not OUTPUT_TEMPLATE.is_file():
        raise RuleError(f"缺少固定输出模板：{OUTPUT_TEMPLATE.name}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUTPUT_TEMPLATE, output_path)
    wb = load_workbook(output_path, data_only=False, read_only=False)
    original_sheets = list(wb.sheetnames)
    original_layout_signature = template_layout_signature(wb)
    clear_broken_formula_cells(wb, warnings)
    write_template_main_sheet(wb, projects, master_costs)
    if summary_source is not None:
        copy_summary_to_template(wb, summary_source, effective_records, month, warnings)
    write_template_material_sheet(
        wb,
        [record for record in effective_records if record.fee_code == DIRECT_MATERIAL_CODE],
        month,
    )
    write_template_detail_sheet(wb, output_year(effective_records))
    write_template_hours_sheets(wb, hours_path, hours_layout, hours_result, projects, month)
    detail_records_by_fee: dict[str, list[LedgerRecord]] = defaultdict(list)
    for record in expense_detail_records:
        detail_records_by_fee[record.fee_code].append(record)

    # 固定模板是输出契约：只迁移模板已预留的费用分配表块，不新增工作表或
    # 自定义表格；迁移后的表头统一位于实际明细内容后的第4行。
    allocation_header_rows: dict[str, int] = {}
    for rule in COST_RULES["fee_items"]:
        ws = safe_sheet(wb, rule["detail_sheet"])
        if rule["code"] == code_key("FYXM02_SYS"):
            # F82:H102 是模板中的可见运费结果样式源，生成时会把整个表块
            # 原样迁移到实际明细内容之后，保留3个空白行。
            header_row = 82
            if ws.max_row < 102 or ws.max_column < 8:
                raise RuleError("固定模板的1.运费缺少可见费用分配预留区")
            allocation_header_rows[rule["code"]] = header_row
        else:
            table = standard_fee_allocation_table(ws)
            if table is None:
                raise RuleError(f"固定模板的{ws.title}缺少费用分配结果预留区")
            allocation_header_rows[rule["code"]] = table[0]

    write_template_expense_source(
        safe_sheet(wb, DIRECT_MATERIAL_RULE["detail_sheet"]),
        detail_records_by_fee.get(DIRECT_MATERIAL_CODE, []),
        month,
    )
    write_template_material_allocation(
        wb,
        projects,
        material_direct_totals,
        material_public_totals,
        master_costs[norm("直接材料")],
        month,
    )
    material_ws = safe_sheet(wb, "14.材料明细账")
    material_header = relocate_material_allocation_tables(
        material_ws,
        blank_rows=3,
        content_last_row=fee_source_content_last_row(
            detail_records_by_fee.get(DIRECT_MATERIAL_CODE, [])
        ),
    )
    if not material_header:
        raise RuleError("固定模板的14.材料明细账缺少三组材料分配结果预留区域")
    original_layout_signature["14.材料明细账"] = sheet_layout_signature(material_ws)
    missing_detail_names: list[str] = []
    for rule in COST_RULES["fee_items"]:
        ws = safe_sheet(wb, rule["detail_sheet"])
        detail_rows = detail_records_by_fee.get(rule["code"], [])
        # 每个费用页的上方账区都只由用户上传的费用明细账承载，避免模板中
        # 可能存在的历史交易被误认为本次输入。
        content_last_row = fee_source_content_last_row(detail_rows)
        if rule["code"] == code_key("FYXM02_SYS"):
            relocated_header = relocate_freight_visible_table(ws, content_last_row + 4)
        else:
            relocated_header = relocate_standard_fee_allocation_table(
                ws,
                blank_rows=3,
                content_last_row=content_last_row,
            )
        if not relocated_header:
            raise RuleError(f"固定模板的{ws.title}缺少可迁移的费用分配结果预留区域")
        allocation_header_rows[rule["code"]] = relocated_header
        write_template_expense_source(
            ws,
            detail_rows,
            month,
            relocated_header,
        )
        original_layout_signature[rule["detail_sheet"]] = sheet_layout_signature(ws)
        if not detail_rows:
            missing_detail_names.append(rule["name"])
    if missing_detail_names:
        warnings.append(
            "未上传以下费用明细账，上方原始账区已留空但下方分配表仍按核算维度明细账计算："
            + "、".join(missing_detail_names)
        )
    # 费用金额只取核算维度明细账的项目分类汇总；上传的费用明细账只保留在
    # 对应费用页上方，不能覆盖或校验计算金额。
    write_template_other_fee_allocations(
        wb,
        projects,
        master_costs,
        month,
        allocation_header_rows,
    )
    write_template_labor_sheet(wb, projects, labor_totals, month)
    for special_sheet_name in ("2.直接人工", "11.社保明细账", "12.公积金", "13.工资明细账"):
        special_ws = safe_sheet(wb, special_sheet_name)
        relocated_header = relocate_standard_fee_allocation_table(special_ws, blank_rows=3)
        if not relocated_header:
            raise RuleError(f"固定模板的{special_sheet_name}缺少费用分配结果预留区域")
        original_layout_signature[special_sheet_name] = sheet_layout_signature(special_ws)
    quote_internal_sheet_references(wb)
    expected = [f"项目分类明细汇总表{month or 1}月" if name == TEMPLATE_MONTH_SHEET else name for name in original_sheets]
    if wb.sheetnames != expected:
        raise RuleError("输出工作表与固定模板不一致")
    assert_formula_errors_cleared(wb)
    assert_template_layout_preserved(wb, original_layout_signature, month)
    # 静态金额已由引擎核定；仍启用自动计算以刷新模板中保留的正常公式。
    if wb.calculation is None:
        wb.calculation = CalcProperties()
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(output_path)
    return wb.sheetnames


def order_sheets(wb: Any, month: int, ledger_sheets: list[str]) -> None:
    desired = [
        "研发费用材料项目台账", "研发项目主表", f"项目分类明细汇总表{month or 1}月", "研发明细表", "研发人员工时统计",
        "研发人员工时统计合计", "RD工时加公摊", "研发人员分配率（直接人工）",
    ] + ledger_sheets + ["2.直接人工"]
    wb._sheets = [wb[name] for name in desired if name in wb.sheetnames] + [ws for ws in wb.worksheets if ws.title not in desired]


def expected_template_sheets(month: int, ledger_sheets: list[str] | None = None) -> list[str]:
    return [
        "研发费用材料项目台账",
        "研发项目主表",
        f"项目分类明细汇总表{month or 1}月",
        "研发明细表",
        "研发人员工时统计",
        "研发人员工时统计合计",
        "1.运费", "2.直接人工", "3.折旧", "4.业务招待费", "5.差旅费", "7.办公费",
        "6.交通费", "8.会展费", "9.租金", "10.福利费", "11.社保明细账", "12.公积金",
        "13.工资明细账", "14.材料明细账", "15.长期待摊费用", "16.维修费", "17.无形资产摊销",
        "18.招聘费", "19.服务费", "20.专利服务费", "21.水电费", "22.咨询费", "23.测试费",
        "24.职工教育经费", "25.劳动保护费",
    ]


def record_month(record: LedgerRecord) -> int:
    value = record.period
    if isinstance(value, (int, float)) and 1 <= int(value) <= 12:
        return int(value)
    match = re.search(r"(?:^|\D)(1[0-2]|[1-9])(?:月|$)", norm(value))
    if match:
        return int(match.group(1))
    if isinstance(record.posting_date, (date, datetime)):
        return record.posting_date.month
    return 0


def select_effective_records(records: list[LedgerRecord], month: int) -> list[LedgerRecord]:
    """优先使用期间一致的凭证明细，并与汇总表交叉勾稽。

    只有汇总表时，汇总表本身是可追溯的费用来源；一旦同时上传逐笔明细，
    则不允许把 2、3 月交易混入 1 月或与汇总金额不一致后仍继续出表。
    """
    selected: list[LedgerRecord] = []
    for fee_code in FEE_RULES_BY_CODE:
        summary_rows = [row for row in records if row.fee_code == fee_code and row.source_kind == "summary"]
        ledger_rows = [row for row in records if row.fee_code == fee_code and row.source_kind == "ledger"]
        if ledger_rows:
            blank_project_rows = [row for row in ledger_rows if not PROJECT_RE.fullmatch(row.project_code)]
            if blank_project_rows:
                summary_is_public = summary_rows and all(row.project_code == "RD9999" for row in summary_rows)
                if len(blank_project_rows) == len(ledger_rows) and summary_is_public:
                    for row in ledger_rows:
                        row.project_code, row.project_name = "RD9999", "研发公摊"
                else:
                    locations = "、".join(f"{row.sheet}!{row.source_row}" for row in blank_project_rows[:5])
                    raise RuleError(
                        f"{FEE_RULES_BY_CODE[fee_code]['name']}明细账缺少研发项目：{locations}；"
                        "仅当同费用分类汇总表明确为RD9999研发公摊时才可按公摊处理"
                    )
            unknown_periods = [row for row in ledger_rows if not record_month(row)]
            if unknown_periods:
                locations = "、".join(f"{row.sheet}!{row.source_row}" for row in unknown_periods[:5])
                raise RuleError(f"{FEE_RULES_BY_CODE[fee_code]['name']}明细账缺少可识别期间：{locations}")
            scoped_rows = [row for row in ledger_rows if record_month(row) == month]
            if summary_rows:
                summary_total = money(sum((row.amount for row in summary_rows), Decimal("0")))
                detail_total = money(sum((row.amount for row in scoped_rows), Decimal("0")))
                if summary_total != detail_total:
                    raise RuleError(
                        f"{FEE_RULES_BY_CODE[fee_code]['name']}的{month}月分类汇总为{summary_total}，"
                        f"同月明细账为{detail_total}，差额{money(summary_total - detail_total)}；"
                        "请提供同一期间的资料后再测算"
                    )
            selected.extend(scoped_rows)
        else:
            selected.extend(summary_rows)
    return selected


def locate_summary_source(ledger_paths: list[Path]) -> tuple[Path, str] | None:
    for path in ledger_paths:
        wb = load_workbook(path, data_only=True, read_only=False)
        for ws in wb.worksheets:
            headers = {norm(label): col for label, col in cell_texts(ws, 1, 40).items()}
            required = [norm(item) for item in COST_RULES.get("summary_input", {}).get("required_headers", [])]
            if required and all(item in headers for item in required):
                return path, ws.title
    return None


def has_ledger_layout(path: Path) -> bool:
    """判断工作簿是否含凭证明细结构，不依赖文件名或工作表名称。"""
    wb = load_workbook(path, data_only=True, read_only=False)
    for ws in wb.worksheets:
        for row in range(1, min(ws.max_row, 30) + 1):
            headers = cell_texts(ws, row, 80)
            if "摘要" in headers and ("借方" in headers or "贷方" in headers):
                return True
    return False


def dimension_source_signature(path: Path) -> tuple[Any, ...]:
    """为核算维度汇总/组合表生成内容指纹，用于去掉重复上传的同源副本。"""
    warnings: list[str] = []
    summary_source = locate_summary_source([path])
    if summary_source is not None:
        source_book = load_workbook(path, data_only=True, read_only=False)
        records, month = parse_summary_sheet(path, source_book[summary_source[1]], warnings)
    else:
        combination_source = locate_combination_summary_sheet(path)
        if combination_source is None:
            return ("not_dimension_source", str(path.resolve()))
        source_book = load_workbook(path, data_only=True, read_only=False)
        records, month = parse_combination_summary_sheet(
            path,
            source_book[combination_source[1]],
            combination_source[2],
            warnings,
        )
    if not records:
        return ("empty_dimension_source", str(path.resolve()))
    return (
        month,
        tuple(
            (
                record.fee_code,
                record.fee_name,
                record.project_code,
                record.project_name,
                str(record.amount),
            )
            for record in records
        ),
    )


def deduplicate_equivalent_dimension_ledgers(paths: list[Path]) -> list[Path]:
    unique: list[Path] = []
    seen: set[tuple[Any, ...]] = set()
    for path in paths:
        try:
            signature = dimension_source_signature(path)
        except (RuleError, OSError, ValueError):
            signature = ("unreadable_dimension_source", str(path.resolve()))
        if signature in seen:
            continue
        seen.add(signature)
        unique.append(path)
    return unique


def classify_input_workbooks(paths: list[Path]) -> dict[str, Any]:
    """根据表结构归类；文件名只用于组合表提示，不是完整名称硬门槛。"""
    roles: dict[str, list[Path]] = {
        "hours": [],
        "salary": [],
        "dimension_ledger": [],
        "expense_ledgers": [],
    }
    unknown: list[Path] = []
    for path in paths:
        summary_candidate = locate_summary_source([path])
        combination_candidate = locate_combination_summary_sheet(path)
        is_dimension_ledger = summary_candidate is not None or combination_candidate is not None
        if is_dimension_ledger:
            roles["dimension_ledger"].append(path)
            # 同一工作簿可以同时含核算维度汇总表和费用凭证明细表。
            # 原逻辑在识别汇总表后立即continue，会漏掉同文件中的运费明细。
            if has_ledger_layout(path):
                roles["expense_ledgers"].append(path)
            continue
        try:
            _, hours_result, _ = calculate_hours(path, OrderedDict())
            if hours_result.get("people"):
                roles["hours"].append(path)
                continue
        except RuleError:
            pass
        try:
            salary_rows, _, _ = read_salary(path)
            if salary_rows:
                roles["salary"].append(path)
                continue
        except RuleError:
            pass
        if has_ledger_layout(path):
            roles["expense_ledgers"].append(path)
            continue
        unknown.append(path)

    roles["dimension_ledger"] = deduplicate_equivalent_dimension_ledgers(roles["dimension_ledger"])
    core_labels = {
        "工时表": "hours",
        "工资表": "salary",
        "核算维度汇总/组合表": "dimension_ledger",
    }
    missing = [
        label
        for label, key in core_labels.items()
        if not roles[key] and (key != "dimension_ledger" or not roles["expense_ledgers"])
    ]
    duplicates = [label for label, key in core_labels.items() if len(roles[key]) > 1]
    if missing or duplicates or unknown:
        messages: list[str] = []
        if missing:
            messages.append("未识别到" + "、".join(missing))
        if duplicates:
            messages.append("识别到多份" + "、".join(duplicates))
        if unknown:
            messages.append("无法归类：" + "、".join(path.name for path in unknown))
        raise RuleError(
            "；".join(messages)
            + "。请保留一份核心工时、工资和核算维度汇总/组合工作簿；费用明细账可多份上传，表内带费用分类会自动识别"
        )
    return {
        "hours": roles["hours"][0],
        "salary": roles["salary"][0],
        "dimension_ledger": roles["dimension_ledger"][0] if roles["dimension_ledger"] else roles["expense_ledgers"],
        "expense_ledgers": roles["expense_ledgers"],
    }


def aggregate_dimension_detail_records(records: list[LedgerRecord]) -> list[LedgerRecord]:
    grouped: OrderedDict[tuple[str, str, str, str], Decimal] = OrderedDict()
    first_records: dict[tuple[str, str, str, str], LedgerRecord] = {}
    for record in records:
        key = (record.fee_code, record.fee_name, record.project_code, record.project_name)
        grouped[key] = grouped.get(key, Decimal("0")) + record.amount
        first_records.setdefault(key, record)
    aggregated: list[LedgerRecord] = []
    for key, amount in grouped.items():
        first = first_records[key]
        fee_code, fee_name, project_code, project_name = key
        aggregated.append(LedgerRecord(
            source_path=first.source_path,
            sheet=first.sheet,
            fee_code=fee_code,
            fee_name=fee_name,
            project_code=project_code,
            project_name=project_name,
            posting_date=first.posting_date,
            fiscal_year=first.fiscal_year,
            period=record_month(first),
            voucher_word="",
            voucher_no=None,
            summary=f"{fee_name}明细账汇总",
            debit=money(amount),
            credit=Decimal("0"),
            source_kind="summary",
            source_row=0,
        ))
    return aggregated


def read_dimension_detail_workbooks(paths: list[Path], warnings: list[str]) -> tuple[list[LedgerRecord], int, tuple[Path, str, str]]:
    records: list[LedgerRecord] = []
    explicit_months: set[int] = set()
    for path in paths:
        parsed, inferred_month, _ = parse_ledger_file(path, warnings)
        if inferred_month:
            explicit_months.add(inferred_month)
        records.extend(record for record in parsed if record.source_kind == "ledger")
    if not records:
        raise RuleError("未上传核算维度汇总/组合表，且费用明细账未识别到可测算的费用项目")
    row_months = {record_month(record) for record in records if record_month(record)}
    months = explicit_months or row_months
    if len(months) > 1:
        labels = "、".join(f"{month}月" for month in sorted(months))
        raise RuleError(f"费用明细账包含多个明确期间：{labels}；请每次上传同一期间的资料")
    month = next(iter(months), 0)
    warnings.append("未上传核算维度汇总/组合表，已按费用明细账的费用项目和研发项目汇总作为测算金额来源")
    return aggregate_dimension_detail_records(records), month, (paths[0], "", "generated")


def read_dimension_summary_workbook(path: Path | list[Path], warnings: list[str]) -> tuple[list[LedgerRecord], int, tuple[Path, str, str]]:
    """读取核算维度金额来源。

    优先使用项目分类汇总表或“核算维度与科目组合”汇总表；若用户只上传
    带费用项目分类的核算维度明细账，则按费用项目+研发项目汇总后作为
    金额来源，同时仍把原始凭证明细写入各费用页上方账区。
    """
    if isinstance(path, list):
        return read_dimension_detail_workbooks(path, warnings)

    summary_source = locate_summary_source([path])
    if summary_source is not None:
        source_book = load_workbook(path, data_only=True, read_only=False)
        source_ws = source_book[summary_source[1]]
        records, month = parse_summary_sheet(path, source_ws, warnings)
        if records is None or not records:
            raise RuleError("核算维度明细账的项目分类汇总表未识别到可测算的费用项目")
        return records, month, (summary_source[0], summary_source[1], "summary")

    combination_source = locate_combination_summary_sheet(path)
    if combination_source is not None:
        source_book = load_workbook(path, data_only=True, read_only=False)
        source_ws = source_book[combination_source[1]]
        records, month = parse_combination_summary_sheet(
            path, source_ws, combination_source[2], warnings
        )
        if not records:
            raise RuleError("核算维度与科目组合表未识别到可测算的费用项目")
        return records, month, (path, source_ws.title, "combination")

    raise RuleError("未找到项目分类汇总表或核算维度与科目组合表，无法测算费用")


def read_expense_detail_workbooks(
    paths: list[Path],
    warnings: list[str],
) -> tuple[list[LedgerRecord], int]:
    """读取用户上传的费用明细账，仅供复制至对应费用模板页。

    明细金额不参与费用计算；但“期间:YYYY年MM月”是已整理组合表删除
    前置报表信息后唯一可追溯的输出期间。因此单独返回明确期间，只用于
    命名和标注固定模板，不改变金额来源。
    """
    records: list[LedgerRecord] = []
    explicit_months: set[int] = set()
    for path in paths:
        warning_count = len(warnings)
        parsed, inferred_month, _ = parse_ledger_file(path, warnings)
        if inferred_month:
            explicit_months.add(inferred_month)
        detail_rows = [record for record in parsed if record.source_kind == "ledger"]
        if not detail_rows:
            if len(warnings) == warning_count:
                warnings.append(f"{path.name} 未识别到可写入费用页的凭证明细")
            continue
        records.extend(detail_rows)
    if len(explicit_months) > 1:
        labels = "、".join(f"{month}月" for month in sorted(explicit_months))
        raise RuleError(f"费用明细账包含多个明确期间：{labels}；请每次上传同一期间的资料")
    return records, next(iter(explicit_months), 0)


def build_output(
    hours_path: Path,
    salary_path: Path,
    dimension_ledger_path: Path | list[Path],
    expense_detail_paths: list[Path],
    output_path: Path,
) -> dict[str, Any]:
    dimension_paths = dimension_ledger_path if isinstance(dimension_ledger_path, list) else [dimension_ledger_path]
    source_paths = [hours_path, salary_path, *dimension_paths, *expense_detail_paths]
    if any(output_path.resolve() == source_path.resolve() for source_path in source_paths):
        raise RuleError("输入文件只读保留，结果必须保存为新的工作簿")
    warnings: list[str] = []
    salary_project_names = read_salary_project_names(salary_path)
    project_name_hints = OrderedDict(
        (code, Project(code, name)) for code, name in salary_project_names.items()
    )
    layout, hours_result, hour_projects = calculate_hours(hours_path, project_name_hints)
    # 研发项目主表由系统生成：先取工时项目，再补充核算维度中直接归属且
    # 不是已有母项目的末级项目；母项目金额仍按其工时子项目分配。
    projects: OrderedDict[str, Project] = OrderedDict(
        (code, Project(project.code, project.name, project.final_hours, project.rate))
        for code, project in hour_projects.items()
    )
    for code, project_name in salary_project_names.items():
        if code in projects and (not projects[code].name or projects[code].name == code):
            projects[code].name = project_name
        if code in hour_projects and (not hour_projects[code].name or hour_projects[code].name == code):
            hour_projects[code].name = project_name

    parsed_records, ledger_month, summary_source = read_dimension_summary_workbook(dimension_ledger_path, warnings)
    for record in parsed_records:
        code = record.project_code
        if code == PUBLIC_PROJECT_CODE:
            continue
        if code in projects:
            if record.project_name and (not projects[code].name or projects[code].name == code):
                projects[code].name = record.project_name
            continue
        if any(existing.startswith(code + "-") for existing in projects):
            continue
        projects[code] = Project(code, record.project_name or code)
    if len(projects) > 18:
        raise RuleError(f"工时及核算维度合计识别到{len(projects)}个末级项目，超出固定模板18个槽位")

    for item in hours_result["projects"]:
        if item["project_code"] in projects:
            item["project_name"] = projects[item["project_code"]].name

    salary_rows, salary_source, salary_month = read_salary(salary_path)
    hour_codes = list(hour_projects)
    salary_allocations, labor_totals, salary_total, matched_salary, unmatched_salary = allocate_salary(
        salary_rows, hours_result, hour_codes, warnings
    )
    for code in projects:
        labor_totals.setdefault(code, Decimal("0"))

    expense_detail_records, expense_detail_month = read_expense_detail_workbooks(
        expense_detail_paths, warnings
    )
    month = ledger_month or expense_detail_month or salary_month or 1
    hours_month = hours_workbook_month(hours_path, layout.sheet)
    if hours_month and hours_month != month:
        warnings.append(
            f"费用明细账明确为{month}月，工时表表头显示{hours_month}月；"
            f"固定模板期间按{month}月输出，工时数据仍按上传内容分配"
        )
    if salary_month and salary_month != month:
        warnings.append(
            f"费用明细账明确为{month}月，工资表表头显示{salary_month}月；"
            f"固定模板期间按{month}月输出，工资金额仍按上传内容分配"
        )
    effective_records = parsed_records

    records_by_fee: dict[str, list[LedgerRecord]] = defaultdict(list)
    for record in effective_records:
        records_by_fee[record.fee_code].append(record)
    master_costs: dict[str, dict[str, Decimal]] = {
        norm(header): {code: Decimal("0") for code in projects}
        for header in MAIN_HEADERS[2:-1]
    }
    for code, amount in labor_totals.items():
        master_costs[norm("人员人工（含社保+公积金）")][code] = money(amount)

    material_records = records_by_fee.get(DIRECT_MATERIAL_CODE, [])
    material_direct_totals = {code: Decimal("0") for code in projects}
    material_public_totals = {code: Decimal("0") for code in projects}
    direct_material_records = [record for record in material_records if record.project_code != PUBLIC_PROJECT_CODE]
    public_material_records = [record for record in material_records if record.project_code == PUBLIC_PROJECT_CODE]
    if direct_material_records:
        material_direct_totals, _, _, _, _ = allocate_expense_records(
            direct_material_records, projects, warnings, "直接材料"
        )
    if public_material_records:
        material_public_totals, _, _, _, _ = allocate_expense_records(
            public_material_records, projects, warnings, "直接材料公摊"
        )
    for code in projects:
        master_costs[norm("直接材料")][code] = money(
            material_direct_totals.get(code, Decimal("0"))
            + material_public_totals.get(code, Decimal("0"))
        )

    freight_totals: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
    other_total = Decimal("0")
    for rule in COST_RULES["fee_items"]:
        source_records = records_by_fee[rule["code"]]
        if not source_records:
            continue
        allocations, pool, direct, result_total, _ = allocate_expense_records(
            source_records,
            projects,
            warnings,
            rule["name"],
            use_displayed_rates=(rule["code"] == code_key("FYXM02_SYS")),
        )
        target_costs = master_costs[norm(rule["master_header"])]
        for code, amount in allocations.items():
            target_costs[code] = money(target_costs.get(code, Decimal("0")) + amount)
        other_total += result_total
        if rule["code"] == code_key("FYXM02_SYS"):
            for code, amount in allocations.items():
                freight_totals[code] += amount
    for code in projects:
        freight_totals.setdefault(code, Decimal("0"))
        labor_totals.setdefault(code, Decimal("0"))

    audit_data = {
        "hours_projects": len(hour_projects),
        "rate_sum": sum((project.rate for project in hour_projects.values()), Decimal("0")),
        "salary_total": salary_total,
        "matched_salary": matched_salary,
        "labor_total": sum(labor_totals.values(), Decimal("0")),
        "unmatched_salary": unmatched_salary,
        "freight_total": sum(freight_totals.values(), Decimal("0")),
        "other_total": other_total,
        "master_projects": len(projects),
        "salary_source": salary_source,
        "ledger_records": len(expense_detail_records),
    }
    output_sheets = write_template_output(
        hours_path,
        layout,
        hours_result,
        projects,
        master_costs,
        labor_totals,
        material_direct_totals,
        material_public_totals,
        effective_records,
        expense_detail_records,
        summary_source,
        month,
        output_path,
        warnings,
    )
    return {**audit_data, "month": month, "warnings": warnings, "output_sheets": output_sheets}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--reference-input", nargs=2, action="append", default=[])
    parser.add_argument("--rules", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--report-json", type=Path, required=True)
    args = parser.parse_args(argv)
    references: dict[str, list[Path]] = defaultdict(list)
    for key, value in args.reference_input:
        references[key].append(Path(value))
    missing = [key for key in ("salary_file", "dimension_ledger", "expense_detail") if not references.get(key)]
    if missing:
        print("缺少输入：" + "、".join(missing), file=sys.stderr)
        return 2
    try:
        result = build_output(
            args.input,
            references["salary_file"][0],
            references["dimension_ledger"][0],
            references["expense_detail"],
            args.output,
        )
        report = {
            "contract_version": 1,
            "status": "PASS",
            "rule_pack": PACK_INFO,
            "summary": {
                "title": "一键综合核算结果",
                "subtitle": "已按顺序完成工时分配率、直接人工和运费",
                "metrics": [
                    {"key": "projects", "label": "主表项目", "value": result["master_projects"], "format": "integer"},
                    {"key": "employees", "label": "工资人员", "value": len(read_salary(references["salary_file"][0])[0]), "format": "integer"},
                    {"key": "dimension_ledger", "label": "核算维度明细账", "value": 1, "format": "integer"},
                ],
                "totals": [
                    {"key": "labor", "label": "已分配直接人工", "value": float(result["labor_total"]), "format": "number"},
                    {"key": "freight", "label": "运费结果", "value": float(result["freight_total"]), "format": "number", "emphasis": True},
                    {"key": "other", "label": "其他费用合计", "value": float(result["other_total"]), "format": "number", "emphasis": True},
                    {"key": "unmatched", "label": "未分配工资", "value": float(result["unmatched_salary"]), "format": "number"},
                ],
            },
            "warnings": result["warnings"],
            "details": result,
        }
        args.report_json.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        return 0
    except (RuleError, OSError, ValueError) as exc:
        print(f"一键综合核算停止：{exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
