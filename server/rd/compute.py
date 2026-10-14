"""Isolated Excel worker; inputs are read-only, results use the fixed template."""
import json, sys, re, hashlib, zipfile, warnings
warnings.filterwarnings("ignore", message="Workbook contains no default style.*", category=UserWarning)
from pathlib import Path
root = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(root / 'public/py/vendor.zip'))
sys.path.insert(0, str(root / 'public/py'))
from openpyxl import Workbook, load_workbook
from rd_cost_workflow_v1 import engine
engine.OUTPUT_TEMPLATE = root / 'public/py/output_template.xlsx'
from collections import OrderedDict

def check_xlsx(path):
    with zipfile.ZipFile(path) as book:
        if sum(i.file_size for i in book.infolist()) > 512 * 1024 * 1024:
            raise ValueError('Excel 解压后超过 512 MB，请拆分资料。')
        if len(book.infolist()) > 20000:
            raise ValueError('Excel 内部文件过多，请拆分资料。')
        if 'xl/workbook.xml' not in book.namelist():
            raise ValueError('不是有效的 Excel 文件。')

def main():
    mode, manifest_path = sys.argv[1:]
    data = json.loads(Path(manifest_path).read_text())
    if mode == 'image':
        kind, rows = data['kind'], data['rows']
        if kind not in ('hours','salary','balance','ledger') or not rows or len(rows)>1000:
            raise ValueError('图片复核数据无效。')
        book = Workbook(); sheet = book.active
        sheet.title = {'hours':'工时','salary':'工资','balance':'科目','ledger':'明细账'}[kind]
        for row in rows:
            if not isinstance(row,list) or len(row)>100: raise ValueError('图片表格过大。')
            converted=[]
            for value in row:
                if value is not None and not isinstance(value,(str,int,float)): raise ValueError('图片单元格无效。')
                if isinstance(value,str):
                    value=value.strip()
                    if value.startswith(('=','+','@')) and not re.fullmatch(r'[+-]?\d+(?:\.\d+)?',value):
                        raise ValueError('图片表格不能包含公式，请填写原图实际数值。')
                    if re.fullmatch(r'-?\d+(?:\.\d+)?',value) and not re.match(r'^0\d',value): value=float(value)
                converted.append(value)
            sheet.append(converted)
        path=Path(data['output']);book.save(path)
        if kind=='hours': valid=bool(engine.calculate_hours(path,OrderedDict())[1].get('people'))
        elif kind=='salary': valid=bool(engine.read_salary(path)[0])
        elif kind=='balance': valid=engine.locate_summary_source([path]) is not None or engine.locate_combination_summary_sheet(path) is not None
        else: valid=engine.has_ledger_layout(path)
        if not valid: raise ValueError('复核表格未被测算规则识别，请补全表头、月份和必需列。')
        return
    # Template is checked before inspecting any business inputs.
    template=load_workbook(engine.OUTPUT_TEMPLATE,read_only=True)
    try:
        if not {'研发项目主表','研发人员工时统计',engine.TEMPLATE_MONTH_SHEET}.issubset(template.sheetnames):
            raise ValueError('固定输出模板结构不完整。')
    finally: template.close()
    paths=[Path(item['path']) for item in data['files'] if item['compute']]
    if not paths: raise ValueError('请上传 Excel 或已复核图片。')
    template_hash=hashlib.sha256(engine.OUTPUT_TEMPLATE.read_bytes()).hexdigest()
    before={str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in paths}
    for p in paths: check_xlsx(p)
    years=set()
    for p in paths: years.update(engine.input_workbook_period(p)[1])
    if len(years)>1: raise ValueError('上传资料包含多个会计年度，请按年度分批测算。')
    year=next(iter(years),data['year'])
    result=engine.build_outputs(paths,Path(data['output']),accounting_year=year)
    for p in paths:
        if hashlib.sha256(p.read_bytes()).hexdigest()!=before[str(p)]: raise ValueError('输入文件完整性核验失败。')
    if hashlib.sha256(engine.OUTPUT_TEMPLATE.read_bytes()).hexdigest()!=template_hash: raise ValueError('固定模板完整性核验失败。')
    Path(data['report']).write_text(json.dumps({'year':year,**result},ensure_ascii=False,default=str))

if __name__=='__main__':
    try: main()
    except Exception as exc:
        print(str(exc),file=sys.stderr);sys.exit(1)
